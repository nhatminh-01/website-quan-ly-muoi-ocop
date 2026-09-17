"""Safe import of the historical OCOP workbook into PostgreSQL.

The source workbook is preserved row-by-row in staging before normalized data
is published. Unknown administrative units are never invented: they remain
visible as validation errors until the shared administrative catalog maps them.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
import hashlib
import io
import json
import re

from permissions import is_chi_cuc_user
import ocop_registry


OcopImportError = ocop_registry.RegistryError


DEFAULT_SHEET = "Loc"
MAX_FILE_BYTES = 20 * 1024 * 1024


_text = ocop_registry.text


_EXCEL_ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}


def _context_text(value) -> str:
    """Read a subject-level cell, treating broken Excel formulas as blank.

    Historical workbooks use formulas in merged continuation rows.  When a
    referenced row was deleted Excel stores ``#REF!`` in those cells; it is
    not a real subject or administrative unit and must not replace the last
    valid merged-cell context.
    """
    text = _text(value)
    return "" if text.upper() in _EXCEL_ERROR_VALUES else text


_key = ocop_registry.key


def _json_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _star(value, label: str, errors: list[str], *, required=True):
    if value in (None, ""):
        if required:
            errors.append(f"{label}: thiếu hạng sao.")
        return None
    try:
        star = int(float(str(value).strip()))
    except (TypeError, ValueError):
        errors.append(f"{label}: hạng sao không hợp lệ ({value!r}).")
        return None
    if star not in (3, 4, 5):
        errors.append(f"{label}: hạng sao phải là 3, 4 hoặc 5.")
        return None
    return star


def _year(value, label: str, errors: list[str], *, required=True):
    if value in (None, ""):
        if required:
            errors.append(f"{label}: thiếu năm công nhận.")
        return None
    try:
        year = int(float(str(value).strip()))
    except (TypeError, ValueError):
        errors.append(f"{label}: năm không hợp lệ ({value!r}).")
        return None
    if not 1900 <= year <= 9999:
        errors.append(f"{label}: năm nằm ngoài phạm vi hợp lệ.")
        return None
    return year


def _date_value(value, label: str, errors: list[str], *, required=False):
    if value in (None, ""):
        if required:
            errors.append(f"{label}: thiếu ngày.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return date(1899, 12, 30) + timedelta(days=float(value))
        except (OverflowError, ValueError):
            errors.append(f"{label}: ngày Excel không hợp lệ ({value!r}).")
            return None
    token = _text(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(token, fmt).date()
        except ValueError:
            pass
    errors.append(f"{label}: ngày không hợp lệ ({value!r}).")
    return None


_stable_code = ocop_registry.stable_code
_source_key = ocop_registry.source_key


def _lookup_unit(source_name: str, unit_lookup):
    raw = _text(source_name)
    if not raw:
        return None, []
    candidates = [raw]
    folded = raw.casefold()
    if not folded.startswith(("xã ", "phường ", "thị trấn ")):
        candidates.extend((f"Xã {raw}", f"Phường {raw}", f"Thị trấn {raw}"))
    matches = {}
    for candidate in candidates:
        result = unit_lookup(candidate)
        if result:
            name, code = result
            matches[str(code)] = (name, str(code))
    if len(matches) == 1:
        return next(iter(matches.values())), []
    if len(matches) > 1:
        return None, [f"Địa bàn '{raw}' khớp nhiều mã hành chính; cần bổ sung alias rõ ràng."]
    return None, [f"Địa bàn chưa có mã trong danh mục hành chính: {raw}."]


def parse_ocop_workbook(file_bytes: bytes, filename: str, unit_lookup, sheet_name: str = DEFAULT_SHEET):
    """Parse the legacy OCOP workbook into a serializable preview.

    Entity-level cells that are visually merged/continued are inherited only
    when the subject is the same as the preceding product (or the subject cell
    itself is blank). This prevents accidentally carrying metadata to a new
    subject that merely has a blank optional field.
    """
    if not file_bytes:
        raise OcopImportError("File Excel rỗng.")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise OcopImportError("File Excel quá lớn (tối đa 20 MB).")
    if not str(filename or "").lower().endswith(".xlsx"):
        raise OcopImportError("Chỉ hỗ trợ file .xlsx.")
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise OcopImportError("Máy chủ chưa cài openpyxl.") from exc
    try:
        raw_book = load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
        value_book = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise OcopImportError(f"Không đọc được file Excel: {exc}") from exc

    try:
        selected = _text(sheet_name) or DEFAULT_SHEET
        if selected not in value_book.sheetnames:
            raise OcopImportError(
                f"Không tìm thấy sheet '{selected}'. Có: {', '.join(value_book.sheetnames)}"
            )
        raw_sheet = raw_book[selected]
        sheet = value_book[selected]
        rows = []
        entity_context = None
        source_units = set()
        mapped_units = set()
        entity_keys = set()
        product_keys = set()

        raw_rows = raw_sheet.iter_rows(min_row=5, max_col=22, values_only=True)
        value_rows = sheet.iter_rows(min_row=5, max_col=22, values_only=True)
        for excel_row, (raw_values, values) in enumerate(zip(raw_rows, value_rows), 5):
            values = list(values)
            product_name = _text(values[1])
            source_tt = _text(values[0])
            if not product_name:
                continue
            # Summary/footer rows do not carry a product sequence number.
            if not source_tt or not re.fullmatch(r"\d+(?:\.0+)?", source_tt):
                continue

            errors: list[str] = []
            warnings: list[str] = []
            entity_name_raw = _context_text(values[4])
            current_same_entity = bool(
                entity_context and entity_name_raw and _key(entity_name_raw) == _key(entity_context["entity_name"])
            )
            if entity_name_raw:
                if current_same_entity:
                    context = dict(entity_context)
                    for key, index in (
                        ("business_type", 3), ("entity_name", 4), ("source_unit_name", 5),
                        ("address", 6), ("representative_name", 7), ("phone", 8),
                    ):
                        incoming = _context_text(values[index])
                        if incoming:
                            context[key] = incoming
                    entity_context = context
                else:
                    entity_context = {
                        "business_type": _context_text(values[3]),
                        "entity_name": entity_name_raw,
                        "source_unit_name": _context_text(values[5]),
                        "address": _context_text(values[6]),
                        "representative_name": _context_text(values[7]),
                        "phone": _context_text(values[8]),
                    }
            elif entity_context:
                # A blank subject is the workbook's merged-cell continuation convention.
                entity_context = dict(entity_context)
            else:
                entity_context = {
                    "business_type": "", "entity_name": "", "source_unit_name": "",
                    "address": "", "representative_name": "", "phone": "",
                }
                errors.append("Không xác định được chủ thể cho dòng sản phẩm này.")

            context = dict(entity_context)
            entity_name = _text(context.get("entity_name"))
            source_unit_name = _text(context.get("source_unit_name"))
            product_group = _text(values[2])
            if not entity_name:
                errors.append("Thiếu tên chủ thể.")
            if not source_unit_name:
                errors.append("Thiếu xã/phường của chủ thể.")
            if not product_group:
                errors.append("Thiếu nhóm sản phẩm.")
            if not context.get("business_type"):
                warnings.append("Chủ thể chưa có loại hình kinh doanh trong file nguồn.")

            official, unit_messages = _lookup_unit(source_unit_name, unit_lookup)
            if unit_messages:
                errors.extend(unit_messages)
            if official:
                official_unit_name, unit_code = official
                mapped_units.add(unit_code)
            else:
                official_unit_name, unit_code = "", None
            if source_unit_name:
                source_units.add(source_unit_name)

            current_star = _star(values[9], "Hạng sao hiện tại", errors)
            expiry_date = _date_value(values[10], "Ngày hết hạn", errors)
            first_star = _star(values[11], "Đánh giá lần 1", errors)
            first_date = _date_value(values[12], "Ngày công nhận lần 1", errors)
            first_year = _year(values[13], "Năm công nhận lần 1", errors)
            first_decision = _text(values[14])
            first_authority = _text(values[15])
            if first_date and first_year and first_date.year != first_year:
                warnings.append(
                    f"Lần 1: ngày công nhận {first_date.isoformat()} không cùng năm với cột năm {first_year}; giữ nguyên cả hai giá trị nguồn."
                )

            second_present = any(values[i] not in (None, "") for i in range(16, 21))
            second_star = second_date = second_year = None
            second_decision = second_authority = ""
            if second_present:
                second_star = _star(values[16], "Đánh giá lần 2/nâng hạng", errors)
                second_date = _date_value(values[17], "Ngày công nhận lần 2", errors)
                second_year = _year(values[18], "Năm công nhận lần 2", errors)
                second_decision = _text(values[19])
                second_authority = _text(values[20])
                if second_date and second_year and second_date.year != second_year:
                    warnings.append(
                        f"Lần 2: ngày công nhận {second_date.isoformat()} không cùng năm với cột năm {second_year}; giữ nguyên cả hai giá trị nguồn."
                    )

            latest_star = second_star if second_present else first_star
            if current_star is not None and latest_star is not None and current_star != latest_star:
                errors.append(
                    f"Hạng sao hiện tại ({current_star}) khác lần công nhận gần nhất ({latest_star}); không tự sửa dữ liệu nguồn."
                )

            entity_natural = f"{_key(source_unit_name)}|{_key(entity_name)}"
            product_natural = f"{entity_natural}|{_key(product_name)}"
            entity_source_key = _source_key("entity", source_unit_name, entity_name)
            product_source_key = _source_key("product", source_unit_name, entity_name, product_name)
            ma_co_so = _stable_code("CS", entity_natural)
            ma_san_pham = _stable_code("OP", product_natural)
            if entity_natural in entity_keys and not current_same_entity and entity_name_raw:
                # Repeated subjects are valid; this message is informational only.
                pass
            entity_keys.add(entity_natural)
            if product_natural in product_keys:
                errors.append("Sản phẩm bị trùng cùng chủ thể và địa bàn trong file nguồn.")
            product_keys.add(product_natural)

            recognitions = []
            if first_star is not None and first_year is not None:
                recognitions.append({
                    "source_key": _source_key(product_source_key, "recognition", "1"),
                    "recognition_sequence": 1,
                    "evaluation_type": "new",
                    "star_rank": first_star,
                    "recognition_date": first_date.isoformat() if first_date else None,
                    "recognition_year": first_year,
                    "decision_number": first_decision,
                    "decision_authority": first_authority,
                    "expiry_date": expiry_date.isoformat() if (not second_present and expiry_date) else None,
                    "is_current": not second_present,
                    "note": _text(values[21]) if not second_present else "",
                })
            if second_present and second_star is not None and second_year is not None:
                evaluation_type = "upgrade" if first_star is not None and second_star > first_star else "re_evaluation"
                recognitions.append({
                    "source_key": _source_key(product_source_key, "recognition", "2"),
                    "recognition_sequence": 2,
                    "evaluation_type": evaluation_type,
                    "star_rank": second_star,
                    "recognition_date": second_date.isoformat() if second_date else None,
                    "recognition_year": second_year,
                    "decision_number": second_decision,
                    "decision_authority": second_authority,
                    "expiry_date": expiry_date.isoformat() if expiry_date else None,
                    "is_current": True,
                    "note": _text(values[21]),
                })

            canonical = {
                "source_tt": source_tt,
                "source_unit_name": source_unit_name,
                "entity": {
                    "source_key": entity_source_key,
                    "ma_co_so": ma_co_so,
                    "name": entity_name,
                    "business_type": _text(context.get("business_type")),
                    "unit_code": unit_code,
                    "unit_name": official_unit_name,
                    "address": _text(context.get("address")),
                    "representative_name": _text(context.get("representative_name")),
                    "phone": _text(context.get("phone")),
                },
                "product": {
                    "source_key": product_source_key,
                    "ma_san_pham": ma_san_pham,
                    "name": product_name,
                    "group": product_group,
                    "current_star": current_star,
                },
                "recognitions": recognitions,
            }
            raw_data = {
                chr(65 + i): {"raw": _json_value(raw_values[i]), "value": _json_value(values[i])}
                for i in range(22)
            }
            status = "error" if errors else ("warning" if warnings else "valid")
            rows.append({
                "excel_row": excel_row,
                "source_tt": source_tt,
                "source_unit_name": source_unit_name,
                "status": status,
                "errors": errors,
                "warnings": warnings,
                "raw_data": raw_data,
                "canonical": canonical,
            })

        if not rows:
            raise OcopImportError("Không tìm thấy dòng sản phẩm OCOP trong sheet đã chọn.")
        return {
            "filename": filename,
            "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
            "sheet_name": selected,
            "rows": rows,
            "total_rows": len(rows),
            "valid_rows": sum(row["status"] == "valid" for row in rows),
            "warning_rows": sum(row["status"] == "warning" for row in rows),
            "error_rows": sum(row["status"] == "error" for row in rows),
            "source_unit_count": len(source_units),
            "mapped_unit_count": len(mapped_units),
            "unknown_units": sorted({row["source_unit_name"] for row in rows if not row["canonical"]["entity"]["unit_code"]}),
            "entity_count": len(entity_keys),
            "product_count": len(product_keys),
            "recognition_count": sum(len(row["canonical"]["recognitions"]) for row in rows),
        }
    finally:
        raw_book.close()
        value_book.close()


def _json(data) -> str:
    return json.dumps(data, ensure_ascii=False, default=str)


def _assert_publishable(preview):
    if preview.get("error_rows"):
        raise OcopImportError(
            f"Còn {preview['error_rows']} dòng lỗi. Chỉ có thể lưu staging; chưa thể đồng bộ OCOP chính thức."
        )


def commit_ocop_preview(con, session, preview, mode="publish"):
    """Persist staging rows and optionally publish normalized OCOP history.

    ``stage_only`` is deliberately allowed when the workbook contains unknown
    administrative units. ``publish`` is all-or-nothing and requires zero errors.
    Caller owns commit/rollback.
    """
    if not is_chi_cuc_user(session):
        raise OcopImportError("Chỉ tài khoản Chi cục được import dữ liệu OCOP.")
    if mode not in ("stage_only", "publish"):
        raise OcopImportError("Chế độ import OCOP không hợp lệ.")
    if mode == "publish":
        _assert_publishable(preview)

    existing = con.execute(
        "SELECT * FROM ocop_import_batches WHERE file_sha256=? AND sheet_name=? FOR UPDATE",
        (preview["file_sha256"], preview["sheet_name"]),
    ).fetchone()
    if existing and existing["status"] == "published":
        return {"batch_id": existing["id"], "already_published": True, "published_products": 0}

    if existing:
        batch_id = existing["id"]
        con.execute("DELETE FROM ocop_import_rows WHERE batch_id=?", (batch_id,))
        con.execute(
            """UPDATE ocop_import_batches SET filename=?,total_rows=?,valid_rows=?,warning_rows=?,error_rows=?,
               commit_mode=?,status=?,imported_by=?,imported_at=CURRENT_TIMESTAMP,
               published_at=CASE WHEN ?='publish' THEN CURRENT_TIMESTAMP ELSE NULL END WHERE id=?""",
            (preview["filename"], preview["total_rows"], preview["valid_rows"], preview["warning_rows"],
             preview["error_rows"], mode, "published" if mode == "publish" else "committed", session["user_id"], mode, batch_id),
        )
    else:
        cursor = con.execute(
            """INSERT INTO ocop_import_batches
               (filename,file_sha256,sheet_name,total_rows,valid_rows,warning_rows,error_rows,commit_mode,status,imported_by,published_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,CASE WHEN ?='publish' THEN CURRENT_TIMESTAMP ELSE NULL END)
               RETURNING id""",
            (preview["filename"], preview["file_sha256"], preview["sheet_name"], preview["total_rows"],
             preview["valid_rows"], preview["warning_rows"], preview["error_rows"], mode,
             "published" if mode == "publish" else "committed", session["user_id"], mode),
        )
        batch_id = cursor.fetchone()[0]

    for row in preview["rows"]:
        messages = [{"level": "error", "message": item} for item in row["errors"]]
        messages += [{"level": "warning", "message": item} for item in row["warnings"]]
        con.execute(
            """INSERT INTO ocop_import_rows
               (batch_id,excel_row,source_tt,source_unit_name,mapped_unit_code,validation_status,
                validation_messages_json,raw_data_json,canonical_data_json)
               VALUES(?,?,?,?,?,?,CAST(? AS JSONB),CAST(? AS JSONB),CAST(? AS JSONB))""",
            (batch_id, row["excel_row"], row["source_tt"], row["source_unit_name"],
             row["canonical"]["entity"]["unit_code"], row["status"], _json(messages),
             _json(row["raw_data"]), _json(row["canonical"])),
        )

    if mode == "stage_only":
        ocop_registry._audit(con, session, "Import OCOP staging",
            f"Lưu staging OCOP: {preview['total_rows']} dòng, {preview['error_rows']} dòng lỗi",
            "ocop_import_batch", batch_id)
        return {"batch_id": batch_id, "already_published": False, "published_products": 0}

    published_products = 0
    for row in preview["rows"]:
        canonical = row["canonical"]
        ocop_registry.publish(con, session, canonical, source="excel",
                              batch_id=batch_id, source_row=row["excel_row"])
        published_products += 1

    ocop_registry._audit(con, session, "Import OCOP Excel",
        f"Import OCOP: {published_products} sản phẩm, {preview['recognition_count']} lần công nhận",
        "ocop_import_batch", batch_id)
    return {"batch_id": batch_id, "already_published": False, "published_products": published_products}


def list_batches(con, limit=30):
    return [dict(row) for row in con.execute(
        """SELECT b.*,u.username FROM ocop_import_batches b
           JOIN users u ON u.id=b.imported_by ORDER BY b.id DESC LIMIT ?""", (int(limit),)
    ).fetchall()]


def list_recognitions(con, session, filters=None):
    """List recognition history with the same locality scope as OCOP products."""
    import ocop_services
    filters = filters or {}
    scope = ocop_services.get_scope(con, session)
    where = ["1=1"]
    args = []
    if scope:
        where.append("p.ma_don_vi_hanh_chinh=?")
        args.append(scope)
    unit = _text(filters.get("unit"))
    if unit and not scope:
        where.append("p.ma_don_vi_hanh_chinh=?")
        args.append(unit)
    q = _text(filters.get("q"))
    if q:
        where.append("(p.ten_san_pham ILIKE ? OR c.TenCoSo ILIKE ? OR r.decision_number ILIKE ?)")
        pattern = f"%{q}%"
        args.extend([pattern, pattern, pattern])
    year = _text(filters.get("year"))
    if year:
        try:
            year_value = int(year)
        except ValueError as exc:
            raise OcopImportError("Năm tra cứu không hợp lệ.") from exc
        where.append("r.recognition_year=?")
        args.append(year_value)
    rows = con.execute(
        """SELECT r.*,p.ten_san_pham,p.ma_san_pham,p.current_star,p.ma_don_vi_hanh_chinh,
                  c.TenCoSo AS entity_name,d.TenDonVi AS unit_name
           FROM ocop_recognitions r
           JOIN ocop_products p ON p.id=r.product_id
           JOIN DM_CoSo c ON c.Ma_CoSo=p.ma_co_so
           JOIN DM_DonViHanhChinh d ON d.Ma_DonViHanhChinh=p.ma_don_vi_hanh_chinh
           WHERE """ + " AND ".join(where) +
        " ORDER BY r.recognition_year DESC,p.ten_san_pham,r.recognition_sequence DESC,r.id DESC",
        args,
    ).fetchall()
    return [dict(row) for row in rows]


def recognitions_for_product(con, session, product_id):
    import ocop_services
    product = ocop_services.get_product(con, session, product_id)
    rows = con.execute(
        """SELECT * FROM ocop_recognitions WHERE product_id=?
           ORDER BY recognition_sequence DESC,recognition_year DESC,id DESC""", (product_id,)
    ).fetchall()
    return product, [dict(row) for row in rows]
