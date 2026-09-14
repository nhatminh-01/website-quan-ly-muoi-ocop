import unittest
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from weekly_import import detect_weekly_period, parse_weekly_workbook
from workbook_utils import inspect_workbook


def workbook_bytes(*sheet_names, header_value=None, row_values=None):
    book = Workbook()
    book.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        book.create_sheet(name)
    sheet = book[sheet_names[0]]
    if header_value is not None:
        sheet["A1"] = header_value
    if row_values is not None:
        for column, value in enumerate(row_values, 1):
            sheet.cell(2, column).value = value
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class WeeklyImportUiRefinementTests(unittest.TestCase):
    def test_inspect_workbook_returns_sheet_names_and_recommended_sheet(self):
        content = workbook_bytes("Sheet A", "Loc", "DanhMuc")

        inspected = inspect_workbook(content, "CCPTNT OCOP.xlsx", "Loc")

        self.assertEqual(inspected["sheet_names"], ["Sheet A", "Loc", "DanhMuc"])
        self.assertEqual(inspected["recommended_sheet"], "Loc")

    def test_period_detection_uses_date_in_worksheet(self):
        content = workbook_bytes(
            "21.8-Tuan 34", header_value=datetime(2026, 8, 21)
        )

        detected = detect_weekly_period(content, "BaoCao.xlsx", "21.8-Tuan 34")

        self.assertEqual(detected["report_date"], "2026-08-21")
        self.assertEqual(detected["source"], "worksheet")
        self.assertEqual(detected["week_number"], 34)

    def test_day_month_sheet_does_not_guess_year(self):
        content = workbook_bytes("21.8-Tuan 34")

        detected = detect_weekly_period(content, "BaoCao.xlsx", "21.8-Tuan 34")

        self.assertIsNone(detected["report_date"])
        self.assertEqual(detected["week_number"], 34)
        self.assertFalse(detected["period_detected"])

    def test_year_from_filename_can_complete_day_month_sheet(self):
        content = workbook_bytes("21.8-Tuan 34")

        detected = detect_weekly_period(content, "BaoCao_2026.xlsx", "21.8-Tuan 34")

        self.assertEqual(detected["report_date"], "2026-08-21")
        self.assertEqual(detected["source"], "sheet_name+filename")

    def test_sheet_week_mismatch_blocks_preview(self):
        values = [None, "Xã A"] + [0] * 26
        content = workbook_bytes("21.8-Tuan 34", row_values=values)

        preview = parse_weekly_workbook(
            content,
            "2026-08-28",
            "21.8-Tuan 34",
            "BaoCao.xlsx",
            lambda name: ("Xã A", "001") if name == "Xã A" else None,
        )

        self.assertTrue(preview["period_blocked"])
        self.assertIn("Tuần 34", preview["period_warnings"][0])
        self.assertEqual(preview["week_code"], "2026-W35")


if __name__ == "__main__":
    unittest.main()
