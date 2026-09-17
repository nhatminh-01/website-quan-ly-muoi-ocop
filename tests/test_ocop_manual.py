"""Direct OCOP entry shares the Excel registry, projection and audit transaction."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import date, timedelta
import io
import os
import unittest
from unittest.mock import patch

import admin_units
import backend_db
import dashboard_services
import ocop_import
import ocop_manual
import ocop_registry as registry
import ocop_services as svc
import server
import test_ocop_legacy_import as excel_tests
import test_profile_http as http_tests
import user_profiles


def form(**changes):
    values = {
        "entity_name":"Hợp tác xã Nhập trực tiếp", "business_type":"Hợp tác xã", "unit_code":"27595",
        "address":"12 Đường kiểm thử", "representative_name":"Nguyễn Văn A", "phone":"0909000000",
        "email":"ocop@example.gov.vn", "product_name":"Sản phẩm nhập trực tiếp", "product_group":"Gia vị",
        "description":"Mô tả sản phẩm", "recognition_0_star_rank":"3", "recognition_0_evaluation_type":"new",
        "recognition_0_recognition_year":"2026", "recognition_0_recognition_date":"2026-01-02",
        "recognition_0_expiry_date":(date.today() + timedelta(days=30)).isoformat(),
        "recognition_0_decision_number":"100/QĐ", "recognition_0_decision_authority":"UBND Thành phố",
        "recognition_0_note":"Lần đầu",
    }
    values.update(changes)
    return values


def next_recognition(**changes):
    values = form(recognition_0_star_rank="4", recognition_0_evaluation_type="upgrade",
                  recognition_0_recognition_year="2027", recognition_0_recognition_date="2027-01-03",
                  recognition_0_expiry_date="2030-01-03", recognition_0_decision_number="200/QĐ")
    values.update(changes)
    return values


class ManualValidationTests(unittest.TestCase):
    def test_strict_server_validation_rejects_forged_options_and_dates(self):
        cases = (
            {"entity_name":""}, {"product_name":""}, {"business_type":"Tự nhập"},
            {"email":"invalid"}, {"recognition_0_star_rank":"2"}, {"recognition_0_star_rank":"6"},
            {"recognition_0_star_rank":"3.9"}, {"recognition_0_recognition_year":"1899"},
            {"recognition_0_recognition_year":"10000"}, {"recognition_0_recognition_year":"2026.5"},
            {"recognition_0_evaluation_type":"historical"}, {"recognition_0_recognition_date":"2025-01-01"},
            {"recognition_0_recognition_date":"2026-02-30"}, {"recognition_0_expiry_date":"not-a-date"},
        )
        for changes in cases:
            with self.subTest(changes=changes), self.assertRaises(registry.RegistryError):
                registry.validate_payload(ocop_manual.payload_from_form(form(**changes)), manual=True)

    def test_code_and_current_star_cannot_be_supplied_by_the_client(self):
        payload = ocop_manual.payload_from_form(form(ma_co_so="FORGED", ma_san_pham="FORGED", current_star="5", source_batch_id="42"))
        self.assertNotEqual(payload["entity"]["ma_co_so"], "FORGED")
        self.assertNotEqual(payload["product"]["ma_san_pham"], "FORGED")
        self.assertNotIn("current_star", payload["product"])
        self.assertNotIn("source_batch_id", payload)

    def test_optional_contacts_dates_and_year_boundaries(self):
        for year in (1900, 9999):
            payload = registry.validate_payload(ocop_manual.payload_from_form(form(
                email="", address="", phone="", recognition_0_recognition_date="",
                recognition_0_expiry_date="", recognition_0_recognition_year=str(year))), manual=True)
            self.assertEqual(payload["recognitions"][0]["recognition_year"], year)
            self.assertIsNone(payload["recognitions"][0]["expiry_date"])

    def test_repeated_recognition_and_too_many_rows_are_rejected(self):
        payload = ocop_manual.payload_from_form(form())
        payload["recognitions"] *= 2
        with self.assertRaises(registry.RegistryError):
            registry.validate_payload(payload, manual=True)
        data = {f"recognition_{i}_star_rank":"3" for i in range(51)}
        with self.assertRaises(registry.RegistryError):
            ocop_manual.form_recognitions(data)


@unittest.skipUnless(os.getenv("OCOP_TEST_PG_DSN"), "Requires disposable PostgreSQL test databases")
class ManualPostgreSQLTests(unittest.TestCase):
    connection = http_tests.ProfileHTTPTests.connection
    drop_database = http_tests.ProfileHTTPTests.drop_database

    def setUp(self):
        http_tests.ProfileHTTPTests.setUp(self)
        self.con = backend_db.CompatConnection(self.connection())
        self.addCleanup(self.con.close)
        self.session = {"user_id":self.ids["profile_staff"], "role":"staff", "unit_name":"Chi cục"}
        self.con.execute("""INSERT INTO app.user_profiles(user_id,full_name,department)
            VALUES(?,?,?) ON CONFLICT(user_id) DO UPDATE SET full_name=excluded.full_name,department=excluded.department""",
            (self.session["user_id"], "Chuyên viên OCOP", user_profiles.STAFF_DEPARTMENTS[1]))
        self.con.commit()

    def count(self, table):
        return self.con.execute("SELECT COUNT(*) FROM " + table).fetchone()[0]

    def test_admin_staff_routes_and_legacy_backend_denial(self):
        for client in (self.admin, self.staff):
            status, _, html = client.request("GET", "/ocop/manual")
            self.assertEqual(status, 200)
            page = http_tests.Page(html)
            self.assertIn("NHẬP DỮ LIỆU OCOP", html)
            self.assertEqual(page.input("phone")["type"], "tel")
            self.assertFalse(any(node["attrs"].get("name") in ("ma_co_so", "ma_san_pham", "current_star") for node in page.nodes))
        legacy = http_tests.Client(self.http.server_address[1])
        row = server.get_user(self.con, self.ids["legacy_unit"])
        sid = server.new_session(row)
        legacy.cookie = "salt_session=" + sid
        legacy.csrf = server.SESSIONS[sid]["csrf"]
        self.assertEqual(legacy.request("GET", "/ocop/manual")[0], 403)
        self.assertEqual(legacy.request("POST", "/ocop/manual", form())[0], 403)
        with self.assertRaises(svc.OcopError) as blocked:
            ocop_manual.save(self.con, {"user_id":row["id"], "role":"unit"}, form())
        self.assertEqual(blocked.exception.status, 403)
        self.assertEqual(self.count("app.ocop_products"), 0)

    def test_create_is_visible_in_catalog_details_dashboard_history_and_export(self):
        status, headers, _ = self.staff.request("POST", "/ocop/manual", form())
        self.assertEqual(status, 303)
        self.assertIn("Đã lưu dữ liệu OCOP thành công.", self.staff.request("GET", headers["Location"])[2])
        product = self.con.execute("SELECT * FROM app.ocop_products").fetchone()
        self.assertEqual(product["current_star"], 3)
        self.assertEqual(self.count("app.ocop_entities"), 1)
        self.assertEqual(self.count("app.ocop_applications"), 0)
        self.assertEqual(self.count("staging.ocop_import_batches"), 0)
        self.assertEqual(self.count("staging.ocop_import_rows"), 0)
        recognition = self.con.execute("SELECT * FROM app.ocop_recognitions").fetchone()
        for field in ("application_id", "source_batch_id", "source_row"):
            self.assertIsNone(recognition[field])
        self.assertTrue(recognition["is_current"])
        qd = self.con.execute("SELECT * FROM qd5277.PTNT_OCOP").fetchone()
        self.assertEqual((qd["XepHang"], qd["TrangThai"]), ("3*", "HIEULUC"))
        self.assertEqual(svc.list_products(self.con, self.session)["total"], 1)
        self.assertEqual(svc.list_entities(self.con, self.session)["total"], 1)
        dashboard = dashboard_services.get_ocop_dashboard(self.con, self.session)
        self.assertEqual(dashboard["star_3"], 1)
        self.assertEqual(dashboard["total_products"], 1)
        self.assertEqual(dashboard["expiring_products"], 1)
        self.assertIn(form()["product_name"], self.staff.request("GET", "/ocop")[2])
        detail = self.staff.request("GET", f'/ocop/products/{product["id"]}')[2]
        self.assertIn("100/QĐ", detail)
        self.assertIn("LỊCH SỬ CÔNG NHẬN", detail)
        self.assertIn("100/QĐ", self.staff.request("GET", "/ocop/recognitions")[2])
        content, _ = server.export_ocop_xlsx(self.session)
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content))
        try:
            self.assertTrue(any(form()["product_name"] in row for row in book.active.iter_rows(values_only=True)))
        finally:
            book.close()

    def test_duplicate_product_requires_explicit_append_and_replay_is_safe(self):
        first = ocop_manual.save(self.con, self.session, form())
        status, _, html = self.staff.request("POST", "/ocop/manual", next_recognition())
        self.assertEqual(status, 409)
        self.assertIn("Sản phẩm này đã tồn tại.", html)
        self.assertIn("Thêm lần công nhận", html)
        self.assertEqual(self.count("app.ocop_recognitions"), 1)
        result = ocop_manual.save(self.con, self.session, next_recognition(append_product_id=str(first["product_id"])))
        self.assertEqual(result["product_id"], first["product_id"])
        rows = self.con.execute("SELECT recognition_sequence,is_current FROM app.ocop_recognitions ORDER BY recognition_sequence").fetchall()
        self.assertEqual([(r[0], r[1]) for r in rows], [(1, False), (2, True)])
        self.assertEqual(self.con.execute("SELECT current_star FROM app.ocop_products").fetchone()[0], 4)
        self.assertEqual(self.count("qd5277.PTNT_OCOP"), 1)
        with self.assertRaises(registry.RegistryError):
            ocop_manual.save(self.con, self.session, next_recognition(append_product_id=str(first["product_id"])))
        self.assertEqual(self.count("app.ocop_recognitions"), 2)

    def test_reuses_entity_with_normalized_name_without_overwriting_contacts(self):
        first = ocop_manual.save(self.con, self.session, form())
        second = ocop_manual.save(self.con, self.session, form(entity_name="  HỢP TÁC XÃ  Nhập trực tiếp ", product_name="Sản phẩm khác", phone=""))
        self.assertEqual(first["entity_id"], second["entity_id"])
        self.assertEqual(self.count("app.ocop_entities"), 1)
        self.assertEqual(self.count("app.ocop_products"), 2)
        self.assertEqual(self.con.execute("SELECT phone FROM app.ocop_entities").fetchone()[0], form()["phone"])

    def test_product_identity_includes_subject_and_locality(self):
        ocop_manual.save(self.con, self.session, form())
        ocop_manual.save(self.con, self.session, form(entity_name="Chủ thể khác"))
        ocop_manual.save(self.con, self.session, form(unit_code="27673"))
        self.assertEqual(self.count("app.ocop_products"), 3)
        self.assertEqual(self.count("app.ocop_entities"), 3)

    def test_inactive_missing_unit_and_fake_group_are_rejected(self):
        self.con.execute("UPDATE DM_DonViHanhChinh SET TinhTrang=FALSE WHERE Ma_DonViHanhChinh='27595'")
        self.con.commit()
        for values in (form(), form(unit_code="TMP999"), form(unit_code="27673", product_group="Nhóm giả")):
            with self.assertRaises(registry.RegistryError):
                ocop_manual.save(self.con, self.session, values)
        self.assertEqual(self.count("app.ocop_products"), 0)
        self.assertEqual(self.count("app.ocop_entities"), 0)

    def test_ambiguous_entity_requires_choice_and_forged_append_is_rejected(self):
        first = ocop_manual.save(self.con, self.session, form())
        self.con.execute("INSERT INTO DM_CoSo(Ma_CoSo,TenCoSo,LoaiCoSo,DiaChi,Ma_DonViHanhChinh) VALUES('CSAMBIG',?,'Hộ','Địa chỉ khác','27595')", (form()["entity_name"],))
        self.con.execute("""INSERT INTO app.ocop_entities(ma_co_so,created_by,created_at,updated_at)
            VALUES('CSAMBIG',?,?,?)""", (self.session["user_id"], server.now_text(), server.now_text()))
        self.con.commit()
        with self.assertRaises(registry.AmbiguousEntity):
            ocop_manual.save(self.con, self.session, form(product_name="Sản phẩm thứ hai"))
        result = ocop_manual.save(self.con, self.session, form(product_name="Sản phẩm thứ hai", selected_entity_id=str(first["entity_id"])))
        self.assertEqual(result["entity_id"], first["entity_id"])
        with self.assertRaises(registry.RegistryError):
            ocop_manual.save(self.con, self.session, next_recognition(
                selected_entity_id=str(first["entity_id"]), append_product_id=str(result["product_id"])))
        self.assertEqual(self.count("app.ocop_recognitions"), 2)

    def test_stale_or_forged_role_is_checked_against_database(self):
        with self.assertRaises(svc.OcopError):
            ocop_manual.save(self.con, {**self.session, "role":"admin"}, form())
        self.con.execute("UPDATE app.users SET active=FALSE WHERE id=?", (self.session["user_id"],))
        self.con.commit()
        with self.assertRaises(svc.OcopError):
            ocop_manual.save(self.con, self.session, form())
        self.assertEqual(self.count("app.ocop_entities"), 0)

    def test_append_keeps_legacy_subject_type_without_blocking_recognition(self):
        first = ocop_manual.save(self.con, self.session, form())
        self.con.execute("UPDATE DM_CoSo SET LoaiCoSo='Công ty cổ phần' WHERE Ma_CoSo=(SELECT ma_co_so FROM app.ocop_products WHERE id=?)", (first["product_id"],))
        self.con.commit()
        page = self.staff.request("GET", f'/ocop/manual?product_id={first["product_id"]}')
        self.assertEqual(page[0], 200)
        self.assertIn("Công ty cổ phần", page[2])
        ocop_manual.save(self.con, self.session, next_recognition(append_product_id=str(first["product_id"]), business_type="Công ty cổ phần"))
        self.assertEqual(self.count("app.ocop_recognitions"), 2)
        self.assertEqual(self.con.execute("SELECT LoaiCoSo FROM DM_CoSo WHERE Ma_CoSo=(SELECT ma_co_so FROM app.ocop_products WHERE id=?)", (first["product_id"],)).fetchone()[0], "Công ty cổ phần")

    def test_multiple_recognitions_choose_latest_and_keep_older_history(self):
        data = next_recognition()
        for field in ocop_manual.RECOGNITION_FIELDS:
            data["recognition_1_" + field] = form().get("recognition_0_" + field, "")
        result = ocop_manual.save(self.con, self.session, data)
        old = form(append_product_id=str(result["product_id"]), recognition_0_recognition_year="2020", recognition_0_recognition_date="2020-01-01")
        ocop_manual.save(self.con, self.session, old)
        rows = self.con.execute("SELECT recognition_sequence,recognition_year,is_current FROM app.ocop_recognitions ORDER BY recognition_sequence").fetchall()
        self.assertEqual([(r[0], r[1], r[2]) for r in rows], [(1, 2026, False), (2, 2027, True), (3, 2020, False)])
        self.assertEqual(self.con.execute("SELECT current_star FROM app.ocop_products").fetchone()[0], 4)

    def test_expired_and_expiring_status_sync_and_optional_date(self):
        ocop_manual.save(self.con, self.session, form(recognition_0_expiry_date=(date.today() - timedelta(days=1)).isoformat()))
        self.assertEqual(self.con.execute("SELECT TrangThai FROM PTNT_OCOP").fetchone()[0], "HETHAN")
        self.assertEqual(svc.get_ocop_expiry_summary(self.con, self.session)["expired_products"], 1)
        ocop_manual.save(self.con, self.session, form(product_name="Không có ngày", recognition_0_recognition_date="", recognition_0_expiry_date=""))
        self.assertEqual(svc.get_ocop_expiry_summary(self.con, self.session)["missing_expiry"], 1)

    def test_excel_then_manual_then_excel_preserves_history_and_current(self):
        preview = ocop_import.parse_ocop_workbook(excel_tests.workbook_bytes(second_product=False), "fixture.xlsx", admin_units.unit_lookup(self.con))
        ocop_import.commit_ocop_preview(self.con, self.session, preview)
        self.con.commit()
        old_ids = [row[0] for row in self.con.execute("SELECT id FROM app.ocop_recognitions").fetchall()]
        pid = self.con.execute("SELECT id FROM app.ocop_products").fetchone()[0]
        values = next_recognition(entity_name="Hợp tác xã T2", product_name="Sản phẩm A", product_group="Thực phẩm",
                                  recognition_0_star_rank="5", append_product_id=str(pid))
        ocop_manual.save(self.con, self.session, values)
        preview["file_sha256"] = "a" * 64
        ocop_import.commit_ocop_preview(self.con, self.session, preview)
        self.con.commit()
        self.assertEqual(self.count("app.ocop_entities"), 1)
        self.assertEqual(self.count("app.ocop_products"), 1)
        self.assertEqual(self.count("app.ocop_recognitions"), 3)
        self.assertEqual(self.con.execute("SELECT current_star FROM app.ocop_products").fetchone()[0], 5)
        ids = [row[0] for row in self.con.execute("SELECT id FROM app.ocop_recognitions").fetchall()]
        self.assertTrue(set(old_ids).issubset(ids))
        self.assertEqual(self.con.execute("SELECT COUNT(*) FROM app.ocop_recognitions WHERE source_batch_id IS NULL").fetchone()[0], 1)

    def test_manual_then_excel_reuses_subject_product_and_exact_recognition(self):
        values = form(entity_name="Hợp tác xã T2", product_name="Sản phẩm A", recognition_0_recognition_year="2024",
                      recognition_0_recognition_date="2024-01-02", recognition_0_expiry_date="",
                      recognition_0_decision_authority="UBND TPHCM")
        ocop_manual.save(self.con, self.session, values)
        preview = ocop_import.parse_ocop_workbook(excel_tests.workbook_bytes(second_product=False), "fixture.xlsx", admin_units.unit_lookup(self.con))
        ocop_import.commit_ocop_preview(self.con, self.session, preview)
        self.con.commit()
        self.assertEqual(self.count("app.ocop_products"), 1)
        self.assertEqual(self.count("app.ocop_recognitions"), 2)
        self.assertEqual(self.con.execute("SELECT current_star FROM app.ocop_products").fetchone()[0], 4)

    def test_business_or_audit_failure_rolls_back_all_rows(self):
        for target in ("_publish_qd5277", "_audit"):
            with self.subTest(target=target), patch.object(registry, target, side_effect=RuntimeError("test-only failure")):
                with self.assertRaises(RuntimeError):
                    ocop_manual.save(self.con, self.session, form())
            for table in ("app.ocop_entities", "app.ocop_products", "app.ocop_recognitions", "qd5277.PTNT_OCOP"):
                self.assertEqual(self.count(table), 0)
            self.assertEqual(self.con.execute("SELECT COUNT(*) FROM app.audit_logs WHERE action LIKE 'ocop_manual_%' AND success=TRUE").fetchone()[0], 0)

    def test_audit_snapshot_and_sanitized_detail_are_atomic(self):
        result = ocop_manual.save(self.con, self.session, form(password="SENTINEL_PASSWORD", csrf="SENTINEL_CSRF"))
        logs = self.con.execute("SELECT * FROM app.audit_logs WHERE action LIKE 'ocop_manual_%' ORDER BY id").fetchall()
        self.assertEqual(len(logs), 3)
        for row in logs:
            self.assertEqual(row["user_id"], self.session["user_id"])
            self.assertEqual(row["role_snapshot"], "staff")
            self.assertEqual(row["department_snapshot"], user_profiles.STAFF_DEPARTMENTS[1])
            self.assertEqual(row["module"], "ocop")
            self.assertTrue(row["success"])
            self.assertNotIn("SENTINEL", row["detail"])
        self.con.execute("UPDATE app.user_profiles SET department=? WHERE user_id=?", (user_profiles.STAFF_DEPARTMENTS[0], self.session["user_id"]))
        self.con.commit()
        self.assertEqual(self.con.execute("SELECT department_snapshot FROM app.audit_logs WHERE action=?", (registry.MANUAL_PRODUCT_CREATE,)).fetchone()[0], user_profiles.STAFF_DEPARTMENTS[1])
        self.assertIsNotNone(result["product_id"])

    def test_concurrent_duplicate_requests_create_one_product(self):
        def submit():
            con = backend_db.CompatConnection(self.connection())
            try:
                return ocop_manual.save(con, self.session, form())["product_id"]
            except registry.DuplicateProduct:
                return "duplicate"
            finally:
                con.close()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: submit(), range(2)))
        self.assertEqual(results.count("duplicate"), 1)
        self.assertEqual(self.count("app.ocop_products"), 1)
        self.assertEqual(self.count("app.ocop_recognitions"), 1)

    def test_csrf_and_invalid_form_do_not_write_and_escape_error_values(self):
        self.assertEqual(self.staff.request("POST", "/ocop/manual", form(csrf="wrong"))[0], 400)
        status, _, html = self.staff.request("POST", "/ocop/manual", form(entity_name='<script>alert(1)</script>', recognition_0_star_rank="6"))
        self.assertEqual(status, 400)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>alert(1)</script>", html)
        self.assertEqual(self.count("app.ocop_products"), 0)
