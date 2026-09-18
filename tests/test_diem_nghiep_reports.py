import unittest
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from diem_nghiep_reports import (
    CUC_LOCALITIES,
    REPORT_CUC,
    REPORT_SO,
    build_report_data,
    export_diem_nghiep_report,
    select_latest_month_snapshot,
)


def weekly_row(report_date, unit_name, *, area=100, harvest=1000, sold=400, remaining=600,
               households=10, workers=20, process_fine=30, process_iodized=5, damage=0):
    return {
        "report_date": report_date,
        "unit_name": unit_name,
        "area_land": area,
        "area_tarp": 0,
        "harvest_land": harvest,
        "harvest_tarp": 0,
        "sold_land": sold,
        "sold_tarp": 0,
        "remaining_land": remaining,
        "remaining_tarp": 0,
        "processed_fine": process_fine,
        "processed_iodized": process_iodized,
        "households": households,
        "workers": workers,
        "damage_land": damage,
        "damage_tarp": 0,
    }


def snapshot(day, rows):
    return {"date": date.fromisoformat(day), "rows": rows}


class DiemNghiepReportTests(unittest.TestCase):
    def test_latest_snapshot_stays_inside_selected_month(self):
        rows = [
            weekly_row("2026-08-07", CUC_LOCALITIES[0]),
            weekly_row("2026-08-21", CUC_LOCALITIES[0]),
            weekly_row("2026-09-04", CUC_LOCALITIES[0]),
        ]
        selected = select_latest_month_snapshot(rows, 2026, 8)
        self.assertEqual(selected["date"], date(2026, 8, 21))
        self.assertEqual(selected["rows"][0]["report_date"], "2026-08-21")

    def test_missing_month_does_not_fallback_and_disables_export(self):
        report = build_report_data(REPORT_CUC, 2026, 8, {"current": None})
        self.assertIsNone(report["snapshot_date"])
        self.assertEqual(report["rows"], [])
        self.assertFalse(report["can_export"])

    def test_cuc_has_eight_areas_and_total_with_zero_derived_industrial(self):
        current_rows = [weekly_row("2026-08-21", name, area=index) for index, name in enumerate(CUC_LOCALITIES, 1)]
        report = build_report_data(REPORT_CUC, 2026, 8, {"current": snapshot("2026-08-21", current_rows)})
        self.assertTrue(report["can_export"])
        self.assertEqual(len(report["rows"]), 9)
        self.assertEqual(report["rows"][0]["name"], "CẢ TỈNH")
        self.assertEqual(report["rows"][0]["cells"]["C"], Decimal("36"))
        self.assertEqual(report["rows"][1]["cells"]["E"], Decimal("0"))

    def test_so_without_previous_year_keeps_d_e_i_j_blank(self):
        snapshots = {
            "current": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], harvest=100)]),
            "previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], harvest=50)]),
            "previous_previous_month": snapshot("2026-06-26", [weekly_row("2026-06-26", CUC_LOCALITIES[0], harvest=40)]),
            "previous_year_current": None,
            "previous_year_previous": None,
        }
        report = build_report_data(REPORT_SO, 2026, 8, snapshots)
        row = next(item for item in report["rows"] if item["row_number"] == 9)
        self.assertEqual(next(item for item in report["rows"] if item["row_number"] == 6)["stt"], "1")
        self.assertEqual(row["stt"], "2")
        self.assertIsNone(row["cells"]["D"])
        self.assertIsNone(row["cells"]["E"])
        self.assertEqual(row["cells"]["F"], Decimal("10"))
        self.assertEqual(row["cells"]["G"], Decimal("50"))
        self.assertEqual(row["cells"]["H"], Decimal("100"))
        self.assertIsNone(row["cells"]["I"])
        self.assertIsNone(row["cells"]["J"])

    def test_negative_delta_is_null_and_warns(self):
        snapshots = {
            "current": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], harvest=90)]),
            "previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], harvest=100)]),
            "previous_previous_month": None,
            "previous_year_current": None,
            "previous_year_previous": None,
        }
        report = build_report_data(REPORT_SO, 2026, 8, snapshots)
        row = next(item for item in report["rows"] if item["row_number"] == 9)
        self.assertIsNone(row["cells"]["G"])
        self.assertTrue(report["warnings"])

    def test_negative_consumption_delta_explains_both_cumulative_values(self):
        snapshots = {
            "current": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], sold=90)]),
            "previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], sold=100)]),
            "previous_previous_month": None,
            "previous_year_current": None,
            "previous_year_previous": None,
        }
        report = build_report_data(REPORT_SO, 2026, 8, snapshots)
        detail = report["warning_details"][0]
        self.assertEqual(detail["row_number"], 12)
        self.assertEqual(detail["column"], "G")
        self.assertEqual(detail["formula"], "90 - 100 = -10 tấn")
        self.assertIn("ô tháng báo cáo", detail["logic_explanation"])
        self.assertIn("90", detail["message"])
        self.assertIn("100", detail["message"])
        row = next(item for item in report["rows"] if item["row_number"] == 12)
        self.assertIn("G", row["cell_warnings"])

    def test_current_month_stays_independent_when_previous_month_delta_is_negative(self):
        report = build_report_data(
            REPORT_SO,
            2026,
            9,
            {
                "current": snapshot("2026-09-04", [weekly_row("2026-09-04", CUC_LOCALITIES[0], sold=130)]),
                "previous_month": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], sold=120)]),
                "previous_previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], sold=125)]),
                "previous_year_current": None,
                "previous_year_previous": None,
            },
        )
        row = next(item for item in report["rows"] if item["row_number"] == 12)
        self.assertIsNone(row["cells"]["F"])
        self.assertEqual(row["cells"]["G"], Decimal("10"))
        self.assertEqual(row["cells"]["H"], Decimal("130"))
        self.assertEqual(report["warning_details"][0]["column"], "F")
        self.assertIn("lấy độc lập từ snapshot", report["warning_details"][0]["logic_explanation"])

    def test_export_preserves_template_mapping_and_preview_values(self):
        current_rows = [weekly_row("2026-08-21", name, area=index) for index, name in enumerate(CUC_LOCALITIES, 1)]
        report = build_report_data(REPORT_CUC, 2026, 8, {"current": snapshot("2026-08-21", current_rows)})
        content = export_diem_nghiep_report(report)
        exported = load_workbook(BytesIO(content), data_only=False)
        original_path = Path(__file__).resolve().parents[1] / "assets" / "templates" / "diem-nghiep" / "template_bao_cao_cuc_diem_nghiep.xlsx"
        original = load_workbook(original_path, data_only=False)
        self.assertEqual(exported.sheetnames, original.sheetnames)
        self.assertEqual(list(exported["Bao cao Cuc"].merged_cells.ranges), list(original["Bao cao Cuc"].merged_cells.ranges))
        self.assertEqual(exported["Bao cao Cuc"]["C6"].value, "=SUM(C7:C14)")
        self.assertEqual(exported["Bao cao Cuc"]["C7"].value, 1)
        self.assertIn("21/08/2026", exported["Bao cao Cuc"]["A2"].value)
        self.assertEqual(exported["_MAPPING"]["B1"].value, "CUC_BIEU_8A_DIEM_NGHIEP")

    def test_so_export_keeps_sections_and_blank_source_cells(self):
        report = build_report_data(
            REPORT_SO,
            2026,
            8,
            {"current": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0])])},
        )
        content = export_diem_nghiep_report(report)
        book = load_workbook(BytesIO(content), data_only=False)
        ws = book["Bao cao So"]
        self.assertEqual(ws["A35"].value, "III")
        self.assertIsNone(ws["D6"].value)
        self.assertIsNotNone(ws["I6"].value)
        self.assertIn("08/2026", ws["A2"].value)

    def test_so_export_marks_warning_cell_and_adds_excel_explanation(self):
        report = build_report_data(
            REPORT_SO,
            2026,
            8,
            {
                "current": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], sold=90)]),
                "previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], sold=100)]),
                "previous_previous_month": None,
                "previous_year_current": None,
                "previous_year_previous": None,
            },
        )
        content = export_diem_nghiep_report(report)
        book = load_workbook(BytesIO(content), data_only=False)
        cell = book["Bao cao So"]["G12"]
        self.assertTrue(cell.fill.fgColor.rgb.endswith("FFC7CE"))
        self.assertIsNotNone(cell.comment)
        self.assertIn("90", cell.comment.text)
        self.assertIn("100", cell.comment.text)
        self.assertIn("Cách tính: 90 - 100 = -10 tấn", cell.comment.text)
        self.assertIn("Giải thích logic:", cell.comment.text)

    def test_september_export_marks_previous_month_reference_not_current_month(self):
        report = build_report_data(
            REPORT_SO,
            2026,
            9,
            {
                "current": snapshot("2026-09-04", [weekly_row("2026-09-04", CUC_LOCALITIES[0], sold=130)]),
                "previous_month": snapshot("2026-08-21", [weekly_row("2026-08-21", CUC_LOCALITIES[0], sold=120)]),
                "previous_previous_month": snapshot("2026-07-24", [weekly_row("2026-07-24", CUC_LOCALITIES[0], sold=125)]),
                "previous_year_current": None,
                "previous_year_previous": None,
            },
        )
        book = load_workbook(BytesIO(export_diem_nghiep_report(report)), data_only=False)
        ws = book["Bao cao So"]
        self.assertTrue(ws["F12"].fill.fgColor.rgb.endswith("FFC7CE"))
        self.assertFalse(ws["G12"].fill.fgColor.rgb.endswith("FFC7CE"))
        self.assertEqual(ws["G12"].value, 10)
        self.assertIn("Chính thức tháng trước", ws["F4"].comment.text)
        self.assertIn("không phải số lũy kế", ws["F4"].comment.text)


if __name__ == "__main__":
    unittest.main()
