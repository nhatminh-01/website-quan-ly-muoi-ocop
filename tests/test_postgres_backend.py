"""Opt-in PostgreSQL integration: each test creates and drops its own database.

OCOP_TEST_PG_DSN is a PostgreSQL maintenance connection with CREATEDB privilege.
Never uses the application's configured database or existing user databases.
"""
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from unittest.mock import patch
from uuid import uuid4

import server
import backend_db
import ocop_db
import ocop_services as svc
from weekly_import import WeeklyImportError, parse_weekly_workbook, commit_weekly_preview
from database.migration import migrate_sqlite_to_postgres as migration
from test_integration import legacy_schema, Client


class PostgreSQLTranslationTests(unittest.TestCase):
    def test_sqlite_null_safe_inequality_is_translated(self):
        sql,_=backend_db.translate_sql('UPDATE ocop_applications SET criteria_set_id=? WHERE criteria_set_id IS NOT ?')
        self.assertIn('criteria_set_id IS DISTINCT FROM %s',sql)


@unittest.skipUnless(os.getenv('OCOP_TEST_PG_DSN'), 'Set OCOP_TEST_PG_DSN to enable disposable PostgreSQL tests')
class PostgreSQLBackendTests(unittest.TestCase):
    def setUp(self):
        import psycopg
        from psycopg import sql
        self.driver=psycopg
        self.dsn=os.environ['OCOP_TEST_PG_DSN']
        self.name='ocop_test_'+uuid4().hex
        with psycopg.connect(self.dsn,autocommit=True) as admin:
            admin.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(self.name)))
        self.addCleanup(self.drop_database)
        with self.connection(autocommit=True) as con:
            for name in ('002_qd5277_schema.sql','004_app_schema.sql','005_monthly_salt_sync.sql',
                         '006_official_admin_units.sql','007_ocop_init_only.sql',
                         '008_ocop_dynamic_criteria.sql','009_staff_role.sql',
                         '010_weekly_salt_imports.sql'):
                con.execute((Path(server.__file__).parent/'database/sql'/name).read_text(encoding='utf-8'))
        self.temp=tempfile.TemporaryDirectory(prefix='ocop-pg-source-')
        self.addCleanup(self.temp.cleanup)
        self.source=Path(self.temp.name)/'source_TEST.db'
        con=sqlite3.connect(self.source); con.row_factory=sqlite3.Row
        con.execute('PRAGMA foreign_keys=ON');con.executescript(legacy_schema())
        password=server.hash_password('Testing-OCOP-2026')
        for uid,user,role,name in ((1,'admin_test','admin','Chi cục'),(2,'unit_a','unit','Xã An Thới Đông'),(3,'staff_test','staff','Chi cục')):
            con.execute('INSERT INTO users VALUES(?,?,?,?,?,1,?)',(uid,user,password,role,name,'2026-01-01'))
        con.execute("INSERT INTO DM_DonViHanhChinh VALUES('27673',NULL,'Xã An Thới Đông','xa',1)")
        con.commit();ocop_db.migrate(con)
        session={'user_id':2,'role':'unit'}
        entity=svc.create_entity(con,session,{'name':'Fixture HTX','facility_type':'Hợp tác xã','address':'Fixture','ma_don_vi_hanh_chinh':'27673'})
        product=svc.create_product(con,session,{'name':'Fixture salt','ma_co_so':entity['ma_co_so'],'criteria_set_id':10})
        application=svc.create_application(con,session,{'product_id':product['id'],'evaluation_type':'new','year':2026})
        svc.application_action(con,session,application['id'],'submit',{'revision':application['revision']})
        con.execute("""INSERT INTO records(id,report_date,unit_name,created_by,area_land,area_tarp,harvest_land,harvest_tarp,price_land,price_tarp,status,created_at,updated_at)
            VALUES(1,'2026-08-28','Xã An Thới Đông',2,2,3,10,20,'1.200','2.000 - 2.500','approved','date','date')""")
        con.commit();con.close()

    def connection(self,autocommit=False):
        return self.driver.connect(self.dsn,dbname=self.name,autocommit=autocommit,
            row_factory=backend_db.hybrid_row,options='-c search_path=app,qd5277,staging,public')

    def drop_database(self):
        from psycopg import sql
        with self.driver.connect(self.dsn,autocommit=True) as admin:
            admin.execute(sql.SQL('DROP DATABASE {} WITH (FORCE)').format(sql.Identifier(self.name)))

    def migrate(self,**kwargs):
        with self.connection() as con:
            return migration.migrate(self.source,target=con,**kwargs)

    def test_ocop_copy_twice_preserves_source_relations_snapshots_and_sequences(self):
        before=hashlib.sha256(self.source.read_bytes()).hexdigest()
        counts=self.migrate();self.assertEqual(counts,self.migrate())
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),before)
        with self.connection() as con:
            for table,n in (('ocop_criteria_sets',26),('ocop_criteria',78),('ocop_entities',1),('ocop_products',1),('ocop_applications',1),('PTNT_OCOP',0)):
                self.assertEqual(con.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0],n,table)
            row=con.execute('SELECT * FROM ocop_applications').fetchone()
            self.assertEqual((row['criteria_set_id'],row['revision']),(10,2))
            self.assertEqual(json.loads(row['submission_snapshot_json'])['criteria_set']['code'],'QD26-10')
            self.assertEqual(con.execute('SELECT COUNT(*) FROM audit_logs WHERE module=%s',('ocop',)).fetchone()[0],4)
            row=con.execute("INSERT INTO app.users(username,role,active,created_at) VALUES('next_user','staff',TRUE,'date') RETURNING id").fetchone()
            self.assertGreater(row[0],3)
            con.execute((Path(server.__file__).parent/'database/sql/009_staff_role.sql').read_text())
            self.assertEqual(con.execute('SELECT COUNT(*) FROM users').fetchone()[0],4)

    def test_salt_foundations_are_one_traditional_method_and_latest_month(self):
        self.migrate()
        with self.connection() as con:
            rows=con.execute('SELECT * FROM DN_SanLuongMuoi ORDER BY PhuongPhapSX').fetchall()
            self.assertEqual(len(rows),1)
            by={r['PhuongPhapSX']:r for r in rows}
            self.assertEqual((by['Truyền thống']['DienTich'],by['Truyền thống']['SanLuong']),(5,30))
            self.assertIsNone(by['Truyền thống']['GiaBanBinhQuan'])
            con.execute("INSERT INTO app.records(report_date,unit_name,created_by,area_land,status,created_at,updated_at,ma_don_vi_hanh_chinh) VALUES('2026-08-01','Xã An Thới Đông',2,999,'approved','d','d','27673')")
            migration.repair_salt(con)
            self.assertEqual(con.execute('SELECT SUM(DienTich) FROM DN_SanLuongMuoi').fetchone()[0],5)

    def test_dry_run_and_identity_collision_roll_back_every_table(self):
        self.migrate(dry_run=True)
        with self.connection() as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM app.users').fetchone()[0],0)
            con.execute("INSERT INTO app.users(id,username,role,created_at) VALUES(1,'different_user','admin','old')")
        with self.assertRaisesRegex(RuntimeError,'ID collision'):
            self.migrate()
        with self.connection() as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM app.users').fetchone()[0],1)
            self.assertEqual(con.execute('SELECT COUNT(*) FROM qd5277.DM_CoSo').fetchone()[0],0)

    def test_missing_ocop_source_tables_do_not_create_workflow_data(self):
        # A separate salt-only source: no fabricated OCOP rows.
        source=Path(self.temp.name)/'salt_only.db'
        with sqlite3.connect(source) as con:
            con.executescript(legacy_schema())
        with self.connection() as con:
            result=migration.migrate(source,target=con)
            self.assertNotIn('ocop_products',result)
            self.assertEqual(con.execute('SELECT COUNT(*) FROM ocop_products').fetchone()[0],0)
            self.assertEqual(con.execute('SELECT COUNT(*) FROM PTNT_OCOP').fetchone()[0],0)

    def test_excel_etl_upsert_keeps_staging_and_aggregates_foundations(self):
        from openpyxl import Workbook
        from psycopg.conninfo import conninfo_to_dict
        from database.etl import import_diem_nghiep as etl
        workbook=Path(self.temp.name)/'fixture.xlsx'
        wb=Workbook();ws=wb.active;ws.title=etl.SHEET_NAME
        for row,name in zip(etl.DATA_ROWS,server.UNITS):
            ws.cell(row,2,name)
            for col,value in ((3,5),(4,2),(5,3),(6,30),(7,10),(8,20),(20,1200),(21,'2.000 - 2.500')):
                ws.cell(row,col,value)
        wb.save(workbook);wb.close()
        options=conninfo_to_dict(self.dsn);options['dbname']=self.name
        with patch.object(etl,'connection_kwargs',lambda:options):
            self.assertEqual(etl.import_workbook(workbook),(8,8))
            with self.connection() as con:
                # An unrelated source sheet must survive repeat import.
                con.execute("INSERT INTO staging.diem_nghiep_2026_w34_raw(source_sheet,excel_row,snapshot_date,source_workbook,value_kinds) VALUES('other-sheet',99,'2026-08-21','keep.xlsx','{}'::jsonb)")
            self.assertEqual(etl.import_workbook(workbook),(8,8))
        with self.connection() as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM staging.diem_nghiep_2026_w34_raw').fetchone()[0],9)
            self.assertEqual(con.execute('SELECT COUNT(*) FROM DN_SanLuongMuoi').fetchone()[0],8)
            self.assertEqual(con.execute("SELECT SUM(DienTich) FROM DN_SanLuongMuoi WHERE PhuongPhapSX='Truyền thống'").fetchone()[0],40)

    def test_weekly_preview_commit_is_qd_shaped_and_rejects_duplicate_file(self):
        from openpyxl import Workbook
        self.migrate()
        workbook=Path(self.temp.name)/'weekly.xlsx'
        wb=Workbook();ws=wb.active;ws.title='Tuan 34'
        ws.append(['STT','Đơn vị'])
        for index,name in enumerate(('Xã An Thới Đông','Xã Thạnh An'),1):
            row=[None]*28;row[0]=index;row[1]=name
            row[2]=5;row[3]=2;row[4]=3;row[5]=30;row[6]=10;row[7]=20
            ws.append(row)
        wb.save(workbook);wb.close()
        preview=parse_weekly_workbook(workbook.read_bytes(),'2026-08-21','Tuan 34',workbook.name,server.canonical_admin_unit)
        self.assertEqual((preview['week_code'],preview['error_rows']),('2026-W34',0))
        with self.connection() as raw:
            con=backend_db.CompatConnection(raw)
            result=commit_weekly_preview(con,{'user_id':3},preview,'skip',server.now_text())
            self.assertEqual(result['inserted'],2)
        with self.connection() as raw:
            rows=raw.execute('SELECT * FROM app.v_dn_sanluongmuoi_weekly_qd5277 ORDER BY Ma_DonViHanhChinh').fetchall()
            self.assertEqual(len(rows),2)
            self.assertEqual((rows[0]['Ma_ThoiGian'],rows[0]['PhuongPhapSX'],rows[0]['DienTich'],rows[0]['SanLuong']),('2026-W34','Truyền thống',5,30))
            con=backend_db.CompatConnection(raw)
            with self.assertRaisesRegex(WeeklyImportError,'đã được import'):
                commit_weekly_preview(con,{'user_id':3},preview,'skip',server.now_text())

    def test_criteria_parent_order_options_and_source_pk_are_preserved(self):
        with sqlite3.connect(self.source) as con:
            con.execute("INSERT INTO ocop_criteria(id,criteria_set_id,parent_id,code,title,item_type,section_code,max_score,sort_order) VALUES(200,10,28,'A.group','Group','group','A',10,1)")
            con.execute("INSERT INTO ocop_criteria(id,criteria_set_id,parent_id,code,title,item_type,section_code,max_score,sort_order) VALUES(100,10,200,'A.child','Child','criterion','A',10,2)")
            con.execute("INSERT INTO ocop_criteria_options(id,criterion_id,label,score,is_eliminating) VALUES(91,100,'Option',5,1)")
        self.migrate();self.migrate()
        with self.connection() as con:
            row=con.execute('SELECT parent_id FROM ocop_criteria WHERE id=100').fetchone()
            self.assertEqual(row[0],200)
            row=con.execute('SELECT criterion_id,is_eliminating FROM ocop_criteria_options WHERE id=91').fetchone()
            self.assertEqual((row[0],row[1]),(100,True))

    def start_http(self):
        self.migrate()
        self.addCleanup(patch.stopall)
        patch.object(server,'DB_PATH',server.DEFAULT_SQLITE_PATH).start()
        patch.dict(os.environ,{'SALT_WEB_BACKEND':'postgres'}).start()
        patch.object(server,'compat_connect',lambda:backend_db.CompatConnection(self.connection(autocommit=True))).start()
        server.SESSIONS.clear()
        class QuietHandler(server.Handler):
            def log_message(self,*args): pass
        http=server.ThreadingHTTPServer(('127.0.0.1',0),QuietHandler)
        thread=threading.Thread(target=http.serve_forever,daemon=True);thread.start()
        def stop():
            http.shutdown();http.server_close();thread.join()
        self.addCleanup(stop)
        return tuple(Client(http.server_address[1]).login(name) for name in ('admin_test','staff_test','unit_a'))

    def test_postgres_http_staff_account_activation_and_inactive_login(self):
        admin,staff,unit=self.start_http()
        self.assertEqual(admin.request('GET','/users')[0],200)
        for path in ('/users','/users/2/edit','/ocop/access'):
            self.assertEqual(staff.request('GET',path)[0],403)
        for action in ('new','2/edit','2/delete','2/activate','2/deactivate'):
            self.assertEqual(staff.request('POST','/users/'+action,{})[0],403)
        self.assertEqual(staff.request('GET','/standard-data')[0],200)
        self.assertEqual(admin.request('POST','/users/2/deactivate',{})[0],303)
        self.assertEqual(unit.request('GET','/dashboard')[0],303)
        body=unit.request('POST','/login',{'username':'unit_a','password':'Testing-OCOP-2026'})[2].decode()
        self.assertIn('Tài khoản không hoạt động.',body)
        self.assertNotIn('Tài khoản không hoạt động.',unit.request('POST','/login',{'username':'unit_a','password':'bad'})[2].decode())
        self.assertEqual(admin.request('POST','/users/2/activate',{})[0],303)
        unit.login('unit_a')

    def test_postgres_http_ocop_mutations_commit_and_snapshot_dates_serialize(self):
        admin,staff,unit=self.start_http()
        for route in ('/ocop','/ocop/criteria','/ocop/criteria/10','/ocop/products/new'):
            status,_,body=staff.request('GET',route);self.assertEqual(status,200,body.decode())
        with self.connection() as con:
            app=con.execute('SELECT * FROM ocop_applications').fetchone()
        result=staff.request('POST',f'/ocop/applications/{app["id"]}/return',{'revision':app['revision'],'comment':'Bổ sung'})
        self.assertEqual(result[0],303,result[2].decode())
        with self.connection() as con:
            app=con.execute('SELECT * FROM ocop_applications').fetchone()
            self.assertEqual(app['status'],'returned')
        result=unit.request('POST',f'/ocop/applications/{app["id"]}/submit',{'revision':app['revision']})
        self.assertEqual(result[0],303,result[2].decode())
        with self.connection() as con:
            app=con.execute('SELECT * FROM ocop_applications').fetchone()
            self.assertEqual(app['status'],'submitted')
            self.assertEqual(json.loads(app['submission_snapshot_json'])['criteria_set']['code'],'QD26-10')
            self.assertEqual(con.execute('SELECT COUNT(*) FROM PTNT_OCOP').fetchone()[0],0)
