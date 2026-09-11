"""Regression tests for OCOP historical Excel import (T2)."""
from __future__ import annotations

from datetime import date
import io
import os
from pathlib import Path
import unittest
from uuid import uuid4

import backend_db
import ocop_import


ROOT = Path(__file__).resolve().parents[1]
MIGRATIONS = (
    "002_qd5277_schema.sql",
    "004_app_schema.sql",
    "005_monthly_salt_sync.sql",
    "006_official_admin_units.sql",
    "007_ocop_init_only.sql",
    "008_ocop_dynamic_criteria.sql",
    "009_staff_role.sql",
    "010_weekly_salt_imports.sql",
    "011_weekly_foundation.sql",
    "012_admin_units.sql",
    "013_ocop_legacy_import.sql",
)


def workbook_bytes(*, unit="Tân Nhựt", second_product=True):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Loc"
    ws.cell(3, 1, "TT")
    ws.cell(3, 2, "Tên sản phẩm")
    ws.cell(4, 6, "Xã, phường")

    row = [
        1, "Sản phẩm A", "Thực phẩm", "HTX", "Hợp tác xã T2", unit,
        "1 Đường thử nghiệm", "Nguyễn Văn A", "0909000000",
        4, date(2029, 1, 1), 3, date(2024, 1, 2), 2024, "100/QĐ", "UBND TPHCM",
        4, date(2026, 2, 3), 2026, "200/QĐ", "UBND TPHCM", "Đánh giá lại/nâng hạng",
    ]
    for col, value in enumerate(row, 1):
        ws.cell(5, col, value)
    if second_product:
        row2 = [
            2, "Sản phẩm B", "Thực phẩm", None, None, None,
            None, None, None,
            3, date(2028, 5, 1), 3, date(2025, 5, 1), 2025, "300/QĐ", "UBND TPHCM",
            None, None, None, None, None, None,
        ]
        for col, value in enumerate(row2, 1):
            ws.cell(6, col, value)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def unit_lookup(name):
    if str(name).strip() == "Xã Tân Nhựt":
        return "Xã Tân Nhựt", "27595"
    return None


class OcopWorkbookParserTests(unittest.TestCase):
    def test_parser_preserves_history_and_merged_subject_context(self):
        preview = ocop_import.parse_ocop_workbook(
            workbook_bytes(), "CCPTNT OCOP.xlsx", unit_lookup, "Loc"
        )
        self.assertEqual(preview["product_count"], 2)
        self.assertEqual(preview["entity_count"], 1)
        self.assertEqual(preview["recognition_count"], 3)
        self.assertEqual(preview["error_rows"], 0)
        first = preview["rows"][0]["canonical"]
        second = preview["rows"][1]["canonical"]
        self.assertEqual(first["entity"]["unit_code"], "27595")
        self.assertEqual(second["entity"]["name"], "Hợp tác xã T2")
        self.assertEqual(first["recognitions"][1]["evaluation_type"], "upgrade")
        self.assertTrue(first["recognitions"][1]["is_current"])
        self.assertFalse(first["recognitions"][0]["is_current"])

    def test_unknown_unit_blocks_publish_but_preview_keeps_source(self):
        preview = ocop_import.parse_ocop_workbook(
            workbook_bytes(unit="Địa bàn chưa có", second_product=False),
            "CCPTNT OCOP.xlsx", unit_lookup, "Loc",
        )
        self.assertEqual(preview["error_rows"], 1)
        self.assertIn("Địa bàn chưa có", preview["unknown_units"])
        with self.assertRaises(ocop_import.OcopImportError):
            ocop_import._assert_publishable(preview)


@unittest.skipUnless(os.getenv("OCOP_TEST_PG_DSN"),
                     "Set OCOP_TEST_PG_DSN to enable disposable PostgreSQL tests")
class OcopLegacyPostgreSQLTests(unittest.TestCase):
    def setUp(self):
        import psycopg
        from psycopg import sql

        self.psycopg = psycopg
        self.admin_dsn = os.environ["OCOP_TEST_PG_DSN"]
        self.database = "ocop_legacy_" + uuid4().hex
        with psycopg.connect(self.admin_dsn, autocommit=True) as admin:
            admin.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(self.database)))
        self.addCleanup(self.drop_database)
        raw = self.connection()
        try:
            for name in MIGRATIONS:
                raw.execute((ROOT / "database" / "sql" / name).read_text(encoding="utf-8"))
        finally:
            raw.close()

    def drop_database(self):
        from psycopg import sql
        with self.psycopg.connect(self.admin_dsn, autocommit=True) as admin:
            admin.execute(sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(self.database)))

    def connection(self):
        return self.psycopg.connect(
            self.admin_dsn,
            dbname=self.database,
            autocommit=True,
            row_factory=backend_db.hybrid_row,
            options="-c search_path=app,qd5277,staging,public",
        )

    def compat(self):
        return backend_db.CompatConnection(self.connection())

    def create_user(self, con, username, role, unit_code=None):
        uid = con.execute(
            "INSERT INTO users(username,role,active,created_at) VALUES(?,?,TRUE,'2026-09-11 00:00:00')",
            (username, role),
        ).lastrowid
        if unit_code:
            con.execute(
                "INSERT INTO user_admin_units(user_id,ma_don_vi_hanh_chinh) VALUES(?,?)",
                (uid, unit_code),
            )
            con.execute(
                "UPDATE users SET unit_name=(SELECT TenDonVi FROM DM_DonViHanhChinh WHERE Ma_DonViHanhChinh=?) WHERE id=?",
                (unit_code, uid),
            )
        con.commit()
        return uid

    def test_publish_is_transactional_idempotent_and_does_not_fake_applications(self):
        con = self.compat()
        try:
            staff_id = self.create_user(con, "ocop_staff_t2", "staff")
            session = {"user_id": staff_id, "role": "staff", "unit_name": "Chi cục"}
            import admin_units
            preview = ocop_import.parse_ocop_workbook(
                workbook_bytes(), "CCPTNT OCOP.xlsx", admin_units.unit_lookup(con), "Loc"
            )
            result = ocop_import.commit_ocop_preview(con, session, preview, "publish")
            con.commit()
            self.assertEqual(result["published_products"], 2)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_entities").fetchone()[0], 1)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_products").fetchone()[0], 2)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_recognitions").fetchone()[0], 3)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM PTNT_OCOP").fetchone()[0], 2)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_applications").fetchone()[0], 0)

            current = con.execute(
                """SELECT p.ten_san_pham,p.current_star,r.star_rank,r.recognition_sequence,r.decision_number
                   FROM ocop_products p JOIN ocop_recognitions r ON r.product_id=p.id AND r.is_current=TRUE
                   WHERE p.ten_san_pham='Sản phẩm A'"""
            ).fetchone()
            self.assertEqual((current["current_star"], current["star_rank"], current["recognition_sequence"]), (4, 4, 2))
            self.assertEqual(current["decision_number"], "200/QĐ")
            qd = con.execute(
                "SELECT XepHang,Ma_ThoiGian,TrangThai FROM PTNT_OCOP WHERE TenSanPham='Sản phẩm A'"
            ).fetchone()
            self.assertEqual(qd["XepHang"], "4*")
            self.assertEqual(qd["Ma_ThoiGian"], "2026")
            self.assertEqual(qd["TrangThai"], "HIEULUC")

            again = ocop_import.commit_ocop_preview(con, session, preview, "publish")
            con.commit()
            self.assertTrue(again["already_published"])
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_recognitions").fetchone()[0], 3)
        finally:
            con.close()

    def test_stage_only_accepts_unmapped_source_without_publishing(self):
        con = self.compat()
        try:
            admin_id = self.create_user(con, "ocop_admin_t2", "admin")
            session = {"user_id": admin_id, "role": "admin", "unit_name": "Chi cục"}
            import admin_units
            preview = ocop_import.parse_ocop_workbook(
                workbook_bytes(unit="Địa bàn chưa có", second_product=False),
                "unknown.xlsx", admin_units.unit_lookup(con), "Loc",
            )
            result = ocop_import.commit_ocop_preview(con, session, preview, "stage_only")
            con.commit()
            self.assertEqual(result["published_products"], 0)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_import_rows").fetchone()[0], 1)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM ocop_products").fetchone()[0], 0)
            row = con.execute("SELECT validation_status,mapped_unit_code FROM ocop_import_rows").fetchone()
            self.assertEqual(row["validation_status"], "error")
            self.assertIsNone(row["mapped_unit_code"])
        finally:
            con.close()


if __name__ == "__main__":
    unittest.main()
