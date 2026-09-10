"""Preview and commit weekly salt workbooks without publishing QD 5277 data."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import re


class WeeklyImportError(ValueError):
    pass


NUMERIC_COLUMNS = {
    "dien_tich": 3, "area_land": 4, "area_tarp": 5,
    "san_luong": 6, "harvest_land": 7, "harvest_tarp": 8,
    "sold_total": 9, "sold_land": 10, "sold_tarp": 11,
    "remaining_total": 12, "remaining_land": 13, "remaining_tarp": 14,
    "processed_total": 15, "processed_fine": 16, "processed_iodized": 17,
    "households": 18, "workers": 19,
    "damage_total": 23, "damage_land": 24, "damage_tarp": 25,
}


def _text(value):
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _number(value, label):
    if value in (None, "", "-"):
        return Decimal("0")
    if isinstance(value, bool):
        raise WeeklyImportError(f"{label}: giá trị đúng/sai không phải số.")
    token = str(value).strip().replace(" ", "")
    if "." in token and "," in token:
        token = token.replace(".", "").replace(",", ".") if token.rfind(",") > token.rfind(".") else token.replace(",", "")
    elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", token):
        token = token.replace(".", "")
    elif "," in token:
        token = token.replace(",", ".")
    try:
        number = Decimal(token)
    except InvalidOperation as exc:
        raise WeeklyImportError(f"{label}: {value!r} không phải số.") from exc
    if not number.is_finite() or number < 0:
        raise WeeklyImportError(f"{label}: phải là số hữu hạn không âm.")
    return number.quantize(Decimal(".01"))


def _week_code(report_date):
    day = date.fromisoformat(str(report_date))
    iso_year, iso_week, _ = day.isocalendar()
    return day, f"{iso_year}-W{iso_week:02d}"


def parse_weekly_workbook(file_bytes, report_date, sheet_name, filename, unit_lookup):
    """Return a serializable preview; invalid rows are shown but cannot commit."""
    try:
        day, week_code = _week_code(report_date)
    except ValueError as exc:
        raise WeeklyImportError("Ngày kết thúc tuần không hợp lệ.") from exc
    try:
        from openpyxl import load_workbook
        raw_book = load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
        value_book = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise WeeklyImportError(f"Không đọc được file Excel: {exc}") from exc
    try:
        selected = sheet_name.strip() if sheet_name else value_book.active.title
        if selected not in value_book.sheetnames:
            raise WeeklyImportError(
                f"Không tìm thấy sheet '{selected}'. Có: {', '.join(value_book.sheetnames)}"
            )
        raw_sheet, sheet = raw_book[selected], value_book[selected]
        rows, seen_units = [], set()
        # In read-only mode, repeated worksheet.cell(row, column) calls can
        # re-scan the XML stream. Some source sheets are formatted down to row
        # 1000, so read both workbooks linearly once instead.
        raw_rows = raw_sheet.iter_rows(min_row=1, max_col=28, values_only=True)
        value_rows = sheet.iter_rows(min_row=1, max_col=28, values_only=True)
        for row_number, (raw_values, values_row) in enumerate(zip(raw_rows, value_rows), 1):
            unit_raw = _text(values_row[1])
            folded = unit_raw.casefold()
            if not unit_raw or "tổng cộng" in folded or not folded.startswith(("xã ", "phường ", "thị trấn ")):
                continue
            errors, warnings = [], []
            official = unit_lookup(unit_raw)
            if not official:
                errors.append(f"Đơn vị chưa có trong danh mục chính thức: {unit_raw}.")
                unit_name, unit_code = unit_raw, None
            else:
                unit_name, unit_code = official
                if unit_code in seen_units:
                    errors.append(f"Đơn vị xuất hiện nhiều lần trong cùng sheet: {unit_name}.")
                seen_units.add(unit_code)
            values = {}
            for name, column in NUMERIC_COLUMNS.items():
                try:
                    values[name] = float(_number(values_row[column - 1], f"{name} (cột {column})"))
                except WeeklyImportError as exc:
                    errors.append(str(exc))
                    values[name] = 0.0
            for total, left, right, label in (
                ("dien_tich", "area_land", "area_tarp", "Diện tích"),
                ("san_luong", "harvest_land", "harvest_tarp", "Sản lượng thu hoạch"),
                ("sold_total", "sold_land", "sold_tarp", "Sản lượng tiêu thụ"),
                ("remaining_total", "remaining_land", "remaining_tarp", "Sản lượng còn lại"),
                ("processed_total", "processed_fine", "processed_iodized", "Sản lượng chế biến"),
                ("damage_total", "damage_land", "damage_tarp", "Thiệt hại"),
            ):
                difference = values[total] - values[left] - values[right]
                if abs(difference) > 0.01:
                    warnings.append(f"{label}: tổng lệch chi tiết {difference:+.2f}.")
            canonical = {
                "week_code": week_code, "report_date": day.isoformat(),
                "unit_name": unit_name, "ma_don_vi_hanh_chinh": unit_code,
                "phuong_phap_sx": "Truyền thống", "gia_ban_binh_quan": None,
                **values,
                "households": int(values["households"]), "workers": int(values["workers"]),
                "price_land": _text(values_row[19]),
                "price_tarp": _text(values_row[20]),
                "note": _text(values_row[26]),
            }
            raw_data = {
                str(column): {
                    "raw": _json_value(raw_values[column - 1]),
                    "value": _json_value(values_row[column - 1]),
                }
                for column in range(1, 29)
            }
            rows.append({
                "excel_row": row_number, "unit_name_raw": unit_raw,
                "status": "error" if errors else ("warning" if warnings else "valid"),
                "errors": errors, "warnings": warnings,
                "raw_data": raw_data, "canonical": canonical,
            })
        if not rows:
            raise WeeklyImportError("Không tìm thấy dòng xã/phường trong sheet đã chọn.")
        return {
            "filename": filename, "file_sha256": hashlib.sha256(file_bytes).hexdigest(),
            "sheet_name": selected, "week_code": week_code, "report_date": day.isoformat(),
            "rows": rows,
            "valid_rows": sum(row["status"] == "valid" for row in rows),
            "warning_rows": sum(row["status"] == "warning" for row in rows),
            "error_rows": sum(row["status"] == "error" for row in rows),
        }
    finally:
        raw_book.close()
        value_book.close()


def commit_weekly_preview(con, session, preview, mode, now):
    """Persist a validated preview. Caller owns commit/rollback."""
    if preview.get("error_rows"):
        raise WeeklyImportError("Preview còn dòng lỗi; chưa thể import.")
    if mode not in ("skip", "update"):
        raise WeeklyImportError("Chế độ xử lý dữ liệu cũ không hợp lệ.")
    duplicate = con.execute(
        "SELECT id FROM salt_import_batches WHERE file_sha256=? AND sheet_name=? AND week_code=?",
        (preview["file_sha256"], preview["sheet_name"], preview["week_code"]),
    ).fetchone()
    if duplicate:
        raise WeeklyImportError("File này đã được import cho cùng tuần và sheet.")
    cursor = con.execute(
        """INSERT INTO salt_import_batches
           (filename,file_sha256,sheet_name,week_code,report_date,import_mode,total_rows,
            warning_rows,imported_by,imported_at)
           VALUES(?,?,?,?,?,?,?,?,?,?)""",
        (preview["filename"], preview["file_sha256"], preview["sheet_name"],
         preview["week_code"], preview["report_date"], mode, len(preview["rows"]),
         preview["warning_rows"], session["user_id"], now),
    )
    batch_id = cursor.lastrowid
    inserted = updated = skipped = 0
    record_columns = (
        "batch_id", "week_code", "report_date", "unit_name", "ma_don_vi_hanh_chinh",
        "phuong_phap_sx", "dien_tich", "san_luong", "gia_ban_binh_quan",
        "area_land", "area_tarp", "harvest_land", "harvest_tarp", "sold_land",
        "sold_tarp", "remaining_land", "remaining_tarp", "processed_fine",
        "processed_iodized", "households", "workers", "price_land", "price_tarp",
        "damage_land", "damage_tarp", "note", "created_by", "created_at", "updated_at",
    )
    for row in preview["rows"]:
        canonical = row["canonical"]
        con.execute(
            """INSERT INTO salt_weekly_import_rows
               (batch_id,excel_row,unit_name_raw,ma_don_vi_hanh_chinh,validation_status,
                validation_messages_json,raw_data_json,canonical_data_json)
               VALUES(?,?,?,?,?,?,?,?)""",
            (batch_id, row["excel_row"], row["unit_name_raw"], canonical["ma_don_vi_hanh_chinh"],
             row["status"], json.dumps(row["warnings"], ensure_ascii=False),
             json.dumps(row["raw_data"], ensure_ascii=False),
             json.dumps(canonical, ensure_ascii=False)),
        )
        existing = con.execute(
            "SELECT id FROM salt_weekly_records WHERE ma_don_vi_hanh_chinh=? AND week_code=?",
            (canonical["ma_don_vi_hanh_chinh"], preview["week_code"]),
        ).fetchone()
        payload = {column: canonical.get(column) for column in record_columns}
        payload.update(batch_id=batch_id, created_by=session["user_id"], created_at=now, updated_at=now)
        if existing and mode == "skip":
            skipped += 1
            continue
        if existing:
            payload["id"] = existing["id"]
            assignments = ",".join(f"{column}=:{column}" for column in record_columns if column not in ("created_by", "created_at"))
            con.execute(f"UPDATE salt_weekly_records SET {assignments} WHERE id=:id", payload)
            updated += 1
        else:
            names = ",".join(record_columns)
            placeholders = ",".join(f":{column}" for column in record_columns)
            con.execute(f"INSERT INTO salt_weekly_records({names}) VALUES({placeholders})", payload)
            inserted += 1
    con.execute(
        "UPDATE salt_import_batches SET imported_rows=?,updated_rows=?,skipped_rows=? WHERE id=?",
        (inserted, updated, skipped, batch_id),
    )
    return {"batch_id": batch_id, "inserted": inserted, "updated": updated, "skipped": skipped}
