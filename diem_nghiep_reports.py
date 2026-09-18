"""Resolve and export the fixed Diêm nghiệp reports.

The weekly salt tables remain the source of truth.  This module deliberately
keeps report rules outside the HTTP page so the HTML preview and the Excel
export consume the same normalized object.
"""

from __future__ import annotations

from calendar import monthrange
from copy import copy
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from pathlib import Path
from typing import Any
import unicodedata


REPORT_CUC = "CUC"
REPORT_SO = "SO"
REPORT_TYPES = (REPORT_CUC, REPORT_SO)

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "assets" / "templates" / "diem-nghiep"
TEMPLATE_PATHS = {
    REPORT_CUC: TEMPLATE_DIR / "template_bao_cao_cuc_diem_nghiep.xlsx",
    REPORT_SO: TEMPLATE_DIR / "template_bao_cao_so_diem_nghiep.xlsx",
}

CUC_LOCALITIES = (
    "Xã An Thời Đông",
    "Xã Thạnh An",
    "Xã Cần Giờ",
    "Xã Long Điền",
    "Xã Long Sơn",
    "Phường Phước Thắng",
    "Phường Long Hương",
    "Phường Bà Rịa",
)

CUC_COLUMNS = (
    "A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"
)

# This is the visible-sheet mapping from template_bao_cao_so_diem_nghiep.xlsx.
# Rows not listed here are section/heading rows or fields deliberately sourced
# from a future PTNT/master-data input and therefore remain blank.
SO_ROW_SPECS = (
    (6, "AREA_TOTAL", "DELTA_MONTH", "SNAPSHOT", "Diện tích sản xuất muối", "ha"),
    (7, "AREA_MANUAL", "DELTA_MONTH", "AREA_TOTAL", "- Diện tích muối sản xuất thủ công", "ha"),
    (8, "AREA_INDUSTRIAL", "DERIVED_ZERO", "DERIVED_ZERO", "- Diện tích muối sản xuất công nghiệp", "ha"),
    (9, "HARVEST_TOTAL", "DELTA_MONTH", "SNAPSHOT", "Sản lượng muối", "tấn"),
    (10, "HARVEST_MANUAL", "DELTA_MONTH", "HARVEST_TOTAL", "- Sản lượng muối sản xuất thủ công", "tấn"),
    (11, "HARVEST_INDUSTRIAL", "DERIVED_ZERO", "DERIVED_ZERO", "- Sản lượng muối sản xuất công nghiệp", "tấn"),
    (12, "CONSUME_TOTAL", "DELTA_MONTH", "SNAPSHOT", "- Sản lượng muối tiêu thụ", "tấn"),
    (13, "HOUSEHOLDS", "SNAPSHOT", "SNAPSHOT", "Hộ sản xuất muối", "hộ"),
    (15, "LOST_AREA", "DELTA_MONTH", "SNAPSHOT", "- Diện tích sản xuất muối bị mất trắng", "ha"),
    (18, "VALUE_SHARE_MANUAL", "MASTER_DATA", "MASTER_DATA", "- Muối thủ công", "%"),
    (19, "VALUE_SHARE_INDUSTRIAL", "MASTER_DATA", "MASTER_DATA", "- Muối công nghiệp", "%"),
    (21, "LAND_MANUAL", "MASTER_DATA", "MASTER_DATA", "- Diện tích đất làm muối thủ công", "ha"),
    (22, "LAND_INDUSTRIAL", "MASTER_DATA", "MASTER_DATA", "- Diện tích đất làm muối công nghiệp", "ha"),
    (24, "AREA_MANUAL", "DELTA_MONTH", "AREA_TOTAL", "- Diện tích sản xuất muối thủ công", "ha"),
    (25, "AREA_INDUSTRIAL", "DERIVED_ZERO", "DERIVED_ZERO", "- Diện tích sản xuất muối công nghiệp", "ha"),
    (26, "HARVEST_TOTAL", "DELTA_MONTH", "SNAPSHOT", "Sản lượng muối sản xuất", "1.000 tấn"),
    (27, "HARVEST_MANUAL", "DELTA_MONTH", "HARVEST_TOTAL", "- Muối thủ công", "1.000 tấn"),
    (28, "HARVEST_INDUSTRIAL", "DERIVED_ZERO", "DERIVED_ZERO", "- Muối công nghiệp", "1.000 tấn"),
    (30, "SAFE_TECH_MANUAL", "MASTER_DATA", "MASTER_DATA", "- Diện tích sản xuất muối thủ công", "ha"),
    (31, "SAFE_TECH_INDUSTRIAL", "MASTER_DATA", "MASTER_DATA", "- Diện tích sản xuất muối công nghiệp", "ha"),
    (33, "WAREHOUSE_COUNT", "MASTER_DATA", "MASTER_DATA", "- Kho dự trữ, bảo quản muối", "kho"),
    (34, "CLEAN_PROCESSING_COUNT", "MASTER_DATA", "MASTER_DATA", "- Chế biến muối sạch", "cơ sở"),
    (36, "HTX_SALT_COUNT", "MASTER_DATA", "MASTER_DATA", "Số hợp tác xã diêm nghiệp", "HTX"),
    (37, "LHHTX_SALT_COUNT", "MASTER_DATA", "MASTER_DATA", "Số liên hiệp HTX diêm nghiệp", "LHHTX"),
    (38, "THT_SALT_COUNT", "MASTER_DATA", "MASTER_DATA", "Số tổ hợp tác diêm nghiệp", "THT"),
    (39, "ENTERPRISE_SALT_COUNT", "MASTER_DATA", "MASTER_DATA", "Số doanh nghiệp diêm nghiệp", "DN"),
    (40, "FARM_SALT_COUNT", "MASTER_DATA", "MASTER_DATA", "Số trang trại sản xuất muối", "trang trại"),
)

# STT is a business/template field, not the worksheet row number.
SO_STT_BY_ROW = {
    6: "1", 7: "", 8: "",
    9: "2", 10: "", 11: "", 12: "",
    13: "3", 14: "4", 15: "",
    17: "1", 18: "", 19: "",
    20: "2", 21: "", 22: "",
    23: "3", 24: "", 25: "",
    26: "4", 27: "", 28: "",
    29: "5", 30: "", 31: "",
    32: "6", 33: "", 34: "",
    36: "1", 37: "2", 38: "3", 39: "4", 40: "5",
}

REPORT_TERM_DEFINITIONS = (
    (
        "Chính thức tháng trước",
        "Số liệu riêng của tháng trước đã được chốt hoặc xác nhận; không phải số lũy kế.",
    ),
    (
        "Ước tháng báo cáo",
        "Số liệu riêng của tháng đang báo cáo, thường được suy ra từ các snapshot báo cáo tuần khi chưa có số chính thức.",
    ),
    (
        "Lũy kế",
        "Tổng số từ đầu niên vụ đến ngày của snapshot báo cáo tuần.",
    ),
    (
        "Snapshot báo cáo tuần",
        "Bộ số liệu được chốt tại một ngày báo cáo tuần; hệ thống chọn snapshot mới nhất trong tháng làm nguồn của tháng đó.",
    ),
    (
        "Delta lũy kế",
        "Lũy kế kỳ này trừ lũy kế kỳ trước. Chỉ dùng để suy ra số riêng của tháng khi nguồn không có trường tháng trực tiếp.",
    ),
)

REPORT_HEADER_COMMENTS = {
    "D4": "Chính thức tháng cùng kỳ năm trước: số riêng của cùng tháng năm trước đã được chốt.",
    "E4": "Lũy kế chính thức đến cùng kỳ năm trước: tổng đã chốt từ đầu niên vụ đến cùng kỳ năm trước.",
    "F4": "Chính thức tháng trước: số riêng của tháng trước đã được chốt. Đây không phải số lũy kế.",
    "G4": "Ước tháng báo cáo: số riêng của tháng đang báo cáo, có thể được suy ra từ các snapshot báo cáo tuần.",
    "H4": "Ước lũy kế đến tháng báo cáo: tổng lũy kế của snapshot báo cáo tuần mới nhất trong tháng báo cáo.",
}


def _as_date(value: Any) -> date | None:
    if isinstance(value, date):
        return value
    if value is None:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return number if number.is_finite() else None


def _format_report_number(value: Decimal | None) -> str:
    """Format a report number with Vietnamese separators for explanations."""
    number = _decimal(value)
    if number is None:
        return ""
    if number == number.to_integral_value():
        text = f"{int(number):,}"
    else:
        text = f"{number:,.2f}".rstrip("0").rstrip(".")
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def _snapshot_date_text(snapshot: dict | None) -> str:
    value = (snapshot or {}).get("date")
    parsed = _as_date(value)
    return parsed.strftime("%d/%m/%Y") if parsed else ""


def _row_value(row: Any, name: str) -> Any:
    if row is None:
        return None
    try:
        return row[name]
    except (KeyError, TypeError):
        pass
    if isinstance(row, dict):
        wanted = name.casefold()
        for key, value in row.items():
            if str(key).casefold() == wanted:
                return value
    return None


def _source_metric(row: Any, metric: str) -> Decimal | None:
    direct = {
        "households": "households",
        "workers": "workers",
        "area_land": "area_land",
        "area_tarp": "area_tarp",
        "harvest_land": "harvest_land",
        "harvest_tarp": "harvest_tarp",
        "sold_land": "sold_land",
        "sold_tarp": "sold_tarp",
        "remaining_land": "remaining_land",
        "remaining_tarp": "remaining_tarp",
        "processed_fine": "processed_fine",
        "processed_iodized": "processed_iodized",
        "damage_land": "damage_land",
        "damage_tarp": "damage_tarp",
    }
    if metric in direct:
        return _decimal(_row_value(row, direct[metric]))
    pairs = {
        "area_total": ("area_land", "area_tarp"),
        "harvest_total": ("harvest_land", "harvest_tarp"),
        "consume_total": ("sold_land", "sold_tarp"),
        "stock_carry_forward": ("remaining_land", "remaining_tarp"),
        "process_total": ("processed_fine", "processed_iodized"),
        "lost_area": ("damage_land", "damage_tarp"),
    }
    if metric in pairs:
        values = [_decimal(_row_value(row, field)) for field in pairs[metric]]
        values = [value for value in values if value is not None]
        return sum(values, Decimal("0")) if values else None
    return None


def _normal_name(value: Any) -> str:
    text = "" if value is None else str(value).strip().casefold()
    text = "".join(char for char in unicodedata.normalize("NFD", text)
                    if unicodedata.category(char) != "Mn")
    text = re.sub(r"\s+", " ", text)
    return text


def _snapshot_rows(snapshot: dict | None) -> list[Any]:
    return list((snapshot or {}).get("rows") or [])


def _snapshot_metric(snapshot: dict | None, metric: str) -> Decimal | None:
    values = [_source_metric(row, metric) for row in _snapshot_rows(snapshot)]
    values = [value for value in values if value is not None]
    return sum(values, Decimal("0")) if values else None


def select_latest_month_snapshot(rows: list[Any], year: int, month: int) -> dict | None:
    """Return only the latest report date inside the requested month."""
    start = date(year, month, 1)
    end = date(year, month, monthrange(year, month)[1])
    dated = [(_as_date(_row_value(row, "report_date")), row) for row in rows]
    dated = [(report_date, row) for report_date, row in dated if report_date and start <= report_date <= end]
    if not dated:
        return None
    latest = max(report_date for report_date, _ in dated)
    return {"date": latest, "rows": [row for report_date, row in dated if report_date == latest]}


def _load_month_snapshot(con, year: int, month: int) -> dict | None:
    start = date(year, month, 1)
    end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    rows = con.execute(
        "SELECT * FROM salt_weekly_records WHERE report_date >= ? AND report_date < ? "
        "ORDER BY report_date DESC, id DESC", (start, end)
    ).fetchall()
    return select_latest_month_snapshot([dict(row) for row in rows], year, month)


def _previous_period(year: int, month: int, offset: int = 1) -> tuple[int, int]:
    index = year * 12 + month - 1 - offset
    return index // 12, index % 12 + 1


def _lookup_locality(rows: list[Any], label: str) -> Any | None:
    wanted = _normal_name(label)
    for row in rows:
        name = _normal_name(_row_value(row, "unit_name"))
        if name == wanted:
            return row
        if name.removeprefix("xa ").removeprefix("phuong ") == wanted.removeprefix("xa ").removeprefix("phuong "):
            return row
    return None


def _cuc_source_cells(row: Any | None) -> dict[str, Decimal | int | None]:
    area = _source_metric(row, "area_total")
    harvest = _source_metric(row, "harvest_total")
    consume = _source_metric(row, "consume_total")
    process_fine = _source_metric(row, "processed_fine")
    process_iodized = _source_metric(row, "processed_iodized")
    return {
        "C": area,
        "D": area,
        "E": Decimal("0") if row is not None else None,
        "F": harvest,
        "G": harvest,
        "H": Decimal("0") if row is not None else None,
        "I": consume,
        "J": consume,
        "K": Decimal("0") if row is not None else None,
        "L": _source_metric(row, "stock_carry_forward"),
        "M": (process_fine + process_iodized
              if process_fine is not None and process_iodized is not None
              else None),
        "N": process_fine,
        "O": process_iodized,
        "P": _source_metric(row, "households"),
        "Q": _source_metric(row, "workers"),
    }


def _sum_cells(rows: list[dict], column: str) -> Decimal | None:
    values = [row["cells"].get(column) for row in rows]
    values = [value for value in values if value is not None]
    return sum(values, Decimal("0")) if values else None


def _price_numbers(value: Any) -> list[Decimal]:
    numbers = []
    for token in re.findall(r"\d+(?:[.,]\d+)*", str(value or "")):
        if "." in token and "," in token:
            token = token.replace(".", "").replace(",", ".") if token.rfind(",") > token.rfind(".") else token.replace(",", "")
        elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", token):
            token = token.replace(".", "")
        elif "," in token:
            token = token.replace(",", ".")
        try:
            number = Decimal(token)
        except InvalidOperation:
            continue
        if number.is_finite():
            numbers.append(number)
    return numbers


def _cuc_price_display(snapshot: dict | None) -> str | None:
    prices = []
    for row in _snapshot_rows(snapshot):
        prices.extend(_price_numbers(_row_value(row, "price_land")))
        prices.extend(_price_numbers(_row_value(row, "price_tarp")))
    if not prices:
        return None
    low, high = min(prices), max(prices)
    def render(value):
        return f"{int(value):,}".replace(",", ".") if value == value.to_integral_value() else str(value)
    if low == high:
        return f"Giá muối bình quân: {render(low)} đồng/ký"
    return f"Giá muối bình quân: {render(low)} đến {render(high)} đồng/ký"


def _build_cuc(snapshot: dict | None, warnings: list[str]) -> list[dict]:
    if snapshot is None:
        return []
    source_rows = _snapshot_rows(snapshot)
    detail_rows = []
    for index, label in enumerate(CUC_LOCALITIES, 1):
        source = _lookup_locality(source_rows, label)
        if source is None:
            warnings.append(f"Thiếu dữ liệu địa bàn {label} trong snapshot.")
        detail_rows.append({
            "stt": index,
            "name": label,
            "cells": _cuc_source_cells(source),
        })
    total = {
        "stt": "",
        "name": "CẢ TỈNH",
        "cells": {column: _sum_cells(detail_rows, column) for column in CUC_COLUMNS[2:]},
    }
    return [total, *detail_rows]


def _difference(current: Decimal | None, previous: Decimal | None, warnings: list[str], label: str,
                *, warning_details: list[dict] | None = None, row_number: int | None = None,
                column: str | None = None, unit: str = "", scale: Decimal = Decimal("1"),
                current_snapshot: dict | None = None, previous_snapshot: dict | None = None) -> Decimal | None:
    if current is None or previous is None:
        return None
    value = current - previous
    if value < 0:
        warnings.append(f"{label} chưa thể tính do số liệu lũy tiến có điều chỉnh.")
        if warning_details is not None:
            display_current = current / scale
            display_previous = previous / scale
            display_delta = value / scale
            if column == "F":
                logic_note = (
                    "Đây là ô tham chiếu tháng trước; để ra số tháng, hệ thống phải lấy "
                    "lũy kế tháng đó trừ lũy kế tháng liền trước. Số liệu tháng báo cáo "
                    "hiện tại vẫn được lấy độc lập từ snapshot báo cáo tuần của kỳ hiện tại."
                )
            elif column == "G":
                logic_note = (
                    "Đây là ô tháng báo cáo; hệ thống lấy trực tiếp chênh lệch giữa snapshot "
                    "báo cáo tuần của kỳ hiện tại và snapshot kỳ trước."
                )
            else:
                logic_note = (
                    "Hệ thống không dùng kết quả của một báo cáo tháng khác; phép tính này "
                    "được thực hiện trực tiếp từ các snapshot báo cáo tuần."
                )
            current_date_text = _snapshot_date_text(current_snapshot)
            previous_date_text = _snapshot_date_text(previous_snapshot)
            current_date_note = f" ngày {current_date_text}" if current_date_text else ""
            previous_date_note = f" ngày {previous_date_text}" if previous_date_text else ""
            formula = (
                f"{_format_report_number(display_current)} - "
                f"{_format_report_number(display_previous)} = "
                f"{_format_report_number(display_delta)} {unit}"
            )
            warning_details.append({
                "row_number": row_number,
                "column": column,
                "label": label.lstrip("- ").strip(),
                "unit": unit,
                "current_value": display_current,
                "previous_value": display_previous,
                "delta": display_delta,
                "current_date": _snapshot_date_text(current_snapshot),
                "previous_date": _snapshot_date_text(previous_snapshot),
                "formula": formula,
                "logic_explanation": logic_note,
                "message": (
                    f"Lũy kế kỳ này{current_date_note} là {_format_report_number(display_current)} {unit}, "
                    f"thấp hơn lũy kế kỳ trước{previous_date_note} là {_format_report_number(display_previous)} {unit}; "
                    f"chênh lệch = {_format_report_number(display_delta)} {unit}. "
                    "Hệ thống để trống ô tháng và yêu cầu kiểm tra lại số liệu nguồn."
                ),
            })
        return None
    return value


def _so_value(rule: str, metric: str, snapshots: dict[str, dict | None], column: str,
              warnings: list[str], label: str, scale: Decimal = Decimal("1"), *,
              warning_details: list[dict] | None = None, row_number: int | None = None,
              unit: str = "") -> Decimal | None:
    if rule == "MASTER_DATA":
        return None
    if rule == "DERIVED_ZERO":
        return Decimal("0")
    source_metric = {
        "AREA_TOTAL": "area_total",
        "AREA_MANUAL": "area_total",
        "HARVEST_TOTAL": "harvest_total",
        "HARVEST_MANUAL": "harvest_total",
        "CONSUME_TOTAL": "consume_total",
        "HOUSEHOLDS": "households",
        "LOST_AREA": "lost_area",
    }.get(metric, metric.casefold())
    if rule == "DELTA_MONTH":
        if column == "D":
            current_snapshot = snapshots.get("previous_year_current")
            previous_snapshot = snapshots.get("previous_year_previous")
        elif column == "E":
            value = _snapshot_metric(snapshots.get("previous_year_current"), source_metric)
        elif column == "F":
            current_snapshot = snapshots.get("previous_month")
            previous_snapshot = snapshots.get("previous_previous_month")
        elif column == "G":
            current_snapshot = snapshots.get("current")
            previous_snapshot = snapshots.get("previous_month")
        else:
            value = _snapshot_metric(snapshots.get("current"), source_metric)
        if column in ("D", "F", "G"):
            value = _difference(
                _snapshot_metric(current_snapshot, source_metric),
                _snapshot_metric(previous_snapshot, source_metric),
                warnings,
                label,
                warning_details=warning_details,
                row_number=row_number,
                column=column,
                unit=unit,
                scale=scale,
                current_snapshot=current_snapshot,
                previous_snapshot=previous_snapshot,
            )
    else:  # SNAPSHOT and the normalized *_TOTAL rules
        source_by_column = {
            "D": "previous_year_current", "E": "previous_year_current",
            "F": "previous_month", "G": "current", "H": "current",
        }
        value = _snapshot_metric(snapshots.get(source_by_column.get(column, "current")), source_metric)
    return value / scale if value is not None else None


def _build_so(snapshots: dict[str, dict | None], warnings: list[str], warning_details: list[dict]) -> list[dict]:
    rows = []
    for row_number, metric, month_rule, cumulative_rule, label, unit in SO_ROW_SPECS:
        cells: dict[str, Decimal | None] = {}
        for column in ("D", "E", "F", "G", "H"):
            rule = month_rule if column in ("D", "F", "G") else cumulative_rule
            scale = Decimal("1000") if metric.endswith("_KT") or (row_number in (26, 27, 28) and rule != "MASTER_DATA") else Decimal("1")
            normalized_metric = metric.replace("_KT", "")
            cells[column] = _so_value(
                rule,
                normalized_metric,
                snapshots,
                column,
                warnings,
                label,
                scale,
                warning_details=warning_details,
                row_number=row_number,
                unit=unit,
            )
        cells["I"] = (cells["G"] / cells["D"] * Decimal("100")
                       if cells["D"] not in (None, Decimal("0")) and cells["G"] is not None else None)
        cells["J"] = (cells["H"] / cells["E"] * Decimal("100")
                       if cells["E"] not in (None, Decimal("0")) and cells["H"] is not None else None)
        rows.append({"row_number": row_number, "stt": SO_STT_BY_ROW.get(row_number, ""),
                     "label": label, "unit": unit, "cells": cells,
                     "cell_warnings": {
                         detail["column"]: detail
                         for detail in warning_details
                         if detail["row_number"] == row_number
                     }})
    return rows


def build_report_data(report_type: str, year: int, month: int, snapshots: dict[str, dict | None]) -> dict:
    """Build the normalized object consumed by both UI and exporter."""
    report_type = str(report_type or "").upper()
    if report_type not in REPORT_TYPES:
        raise ValueError("Loại báo cáo không hợp lệ.")
    if not 1 <= int(month) <= 12:
        raise ValueError("Tháng báo cáo không hợp lệ.")
    year, month = int(year), int(month)
    warnings: list[str] = []
    warning_details: list[dict] = []
    current = snapshots.get("current")
    if report_type == REPORT_CUC:
        rows = _build_cuc(current, warnings)
    else:
        rows = _build_so(snapshots, warnings, warning_details) if current else []
    previous_year_available = snapshots.get("previous_year_current") is not None
    return {
        "report_type": report_type,
        "year": year,
        "month": month,
        "snapshot_date": (current or {}).get("date"),
        "previous_month_snapshot_date": (snapshots.get("previous_month") or {}).get("date"),
        "previous_year_snapshot_date": (snapshots.get("previous_year_current") or {}).get("date"),
        "previous_year_available": previous_year_available,
        "price_display": _cuc_price_display(current) if report_type == REPORT_CUC else None,
        "rows": rows,
        "warnings": list(dict.fromkeys(warnings)),
        "warning_details": warning_details,
        "can_export": current is not None,
    }


def resolve_diem_nghiep_report(con, *, report_type: str, year: int, month: int) -> dict:
    """Resolve a report from the weekly source without falling across months."""
    year, month = int(year), int(month)
    periods = {
        "current": (year, month),
        "previous_month": _previous_period(year, month),
        "previous_previous_month": _previous_period(year, month, 2),
        "previous_year_current": (year - 1, month),
        "previous_year_previous": _previous_period(year - 1, month),
    }
    snapshots = {name: _load_month_snapshot(con, *period) for name, period in periods.items()}
    return build_report_data(report_type, year, month, snapshots)


def available_report_periods(con) -> list[tuple[int, int]]:
    rows = con.execute("SELECT report_date FROM salt_weekly_records GROUP BY report_date ORDER BY report_date DESC").fetchall()
    seen = set()
    result = []
    for row in rows:
        report_date = _as_date(_row_value(row, "report_date"))
        if report_date and (report_date.year, report_date.month) not in seen:
            seen.add((report_date.year, report_date.month))
            result.append((report_date.year, report_date.month))
    return result


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def _set_calculation_on_load(book) -> None:
    calculation = getattr(book, "calculation", None)
    if calculation is not None:
        calculation.fullCalcOnLoad = True
        calculation.forceFullCalc = True
        calculation.calcMode = "auto"


def _replace_text(value: Any, replacements: dict[str, str]) -> Any:
    if not isinstance(value, str):
        return value
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value


def _update_cuc_header(ws, report_date: date | None) -> None:
    if not report_date:
        return
    display = report_date.strftime("%d/%m/%Y")
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if isinstance(cell.value, str):
                value = _replace_text(cell.value, {"[NGÀY CHỐT BÁO CÁO]": display})
                if "[NGÀY CHỐT BÁO CÁO]" not in value and "ĐẾN" in value:
                    value = re.sub(r"ĐẾN\s+\d{1,2}\s+THÁNG\s+\d{1,2}\s+NĂM\s+\d{4}",
                                   f"ĐẾN {report_date.day:02d}/{report_date.month:02d}/{report_date.year}", value,
                                   flags=re.IGNORECASE)
                cell.value = value


def _update_so_header(ws, year: int, month: int) -> None:
    prev_year, prev_month = _previous_period(year, month)
    labels = {
        "D4": f"Chính thức tháng {month} năm {year - 1}",
        "E4": f"Lũy kế chính thức đến tháng {month} năm {year - 1}",
        "F4": f"Chính thức tháng {prev_month} năm {prev_year}",
        "G4": f"Ước tháng {month} năm {year}",
        "H4": f"Ước lũy kế đến tháng {month} năm {year}",
        "I4": f"Ước tháng {month} năm {year} / Chính thức tháng {month} năm {year - 1} (%)",
        "J4": f"Ước lũy kế đến tháng {month} năm {year} / Lũy kế chính thức đến tháng {month} năm {year - 1} (%)",
    }
    for address, value in labels.items():
        ws[address] = value
    ws["A2"] = f"Kỳ báo cáo: {month:02d}/{year}"


def export_diem_nghiep_report(report_data: dict) -> bytes:
    """Fill the supplied fixed template while preserving its formatting."""
    try:
        from openpyxl import load_workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import PatternFill
    except Exception as exc:
        raise RuntimeError("Máy chủ chưa cài openpyxl.") from exc
    report_type = report_data["report_type"]
    template_path = TEMPLATE_PATHS[report_type]
    if not template_path.is_file():
        raise FileNotFoundError(str(template_path))
    book = load_workbook(template_path, data_only=False)
    ws = book["Bao cao Cuc" if report_type == REPORT_CUC else "Bao cao So"]
    if report_type == REPORT_CUC:
        _update_cuc_header(ws, _as_date(report_data.get("snapshot_date")))
        ws["A16"] = report_data.get("price_display")
        for row_number, row in zip(range(6, 15), report_data.get("rows", [])):
            # The total row is intentionally left untouched: its C:Q formulas
            # are part of the supplied template and must remain formulas.
            if row_number == 6:
                continue
            ws.cell(row_number, 1).value = row.get("stt")
            ws.cell(row_number, 2).value = row.get("name")
            for column in CUC_COLUMNS[2:]:
                ws[f"{column}{row_number}"] = _excel_value(row["cells"].get(column))
    else:
        _update_so_header(ws, int(report_data["year"]), int(report_data["month"]))
        for row in report_data.get("rows", []):
            row_number = row["row_number"]
            for column in ("D", "E", "F", "G", "H"):
                ws[f"{column}{row_number}"] = _excel_value(row["cells"].get(column))
            # The template already contains the correct IF/blank semantics;
            # ensure the formulas survive even when a source cell is blank.
            ws[f"I{row_number}"] = f'=IF(OR(D{row_number}="",G{row_number}="",D{row_number}=0),"",G{row_number}/D{row_number}*100)'
            ws[f"J{row_number}"] = f'=IF(OR(E{row_number}="",H{row_number}="",E{row_number}=0),"",H{row_number}/E{row_number}*100)'
        error_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
        for detail in report_data.get("warning_details", []):
            address = f'{detail["column"]}{detail["row_number"]}'
            cell = ws[address]
            cell.fill = error_fill
            cell_font = copy(cell.font)
            cell_font.bold = True
            cell_font.color = "9C0006"
            cell.font = cell_font
            cell.comment = Comment(
                f'{detail["message"]}\nCách tính: {detail["formula"]}\nGiải thích logic: {detail["logic_explanation"]}',
                "Hệ thống",
            )
        for address, explanation in REPORT_HEADER_COMMENTS.items():
            if ws[address].comment is None:
                ws[address].comment = Comment(explanation, "Hệ thống")
    _set_calculation_on_load(book)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def report_filename(report_type: str, year: int, month: int) -> str:
    label = "Cuc" if report_type == REPORT_CUC else "So"
    return f"Bao_cao_Diem_nghiep_{label}_{int(month):02d}-{int(year)}.xlsx"
