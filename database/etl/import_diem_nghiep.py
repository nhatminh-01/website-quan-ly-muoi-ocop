"""Import the specified Diem nghiep worksheet into staging and QD 5277."""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook
from psycopg.types.json import Jsonb
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from database.etl.mappings import TIME_CODE, lookup_admin_code
from salt_normalization import sync_methods, method_values
from backend_db import CompatConnection, hybrid_row


SHEET_NAME = "21.8-Tuan 34"
DATA_ROWS = range(6, 14)
SNAPSHOT_DATE = date(2026, 8, 21)
COLUMN_LETTERS = [get_column_letter(i).lower() for i in range(1, 29)]


def serialize_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def value_kind(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, str) and value == "-":
        return "dash"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float, Decimal)):
        return "number"
    return "text"


def numeric_cached(cell_value: Any, *, row: int, column: str) -> Decimal:
    if isinstance(cell_value, bool) or not isinstance(cell_value, (int, float, Decimal)):
        raise ValueError(
            f"Giá trị số bắt buộc bị thiếu/sai kiểu tại {SHEET_NAME}!{column}{row}: {cell_value!r}"
        )
    return Decimal(str(cell_value))


def connection_kwargs() -> dict[str, Any]:
    return {
        "host": os.getenv("PGHOST", "localhost"),
        "port": int(os.getenv("PGPORT", "5432")),
        "dbname": os.getenv("PGDATABASE", "ptnt_qd5277_dev"),
        "user": os.getenv("PGUSER", "postgres"),
        "password": os.getenv("PGPASSWORD") or None,
    }


def import_workbook(workbook_path: Path) -> tuple[int, int]:
    formulas = load_workbook(workbook_path, data_only=False, read_only=True)
    cached = load_workbook(workbook_path, data_only=True, read_only=True)
    if SHEET_NAME not in formulas.sheetnames:
        raise ValueError(f"Không tìm thấy sheet {SHEET_NAME!r}")
    ws_formula = formulas[SHEET_NAME]
    ws_cached = cached[SHEET_NAME]

    raw_columns = []
    for letter in COLUMN_LETTERS:
        raw_columns.extend([f"{letter}_raw", f"{letter}_cached"])
    insert_columns = [
        "source_sheet", "excel_row", "snapshot_date", "source_workbook",
        *raw_columns, "value_kinds",
    ]
    placeholders = ", ".join(["%s"] * len(insert_columns))
    staging_sql = (
        f"INSERT INTO staging.diem_nghiep_2026_w34_raw "
        f"({', '.join(insert_columns)}) VALUES ({placeholders}) "
        "ON CONFLICT (source_sheet, excel_row) DO UPDATE SET "
        + ", ".join(
            f"{column} = EXCLUDED.{column}"
            for column in insert_columns
            if column not in {"source_sheet", "excel_row"}
        )
        + ", imported_at = CURRENT_TIMESTAMP"
    )

    with psycopg.connect(**connection_kwargs(), row_factory=hybrid_row) as conn:
        with conn.cursor() as cur:
            target_count = 0
            cur.execute("INSERT INTO qd5277.DM_KhoangThoiGian(Ma_ThoiGian,Nam,Thang) VALUES(%s,2026,8) ON CONFLICT(Ma_ThoiGian) DO NOTHING", (TIME_CODE,))
            cur.execute("SET LOCAL search_path=app,qd5277,staging,public")
            for row in DATA_ROWS:
                raw_values: list[Any] = []
                kinds: dict[str, dict[str, str]] = {}
                for col in range(1, 29):
                    letter = get_column_letter(col)
                    raw = ws_formula.cell(row, col).value
                    value = ws_cached.cell(row, col).value
                    raw_values.extend([serialize_value(raw), serialize_value(value)])
                    kinds[letter] = {"raw": value_kind(raw), "cached": value_kind(value)}
                cur.execute(
                    staging_sql,
                    [
                        SHEET_NAME,
                        row,
                        SNAPSHOT_DATE,
                        workbook_path.name,
                        *raw_values,
                        Jsonb(kinds),
                    ],
                )

                # Preserve C/F in staging; standard totals come from D+E / G+H.
                observation = {}
                for suffix, area_col, harvest_col, price_col in (("land",4,7,20),("tarp",5,8,21)):
                    for name, col in (("area_",area_col),("harvest_",harvest_col)):
                        value = ws_cached.cell(row,col).value
                        raw = ws_formula.cell(row,col).value
                        if value in (None, "", "-") and not (isinstance(raw,str) and raw.startswith("=")):
                            value = 0
                        observation[name + suffix] = numeric_cached(value,row=row,column=get_column_letter(col))
                    observation["price_" + suffix] = ws_cached.cell(row,price_col).value
                def execute(query, params):
                    return cur.execute(query.replace("?", "%s"), params)
                sync_methods(execute,lookup_admin_code(CompatConnection(conn),ws_cached.cell(row,2).value),TIME_CODE,observation,postgres=True)
                target_count += len(list(method_values(observation)))
        conn.commit()

    formulas.close()
    cached.close()
    return len(DATA_ROWS), target_count


def main() -> int:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        type=Path,
        default=project_root / "data" / "CCPTNT Diemnghiep.xlsx",
    )
    args = parser.parse_args()
    load_dotenv(project_root / ".env")
    if not args.workbook.is_file():
        parser.error(f"Không tìm thấy workbook: {args.workbook}")
    staging_count, target_count = import_workbook(args.workbook)
    print(json.dumps({"staging_rows": staging_count, "qd5277_rows": target_count}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())

