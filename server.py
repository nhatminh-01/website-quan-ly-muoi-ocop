#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Website nội bộ quản lý số liệu sản xuất muối
- Không cần Flask/Django; dữ liệu được lưu tập trung trên PostgreSQL.
- openpyxl là tùy chọn để xuất Excel .xlsx.

Chạy: python server.py
Mặc định: http://0.0.0.0:8080
"""

from __future__ import annotations

import csv
import argparse
import hashlib
import html
import io
import json
import os
import re
import secrets
import sys
import threading
import time
import logging
from backend_db import compat_connect, load_settings, INTEGRITY_ERRORS
from permissions import (ROLE_ADMIN, ROLE_STAFF, ROLE_UNIT, ROLE_LABELS,
                         is_admin, is_chi_cuc_user, can_manage_users, can_review_records)
from salt_normalization import sync_methods
from weekly_import import (
    WeeklyImportError, detect_weekly_period, parse_weekly_workbook,
    commit_weekly_preview,
)
from workbook_utils import WorkbookInspectionError, inspect_workbook
import repositories
import admin_units
import admin_unit_pages
import user_profiles
import ocop_import
import ocop_import_pages
import ocop_manual
import ocop_registry
import ocop_pages
import ocop_services
import dashboard_services
from datetime import datetime, date
from http import cookies
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, quote, urlencode, urlparse
from email.parser import BytesParser
from email.policy import default

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASSET_FILES = {
    "/assets/login-hero.png": ("login-hero.png", "image/png"),
    "/assets/login.css": ("login.css", "text/css; charset=utf-8"),
    "/assets/quoc-huy.png": ("quoc-huy.png", "image/png"),
    "/assets/LOGO-CCPTNT-TP.HCM_.jpg": ("LOGO-CCPTNT-TP.HCM_.jpg", "image/jpeg"),
    "/assets/app.css": ("app.css", "text/css; charset=utf-8"),
    "/assets/app.js": ("app.js", "text/javascript; charset=utf-8"),
    "/assets/ocop-manual.js": ("ocop-manual.js", "text/javascript; charset=utf-8"),
}
HOST = os.environ.get("SALT_WEB_HOST", "0.0.0.0")
PORT = int(os.environ.get("SALT_WEB_PORT", "8080"))

STATUS_DRAFT = "draft"
STATUS_SUBMITTED = "submitted"
STATUS_APPROVED = "approved"
STATUS_RETURNED = "returned"

STATUS_LABELS = {
    STATUS_DRAFT: "Bản nháp",
    STATUS_SUBMITTED: "Đã gửi Chi cục",
    STATUS_APPROVED: "Đã duyệt",
    STATUS_RETURNED: "Yêu cầu chỉnh sửa",
}

SESSIONS: dict[str, dict] = {}
SESSION_LOCK = threading.Lock()

NUMERIC_FIELDS = [
    "area_land", "area_tarp",
    "harvest_land", "harvest_tarp",
    "sold_land", "sold_tarp",
    "remaining_land", "remaining_tarp",
    "processed_fine", "processed_iodized",
    "households", "workers",
    "damage_land", "damage_tarp",
]

FIELD_LABELS = {
    "area_land": "Diện tích muối đất (ha)",
    "area_tarp": "Diện tích muối trải bạt (ha)",
    "harvest_land": "Thu hoạch muối đất (tấn)",
    "harvest_tarp": "Thu hoạch muối trải bạt (tấn)",
    "sold_land": "Tiêu thụ muối đất (tấn)",
    "sold_tarp": "Tiêu thụ muối trải bạt (tấn)",
    "remaining_land": "Còn lại muối đất (tấn)",
    "remaining_tarp": "Còn lại muối trải bạt (tấn)",
    "processed_fine": "Muối tinh chế biến (tấn)",
    "processed_iodized": "Muối I-ốt chế biến (tấn)",
    "households": "Số hộ làm muối (hộ)",
    "workers": "Số lao động làm muối (người)",
    "price_land": "Giá bán muối đất (đồng/kg)",
    "price_tarp": "Giá bán muối trải bạt (đồng/kg)",
    "damage_land": "Thiệt hại muối đất (tấn)",
    "damage_tarp": "Thiệt hại muối trải bạt (tấn)",
}

def esc(value):
    return html.escape("" if value is None else str(value), quote=True)


def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def hash_password(password: str, salt: str | None = None) -> str:
    if salt is None:
        salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt), 160_000)
    return f"{salt}${digest.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, expected = stored.split("$", 1)
        actual = hash_password(password, salt).split("$", 1)[1]
        return secrets.compare_digest(actual, expected)
    except Exception:
        return False


def db_conn():
    return compat_connect()


def get_user(con, user_id):
    return repositories.get_user(con, user_id)


def find_user_by_username(con, username):
    return repositories.find_user_by_username(con, username)


def find_active_user(con, username):
    user = find_user_by_username(con, username)
    return user if user and user["active"] else None


def create_user(con, username, password_hash, role, unit_name, created_at):
    return repositories.create_user(con, username, password_hash, role, unit_name, created_at)


def set_user_password(con, user_id, password_hash):
    repositories.set_local_password(con, user_id, password_hash)


def update_user_account(con, user_id, username, role, unit_name, active, password_hash=None):
    repositories.update_user(con, user_id, username, role, unit_name, active, password_hash)


def delete_user_account(con, user_id):
    repositories.delete_user(con, user_id)


def ocop_available(con=None):
    """Return whether the required PostgreSQL application schema is ready."""
    own = con is None
    if own:
        con = db_conn()
    try:
        return repositories.postgres_application_ready(con)
    finally:
        if own:
            con.close()


def invalidate_user_sessions(user_id, keep_sid=None):
    with SESSION_LOCK:
        for sid in list(SESSIONS):
            if SESSIONS[sid]["user_id"] == user_id and sid != keep_sid:
                SESSIONS.pop(sid, None)


def ocop_user_links(con, user_id):
    if not ocop_available(con):
        return 0
    total = 0
    for table, column in (("ocop_entities", "created_by"), ("ocop_products", "created_by"),
                          ("ocop_applications", "created_by"), ("ocop_applications", "reviewer_id"),
                          ("ocop_reviews", "reviewer_id"), ("user_admin_units", "user_id")):
        total += con.execute(f"SELECT COUNT(*) FROM {table} WHERE {column}=?", (user_id,)).fetchone()[0]
    return total


def init_db():
    con = db_conn()
    try:
        if not repositories.postgres_application_ready(con):
            raise RuntimeError("PostgreSQL chưa đủ schema ứng dụng. Áp dụng các migration database/sql đến 012_admin_units.sql trước.")
    finally:
        con.close()


def add_audit(con, record_id, user_id, action, detail=""):
    con.execute(
        "INSERT INTO audit_logs(record_id,user_id,action,detail,created_at) VALUES(?,?,?,?,?)",
        (record_id, user_id, action, detail, now_text()),
    )


def fnum(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return 0.0


def fint(v):
    try:
        return int(float(str(v).replace(",", ".")))
    except (ValueError, TypeError):
        return 0


def _parse_vietnamese_number(value):
    """Đọc số theo cách người dùng Việt Nam thường nhập: 1.500; 1,500; 1500; 1.500,50."""
    if value is None:
        return None
    token = str(value).strip().replace(" ", "")
    if not token:
        return None

    # Có cả dấu chấm và dấu phẩy: ưu tiên kiểu Việt Nam 1.500,50.
    if "." in token and "," in token:
        if token.rfind(",") > token.rfind("."):
            token = token.replace(".", "").replace(",", ".")
        else:
            token = token.replace(",", "")
    elif "." in token:
        parts = token.split(".")
        # 1.500 hoặc 12.500.000 -> dấu chấm là phân cách hàng nghìn.
        if len(parts) > 1 and all(len(x) == 3 for x in parts[1:]):
            token = "".join(parts)
    elif "," in token:
        parts = token.split(",")
        # Giá bán thường nhập nguyên đồng/kg; 1,500 được hiểu là 1500.
        if len(parts) > 1 and all(len(x) == 3 for x in parts[1:]):
            token = "".join(parts)
        else:
            token = token.replace(",", ".")

    try:
        return float(token)
    except ValueError:
        return None


def price_average(value):
    """Chuyển ô giá dạng '1.000 - 1.500' thành giá bình quân 1250."""
    if value is None:
        return 0.0
    text = str(value).strip()
    if not text:
        return 0.0

    tokens = re.findall(r"\d[\d.,]*", text)
    numbers = []
    for token in tokens:
        n = _parse_vietnamese_number(token)
        if n is not None:
            numbers.append(n)

    if not numbers:
        return 0.0
    if len(numbers) == 1:
        return round(numbers[0], 2)
    return round((numbers[0] + numbers[1]) / 2.0, 2)


def canonical_admin_unit(unit_name, con=None, *, active_only=True):
    own = con is None
    if own:
        con = db_conn()
    try:
        return admin_units.unit_lookup(con, active_only=active_only)(unit_name)
    finally:
        if own:
            con.close()


def migrate_temporary_unit_codes(con):
    """Chuyển các bản ghi chuẩn hóa đang dùng TMP... sang mã hành chính chính thức."""
    migrated = 0
    for unit in admin_units.units(con):
        unit_name, official_code = unit["name"], unit["code"]
        old_rows = con.execute(
            """
            SELECT Ma_DonViHanhChinh
            FROM DM_DonViHanhChinh
            WHERE lower(TenDonVi)=lower(?)
              AND Ma_DonViHanhChinh LIKE 'TMP%'
            """,
            (unit_name,),
        ).fetchall()

        for old in old_rows:
            old_code = old["Ma_DonViHanhChinh"]
            standard_rows = con.execute(
                "SELECT * FROM DN_SanLuongMuoi WHERE Ma_DonViHanhChinh=?",
                (old_code,),
            ).fetchall()

            for r in standard_rows:
                con.execute(
                    """
                    INSERT INTO DN_SanLuongMuoi
                    (Ma_DonViHanhChinh, Ma_ThoiGian, PhuongPhapSX, DienTich, SanLuong, GiaBanBinhQuan)
                    VALUES(?,?,?,?,?,?)
                    ON CONFLICT(Ma_DonViHanhChinh, Ma_ThoiGian, PhuongPhapSX) DO UPDATE SET
                        DienTich=excluded.DienTich,
                        SanLuong=excluded.SanLuong,
                        GiaBanBinhQuan=excluded.GiaBanBinhQuan
                    """,
                    (official_code, r["Ma_ThoiGian"], r["PhuongPhapSX"],
                     r["DienTich"], r["SanLuong"], r["GiaBanBinhQuan"]),
                )
                migrated += 1

            con.execute(
                "DELETE FROM DN_SanLuongMuoi WHERE Ma_DonViHanhChinh=?",
                (old_code,),
            )
            con.execute(
                "DELETE FROM DM_DonViHanhChinh WHERE Ma_DonViHanhChinh=?",
                (old_code,),
            )

    return migrated


def ensure_standard_unit(con, unit_name):
    """Resolve the DB code, including historical names; never create a TMP unit."""
    unit = canonical_admin_unit(unit_name, con, active_only=False)
    if not unit:
        raise ValueError(f"Đơn vị chưa có mã hành chính: {unit_name}")
    return unit[1]


def ensure_standard_time(con, report_date):
    """Tạo mã kỳ tháng YYYY-MM cho dữ liệu tổng hợp chuẩn."""
    d = date.fromisoformat(str(report_date))
    code = d.strftime("%Y-%m")
    con.execute(
        """
        INSERT INTO DM_KhoangThoiGian(Ma_ThoiGian, Nam, Thang, VuMua)
        VALUES(?, ?, ?, NULL)
        ON CONFLICT(Ma_ThoiGian) DO NOTHING
        """,
        (code, d.year, d.month),
    )
    return code


def sync_standard_salt_record(con, record):
    """Đồng bộ báo cáo lũy tiến mới nhất thành một dòng muối truyền thống/tháng."""
    if not record:
        return

    unit_code = ensure_standard_unit(con, record["unit_name"])
    time_code = ensure_standard_time(con, record["report_date"])
    month_start = date.fromisoformat(str(record["report_date"])).replace(day=1)
    next_month = date(month_start.year + (month_start.month == 12), (month_start.month % 12) + 1, 1)
    latest = con.execute(
        """SELECT * FROM records
           WHERE unit_name=? AND status='approved'
             AND report_date>=? AND report_date<? AND reporting_mode='cumulative'""" +
        " ORDER BY report_date DESC, id DESC LIMIT 1",
        (record["unit_name"], month_start.isoformat(), next_month.isoformat()),
    ).fetchone()

    sync_methods(con.execute, unit_code, time_code, latest, postgres=True)


def sync_all_approved_records():
    """Có thể gọi thủ công khi cần đồng bộ lại toàn bộ dữ liệu đã duyệt."""
    con = db_conn()
    try:
        rows = con.execute(
            "SELECT * FROM records WHERE status=? ORDER BY report_date, unit_name",
            (STATUS_APPROVED,),
        ).fetchall()
        for row in rows:
            sync_standard_salt_record(con, row)
        con.commit()
        return len(rows)
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def fmt_num(v, decimals=1):
    try:
        n = float(v or 0)
    except Exception:
        n = 0
    if abs(n - round(n)) < 1e-10:
        return f"{int(round(n)):,}".replace(",", ".")
    s = f"{n:,.{decimals}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def record_totals(r):
    area_total = (r["area_land"] or 0) + (r["area_tarp"] or 0)
    harvest_total = (r["harvest_land"] or 0) + (r["harvest_tarp"] or 0)
    sold_total = (r["sold_land"] or 0) + (r["sold_tarp"] or 0)
    remaining_total = (r["remaining_land"] or 0) + (r["remaining_tarp"] or 0)
    processed_total = (r["processed_fine"] or 0) + (r["processed_iodized"] or 0)
    damage_total = (r["damage_land"] or 0) + (r["damage_tarp"] or 0)
    avg_yield = harvest_total / area_total if area_total else 0
    return {
        "area_total": area_total,
        "harvest_total": harvest_total,
        "sold_total": sold_total,
        "remaining_total": remaining_total,
        "processed_total": processed_total,
        "damage_total": damage_total,
        "avg_yield": avg_yield,
    }


def parse_body(handler):
    length = int(handler.headers.get("Content-Length", "0") or 0)
    raw = handler.rfile.read(length).decode("utf-8", errors="replace")
    parsed = parse_qs(raw, keep_blank_values=True)
    return {k: v[-1] if v else "" for k, v in parsed.items()}

def parse_multipart_form(handler):
    content_type = handler.headers.get("Content-Type", "")
    length = int(handler.headers.get("Content-Length", "0") or 0)

    raw_body = handler.rfile.read(length)

    raw_message = (
        f"Content-Type: {content_type}\r\n"
        f"MIME-Version: 1.0\r\n\r\n"
    ).encode("utf-8") + raw_body

    message = BytesParser(policy=default).parsebytes(raw_message)

    fields = {}
    files = {}

    if not message.is_multipart():
        return fields, files

    for part in message.iter_parts():
        name = part.get_param(
            "name",
            header="Content-Disposition"
        )

        if not name:
            continue

        payload = part.get_payload(decode=True) or b""
        filename = part.get_filename()

        if filename:
            files[name] = {
                "filename": filename,
                "content": payload,
                "content_type": part.get_content_type()
            }

        else:
            charset = part.get_content_charset() or "utf-8"

            fields[name] = payload.decode(
                charset,
                errors="replace"
            )

    return fields, files


def excel_text(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()

def new_session(user_row):
    sid = secrets.token_urlsafe(32)
    csrf = secrets.token_urlsafe(24)
    with SESSION_LOCK:
        SESSIONS[sid] = {
            "user_id": user_row["id"],
            "username": user_row["username"],
            "role": user_row["role"],
            "unit_name": user_row["unit_name"],
            "csrf": csrf,
            "created": time.time(),
            "auth_version": user_row["password_hash"],
            "flash": None,
        }
    return sid


def get_session(handler):
    raw = handler.headers.get("Cookie", "")
    jar = cookies.SimpleCookie()
    try:
        jar.load(raw)
    except Exception:
        return None, None
    morsel = jar.get("salt_session")
    if not morsel:
        return None, None
    sid = morsel.value
    with SESSION_LOCK:
        s = SESSIONS.get(sid)
    if s:
        con = db_conn()
        try:
            current = get_user(con, s["user_id"])
        finally:
            con.close()
        if (not current or not current["active"] or current["role"] not in tuple(ROLE_LABELS)
                or s.get("auth_version") != current["password_hash"]):
            with SESSION_LOCK:
                SESSIONS.pop(sid, None)
            return None, None
        # Authorization reflects account changes on the next request.
        s.update(username=current["username"], role=current["role"], unit_name=current["unit_name"])
    return sid, s


def set_flash(session, kind, text):
    session["flash"] = (kind, text)


def take_flash(session):
    if not session:
        return ""
    item = session.get("flash")
    session["flash"] = None
    if not item:
        return ""
    kind, text = item
    return f'<div class="notice {esc(kind)}">{esc(text)}</div>'


def csrf_input(session):
    return f'<input type="hidden" name="csrf" value="{esc(session["csrf"])}">'


def check_csrf(session, data):
    return bool(session and secrets.compare_digest(str(data.get("csrf", "")), str(session.get("csrf", ""))))


def account_display_name(session):
    """Tên trên giao diện; không thay đổi tên đơn vị dùng để lọc báo cáo."""
    username = str(session.get("username") or "").strip()
    account_key = username.casefold().replace("_", "").replace("-", "").replace(" ", "")
    account_names = {
        "chicuc": "CHI CỤC PHÁT TRIỂN NÔNG THÔN THÀNH PHỐ HỒ CHÍ MINH",
        "cangio": "XÃ CẦN GIỜ",
    }
    unit_name = str(session.get("unit_name") or "").strip()
    return account_names.get(account_key, (unit_name or username).upper())


def icon(name):
    paths = {
        "dashboard": '<rect x="3" y="3" width="7" height="7" rx="1.5"/><rect x="14" y="3" width="7" height="7" rx="1.5"/><rect x="3" y="14" width="7" height="7" rx="1.5"/><rect x="14" y="14" width="7" height="7" rx="1.5"/>',
        "table": '<rect x="3" y="4" width="18" height="16" rx="2"/><path d="M3 10h18M3 15h18M9 4v16"/>',
        "plus": '<rect x="4" y="3" width="16" height="18" rx="2"/><path d="M12 8v8M8 12h8"/>',
        "users": '<circle cx="9" cy="8" r="3"/><path d="M3 21v-2a6 6 0 0 1 12 0v2M16 5a3 3 0 0 1 0 6M21 21v-2a6 6 0 0 0-4-5"/>',
        "key": '<circle cx="8" cy="8" r="5"/><path d="m12 12 9 9m-3-3 3-3m-6 0 3-3"/>',
        "search": '<circle cx="10.5" cy="10.5" r="6.5"/><path d="m16 16 5 5"/>',
        "menu": '<path d="M4 6h16M4 12h16M4 18h16"/>',
        "logout": '<path d="M9 4H4v16h5M9 12h12m-4-4 4 4-4 4"/>',
        "file": '<path d="M14 2H5v20h14V7zM14 2v6h5M8 12h8M8 16h8"/>',
        "send": '<path d="m22 2-7 20-4-9-9-4 20-7ZM22 2 11 13"/>',
        "check": '<rect x="4" y="3" width="16" height="18" rx="2"/><path d="m8 12 3 3 6-6"/>',
        "return": '<path d="m8 4-5 5 5 5M3 9h11a7 7 0 0 1 0 14"/>',
        "chart": '<path d="M3 3v18h18M7 15l4-5 4 2 5-7M16 5h4v4"/>',
        "area": '<path d="m12 3 10 5-10 5L2 8 12 3ZM2 12l10 5 10-5M2 16l10 5 10-5"/>',
        "harvest": '<path d="M4 20V9h16v11zM4 9l8-6 8 6M8 13h8M8 17h8"/>',
        "sold": '<path d="M3 5h3l3 12h10l2-9H7"/><circle cx="10" cy="21" r="1"/><circle cx="18" cy="21" r="1"/>',
        "stock": '<path d="m12 3 9 5v10l-9 5-9-5V8l9-5ZM3 8l9 5 9-5M12 13v10M8 5l9 5"/>',
        "home": '<path d="m3 10 9-7 9 7v11H3zM9 21v-8h6v8"/>',
        "download": '<path d="M12 3v12m-5-5 5 5 5-5M4 15v6h16v-6"/>',
        "location": '<path d="M20 10c0 6-8 12-8 12S4 16 4 10a8 8 0 0 1 16 0Z"/><circle cx="12" cy="10" r="2.5"/>',
        "alert": '<path d="m12 3 9 17H3L12 3Z"/><path d="M12 9v4m0 3h.01"/>',
    }
    return f'<svg class="icon" viewBox="0 0 24 24" aria-hidden="true" focusable="false">{paths.get(name, paths["file"])}</svg>'


def sidebar_group(group, items, current):
    links = []
    for href, symbol, label in items:
        active = href == current
        links.append(f'<a class="sidebar-link{" active" if active else ""}" href="{href}" title="{label}"'
                     + (' aria-current="page"' if active else '')
                     + f'>{icon(symbol)}<span class="sidebar-label">{label}</span></a>')
    content = ''.join(links)
    if group not in ("DIÊM NGHIỆP", "OCOP"):
        return f'<div class="sidebar-section">{group}</div>{content}'
    key = "ocop" if group == "OCOP" else "salt"
    active = any(href == current for href, _, _ in items)
    return (f'<details class="sidebar-group" data-group="{key}" data-active="{str(active).lower()}"'
            + (' open' if active else '') + f'><summary title="{group}"><span class="group-label">{group}</span>'
            + '<svg class="group-chevron icon" viewBox="0 0 24 24" aria-hidden="true"><path d="m6 9 6 6 6-6"/></svg>'
            + f'</summary><div class="sidebar-group-links">{content}</div></details>')


def base_page(title, body, session=None, active_path=None):
    top = ""
    if session:
        current = {
            "Tổng quan": "/dashboard",
            "Trang chủ": "/dashboard",
            "Tài khoản": "/users",
            "Sửa tài khoản": "/users",
            "Đổi mật khẩu": "/change-password",
            "Nhập số liệu": "/records/new",
            "Dữ liệu báo cáo": "/records",
            "Chi tiết báo cáo": "/records",
            "Yêu cầu chỉnh sửa": "/records",
            "Import Excel": "/records/new",
            "Import báo cáo tuần": "/import-excel",
            "Xem trước import tuần": "/import-excel",
            "Tra cứu báo cáo tuần": "/salt/weekly",
            "Dữ liệu chuẩn hóa": "/records",
        }.get(title, "/records")
        current = active_path or ("/standard-data" if title == "Dữ liệu chuẩn hóa" else current)
        salt_items = [("/records", "table", "Tra cứu báo cáo Diêm nghiệp")]
        if is_chi_cuc_user(session):
            salt_items.insert(0, ("/import-excel", "download", "Import báo cáo tuần"))
        system_items = []
        if can_manage_users(session):
            system_items.append(("/users", "users", "Tài khoản"))
            system_items.append(("/admin-units", "location", "Danh mục đơn vị hành chính"))
        system_items.extend([("/change-password", "key", "Đổi mật khẩu"),
                             ("/logout", "logout", "Đăng xuất")])
        # Keep the dashboard as a standalone home link so it is always easy
        # to reach without adding it to either business group.
        nav_groups = [("DIÊM NGHIỆP", salt_items)]
        if ocop_available():
            ocop_items = [("/ocop", "table", "Tra cứu / Xuất báo cáo"),
                          ("/ocop/expiry-alerts", "alert", "Cảnh báo hết hạn")]
            if is_chi_cuc_user(session):
                ocop_items.append(("/ocop/import", "download", "Import dữ liệu OCOP"))
                ocop_items.append(("/ocop/manual", "edit", "Nhập dữ liệu trực tiếp"))
            nav_groups.append(("OCOP", ocop_items))
        nav_groups.append(("HỆ THỐNG", system_items))
        home_active = current == "/dashboard"
        home_link = (f'<a class="sidebar-link sidebar-home-link{" active" if home_active else ""}" '
                     'href="/dashboard" title="Bảng giám sát"'
                     + (' aria-current="page"' if home_active else '')
                     + f'>{icon("home")}<span class="sidebar-label">Bảng giám sát</span></a>')
        nav = [home_link] + [sidebar_group(group, items, current) for group, items in nav_groups]
        name = account_display_name(session)
        display_label = esc(name)
        if name == "CHI CỤC PHÁT TRIỂN NÔNG THÔN THÀNH PHỐ HỒ CHÍ MINH":
            display_label = 'CHI CỤC PHÁT TRIỂN NÔNG THÔN<br>THÀNH PHỐ HỒ CHÍ MINH'
        role_label = ROLE_LABELS[session["role"]]
        initial = esc((session['username'] or 'U')[0].upper())
        account_copy = f'<div class="account-copy"><div class="account-name">{esc(session["username"])}</div><div class="account-role">{role_label}</div></div>'
        avatar = f'<span class="account-avatar" aria-hidden="true">{initial}</span>'
        if is_chi_cuc_user(session):
            account_html = (f'<div class="header-account profile-account"><a href="/profile" class="user-profile-link" title="Xem thông tin cá nhân" aria-label="Xem thông tin cá nhân">{account_copy}{avatar}</a>'
                            '<a class="account-logout" href="/logout">Đăng xuất</a></div>')
        else:
            account_html = f'<div class="header-account"><div class="account-copy"><div class="account-name">{esc(session["username"])}</div><div class="account-role">{role_label}</div><a class="account-logout" href="/logout">Đăng xuất</a></div>{avatar}</div>'
        top = f"""
        <a class="skip-link" href="#main-content">Đến nội dung chính</a>
        <header class="site-header">
          <button class="menu-toggle" id="sidebar-toggle" type="button" aria-label="Thu gọn menu" aria-expanded="true" aria-controls="site-sidebar">{icon('menu')}</button>
          <a class="brand-link" href="/dashboard">
            <img class="brand-emblem" src="/assets/LOGO-CCPTNT-TP.HCM_.jpg" alt="Logo Chi cục Phát triển nông thôn Thành phố Hồ Chí Minh" width="56" height="58">
            <div class="brand-copy"><div class="brand-agency">{display_label}</div><div class="brand-subtitle">HỆ THỐNG QUẢN LÝ NGHIỆP VỤ</div></div>
          </a>
          <div class="header-tools">
            <div class="header-clock"><time class="clock-time" id="clock-time">{datetime.now().strftime('%H:%M:%S')}</time><div class="clock-date" id="clock-date">{date.today().strftime('%d/%m/%Y')}</div></div>
            {account_html}
          </div>
        </header>
        <aside class="sidebar" id="site-sidebar" aria-label="Menu chính">
          <nav class="sidebar-nav" aria-label="Chức năng">{''.join(nav)}</nav>
          <div class="sidebar-footer"><img class="sidebar-watermark" src="/assets/LOGO-CCPTNT-TP.HCM_.jpg" alt="" width="148" height="152"><div class="sidebar-footer-line"></div><strong><span>TRUNG TÂM CHUYỂN ĐỔI SỐ</span><span>NÔNG NGHIỆP VÀ MÔI TRƯỜNG</span></strong><p>Theo dõi sản xuất và tổng hợp báo cáo các đơn vị.</p></div>
        </aside>
        <button type="button" id="sidebar-backdrop" class="sidebar-backdrop" aria-label="Đóng menu" tabindex="-1"></button>
        """
        body = f'<main class="app-main" id="main-content">{body}</main>'
    return f"""<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(title)} · Quản lý nghiệp vụ</title><link rel="icon" href="/assets/LOGO-CCPTNT-TP.HCM_.jpg" type="image/jpeg"><link rel="stylesheet" href="/assets/app.css?v=20260917-dashboard-v1"><script src="/assets/app.js?v=20260916-sidebar-scroll-v2" defer></script></head><body>{top}{body}</body></html>"""


# OCOP T2 is part of this single server entry point.  Keep the original page
# renderer as the base and extend it with the import/history navigation.
_base_page_core = base_page
_ocop_render_core = ocop_pages.render


def _t2_base_page(title, body, session=None, active_path=None):
    # Legacy T2 routes still render through the single base page. Their links
    # are intentionally not injected into the product navigation anymore.
    return _base_page_core(title, body, session, active_path=active_path)


def _t2_ocop_render(path, query, session, con, helpers):
    result = _ocop_render_core(path, query, session, con, helpers)
    clean = str(path or "").rstrip("/")
    if result and clean.startswith("/ocop/products/") and not clean.endswith("/edit"):
        parts = clean.strip("/").split("/")
        if len(parts) == 3 and parts[2].isdigit():
            section = ocop_import_pages.product_history_section(con, session, int(parts[2]))
            if section:
                title, body = result
                pos = body.rfind("</div>")
                body = body[:pos] + section + body[pos:] if pos >= 0 else body + section
                result = (title, body)
    return result


base_page = _t2_base_page
ocop_pages.render = _t2_ocop_render


def login_page(message=""):
    notice = f'<div class="notice err">{esc(message)}</div>' if message else ""
    body = f"""
    <div class="login-wrap"><div class="login">
      <img class="login-emblem" src="/assets/LOGO-CCPTNT-TP.HCM_.jpg" alt="Logo Chi cục Phát triển nông thôn Thành phố Hồ Chí Minh" width="70" height="72">
      <h1>Hệ thống quản lý nghiệp vụ</h1>
      <p>Chi cục Phát triển nông thôn và các đơn vị xã/phường.</p>
      {notice}
      <form method="post" action="/login">
        <div class="field"><label>Tên đăng nhập</label><input name="username" autocomplete="username" required autofocus></div>
        <div class="field"><label>Mật khẩu</label><input type="password" name="password" autocomplete="current-password" required></div>
        <button class="btn primary" style="width:100%;margin-top:6px">Đăng nhập</button>
      </form>
      <div class="notice info" style="margin-top:16px"><b>Tài khoản demo:</b> chicuc / 123456<br>Đơn vị ví dụ: an_thoi_dong / 123456<br><b>Đổi mật khẩu ngay khi triển khai thật.</b></div>
    </div></div>
    """
    return base_page("Đăng nhập", body)


def filters_from_query(query, session):
    params = parse_qs(query)
    filters = {
        "q": params.get("q", [""])[0].strip(),
        "unit": params.get("unit", [""])[0].strip(),
        "status": params.get("status", [""])[0].strip(),
        "from": params.get("from", [""])[0].strip(),
        "to": params.get("to", [""])[0].strip(),
    }
    if session["role"] == ROLE_UNIT:
        filters["unit"] = session["unit_name"]
    return filters


def fetch_records(con, session, filters):
    sql = "SELECT r.*, u.username creator_username FROM records r JOIN users u ON u.id=r.created_by WHERE 1=1"
    args = []
    if session["role"] == ROLE_UNIT:
        names = admin_units.report_names(con, session["unit_name"])
        sql += " AND r.unit_name IN (" + ",".join("?" for _ in names) + ")"
        args.extend(names)
    elif filters.get("unit"):
        names = admin_units.report_names(con, filters["unit"])
        sql += " AND r.unit_name IN (" + ",".join("?" for _ in names) + ")"
        args.extend(names)
    if filters.get("status"):
        sql += " AND r.status=?"
        args.append(filters["status"])
    if filters.get("from"):
        sql += " AND r.report_date>=?"
        args.append(filters["from"])
    if filters.get("to"):
        sql += " AND r.report_date<=?"
        args.append(filters["to"])
    if filters.get("q"):
        sql += " AND (r.unit_name LIKE ? OR r.note LIKE ? OR r.reviewer_note LIKE ?)"
        q = f"%{filters['q']}%"
        args.extend([q, q, q])
    sql += " ORDER BY r.report_date DESC, r.unit_name ASC"
    return con.execute(sql, args).fetchall()


def get_units(con):
    return [r["name"] for r in admin_units.units(con, active_only=True, communes_only=True)]


def _dashboard_value(value):
    return "—" if value is None else fmt_num(value)


def _dashboard_change(change, unit):
    if not change:
        return ""
    delta = change.get("delta")
    percent = change.get("percent_change")
    if percent is not None:
        arrow = "▲" if delta > 0 else "▼" if delta < 0 else "•"
        text = f"{arrow} {fmt_num(abs(percent), 1)}% so với kỳ trước"
    elif delta is not None:
        arrow = "▲" if delta > 0 else "▼" if delta < 0 else "•"
        text = f"{arrow} {fmt_num(abs(delta), 1)} {unit} so với kỳ trước"
    else:
        return ""
    return f'<div class="dashboard-card-change">{esc(text)}</div>'


def _dashboard_metric_card(title, metric, unit, symbol, *, change=None,
                           breakdown=False, warning=False):
    tone = " warning" if warning else ""
    parts = [
        f'<article class="dashboard-stat-card{tone}">'
        f'<div class="dashboard-stat-heading"><h3>{esc(title)}</h3>{icon(symbol)}</div>'
        f'<div class="dashboard-stat-value">{_dashboard_value(metric.get("total"))} <small>{esc(unit)}</small></div>'
        f'{_dashboard_change(change, unit)}'
    ]
    if breakdown:
        parts.append(
            '<div class="dashboard-stat-breakdown">'
            f'<span><b>Muối đất</b><strong>{_dashboard_value(metric.get("land"))} <small>{esc(unit)}</small></strong></span>'
            f'<span><b>Muối trải bạt</b><strong>{_dashboard_value(metric.get("tarp"))} <small>{esc(unit)}</small></strong></span>'
            '</div>'
        )
    parts.append('</article>')
    return ''.join(parts)


def _dashboard_summary_card(title, value, symbol, *, href="", warning=False):
    tag = "a" if href else "article"
    tone = " warning" if warning else ""
    link = f' href="{esc(href)}"' if href else ""
    return (
        f'<{tag} class="dashboard-summary-card{tone}{" dashboard-summary-link" if href else ""}"{link}>'
        f'<div class="dashboard-summary-heading"><h3>{esc(title)}</h3>{icon(symbol)}</div>'
        f'<strong>{_dashboard_value(value)}</strong>'
        f'{"<span>Xem danh sách cảnh báo →</span>" if href else ""}'
        f'</{tag}>'
    )


def landing_page(session, query=""):
    """Render the common monitoring dashboard from the Plan 04 data model."""
    params = parse_qs(query)
    selected_unit = params.get("unit", [""])[0].strip()
    selected_week = params.get("week", [""])[0].strip()
    salt_filters = {"unit": selected_unit, "week": selected_week}
    ocop_filters = {"unit": selected_unit}
    con = db_conn()
    try:
        dashboard = dashboard_services.get_dashboard_data(
            con, session, salt_filters=salt_filters, ocop_filters=ocop_filters
        )
        unit_options = dashboard_services.get_dashboard_unit_options(con, session)
    finally:
        con.close()

    salt = dashboard["salt"]
    ocop = dashboard["ocop"]
    if is_chi_cuc_user(session):
        unit_options_html = '<option value="">Toàn thành phố</option>'
        unit_options_html += ''.join(
            f'<option value="{esc(row["code"])}" {"selected" if row["code"] == selected_unit else ""}>'
            f'{esc(row["name"])}</option>'
            for row in unit_options
        )
    else:
        unit_options_html = ''.join(
            f'<option value="{esc(row["code"])}" selected>{esc(row["name"])}</option>'
            for row in unit_options
        )

    if salt:
        selected_period = salt["selected"]
        week_options = ''.join(
            f'<option value="{esc(period["week_code"])}" {"selected" if period["week_code"] == selected_period["week_code"] else ""}>'
            f'{esc(period["week_code"])}</option>'
            for period in salt["periods"]
        )
        current_metrics = salt["current_metrics"]
        changes = salt["changes"]
        salt_period_label = (
            f'{esc(selected_period["week_code"])} · '
            f'{date.fromisoformat(str(selected_period["report_date"])).strftime("%d/%m/%Y")}'
        )
        salt_cards = ''.join([
            _dashboard_metric_card(
                "DIỆN TÍCH SẢN XUẤT", current_metrics["area"], "ha", "area",
                change=changes["area"]["total"], breakdown=True,
            ),
            _dashboard_metric_card(
                "SẢN LƯỢNG THU HOẠCH", current_metrics["harvest"], "tấn", "harvest",
                change=changes["harvest"]["total"], breakdown=True,
            ),
            _dashboard_metric_card(
                "ĐÃ TIÊU THỤ", current_metrics["consumption"], "tấn", "sold",
                change=changes["consumption"]["total"],
            ),
            _dashboard_metric_card(
                "CÒN LẠI", current_metrics["remaining"], "tấn", "stock",
                change=changes["remaining"]["total"],
            ),
        ])
        salt_people = (
            '<div class="dashboard-people">'
            f'<div><span>Số hộ làm muối</span><strong>{_dashboard_value(current_metrics["households"]["total"])}</strong><small>hộ</small></div>'
            f'<div><span>Số lao động</span><strong>{_dashboard_value(current_metrics["workers"]["total"])}</strong><small>người</small></div>'
            '</div>'
        )
    else:
        week_options = '<option value="">Chưa có kỳ dữ liệu</option>'
        salt_period_label = "Chưa có dữ liệu"
        salt_cards = '<div class="dashboard-empty">Chưa có dữ liệu Diêm nghiệp trong phạm vi đang xem.</div>'
        salt_people = ''

    detail_params = {"week": selected_period["week_code"]} if salt else {}
    if selected_unit:
        detail_params["unit"] = selected_unit
    salt_detail_href = "/salt/weekly" + ("?" + urlencode(detail_params) if detail_params else "")
    ocop_updated = date.today().strftime("%d/%m/%Y")
    ocop_cards = ''.join([
        _dashboard_summary_card("SẢN PHẨM OCOP", ocop["managed_products"], "table"),
        _dashboard_summary_card("CHỦ THỂ OCOP", ocop["managed_entities"], "users"),
        _dashboard_summary_card("CÒN HIỆU LỰC", ocop["valid_products"], "check"),
        _dashboard_summary_card(
            "SẮP HẾT HẠN ≤ 3 THÁNG", ocop["expiring_products"], "alert",
            href="/ocop/expiry-alerts", warning=True,
        ),
    ])
    ocop_meta = (
        '<div class="dashboard-ocop-meta">'
        '<div class="dashboard-stars" aria-label="Phân bố hạng sao">'
        f'<span><b>3 sao</b><strong>{_dashboard_value(ocop["star_3"])}</strong></span>'
        f'<span><b>4 sao</b><strong>{_dashboard_value(ocop["star_4"])}</strong></span>'
        f'<span><b>5 sao</b><strong>{_dashboard_value(ocop["star_5"])}</strong></span>'
        '</div>'
        f'<span class="dashboard-badge danger">Đã hết hạn: {_dashboard_value(ocop["expired_products"])}</span>'
        + (
            f'<span class="dashboard-badge muted">Chưa có ngày hết hạn: {_dashboard_value(ocop["missing_expiry"])}</span>'
            if ocop["missing_expiry"] else ""
        )
        + '</div>'
    )

    body = f"""
    <div class="container dashboard-home">
      {take_flash(session)}
      <div class="page-head"><div><h1>Bảng giám sát</h1><div class="subtitle">Tổng hợp nhanh tình hình Diêm nghiệp và OCOP theo phạm vi được phép xem.</div></div></div>
      <form class="dashboard-filter-card" method="get" action="/dashboard">
        <div class="field"><label for="dashboard-unit">Đơn vị</label><select id="dashboard-unit" name="unit">{unit_options_html}</select></div>
        <div class="field"><label for="dashboard-week">Kỳ Diêm nghiệp</label><select id="dashboard-week" name="week">{week_options}</select></div>
        <button class="btn primary" type="submit">Áp dụng</button><a class="btn" href="/dashboard">Đặt lại bộ lọc</a>
      </form>
      <div class="dashboard-module-grid">
        <section class="dashboard-section" aria-labelledby="salt-dashboard-title">
          <div class="dashboard-section-heading"><div><h2 id="salt-dashboard-title">DIÊM NGHIỆP</h2></div><div class="dashboard-section-meta"><strong>{salt_period_label}</strong></div></div>
          <div class="dashboard-stat-grid">{salt_cards}</div>
          {salt_people}
          <div class="dashboard-section-footer"><a class="btn small" href="{esc(salt_detail_href)}">Xem chi tiết Diêm nghiệp →</a></div>
        </section>
        <section class="dashboard-section" aria-labelledby="ocop-dashboard-title">
          <div class="dashboard-section-heading"><div><h2 id="ocop-dashboard-title">OCOP</h2></div><div class="dashboard-section-meta">Tính đến ngày<br><strong>{ocop_updated}</strong></div></div>
          <div class="dashboard-summary-grid">{ocop_cards}</div>
          {ocop_meta}
          <div class="dashboard-section-footer"><a class="btn small" href="/ocop">Xem chi tiết OCOP →</a></div>
        </section>
      </div>
    </div>"""
    return base_page("Tổng quan", body, session)


def dashboard_page(session, query=""):
    """Legacy salt dashboard kept for compatibility; not linked from UI."""
    """Dashboard for effective weekly records, compared with a distinct earlier week."""
    params = parse_qs(query)
    con = db_conn()
    try:
        data = dashboard_services.get_salt_dashboard(
            con, session, {
                "week": params.get("week", [""])[0].strip(),
                "batch": params.get("batch", [""])[0].strip(),
                "unit": params.get("unit", [""])[0].strip(),
            }
        )
    finally:
        con.close()
    if data is None:
        return legacy_dashboard_page(session)
    periods, selected, previous = data["periods"], data["selected"], data["previous"]
    rows, previous_rows = data["rows"], data["previous_rows"]

    current_metrics = data["current_metrics"]
    previous_metrics = data["previous_metrics"]
    current_totals = {
        "dien_tich": current_metrics["area"]["total"],
        "san_luong": current_metrics["harvest"]["total"],
        "sold_total": current_metrics["consumption"]["total"],
        "remaining_total": current_metrics["remaining"]["total"],
        "area_land": current_metrics["area"]["land"],
        "area_tarp": current_metrics["area"]["tarp"],
        "harvest_land": current_metrics["harvest"]["land"],
        "harvest_tarp": current_metrics["harvest"]["tarp"],
        "sold_land": current_metrics["consumption"]["land"],
        "sold_tarp": current_metrics["consumption"]["tarp"],
        "remaining_land": current_metrics["remaining"]["land"],
        "remaining_tarp": current_metrics["remaining"]["tarp"],
        "households": current_metrics["households"]["total"],
        "workers": current_metrics["workers"]["total"],
    }
    previous_totals = {
        "dien_tich": previous_metrics["area"]["total"],
        "san_luong": previous_metrics["harvest"]["total"],
        "sold_total": previous_metrics["consumption"]["total"],
        "remaining_total": previous_metrics["remaining"]["total"],
        "area_land": previous_metrics["area"]["land"],
        "area_tarp": previous_metrics["area"]["tarp"],
        "harvest_land": previous_metrics["harvest"]["land"],
        "harvest_tarp": previous_metrics["harvest"]["tarp"],
        "sold_land": previous_metrics["consumption"]["land"],
        "sold_tarp": previous_metrics["consumption"]["tarp"],
        "remaining_land": previous_metrics["remaining"]["land"],
        "remaining_tarp": previous_metrics["remaining"]["tarp"],
        "households": previous_metrics["households"]["total"],
        "workers": previous_metrics["workers"]["total"],
    }

    def display_num(value):
        return "—" if value is None else fmt_num(value)

    period_label = date.fromisoformat(str(selected["report_date"])).strftime("%d/%m/%Y")
    previous_label = (previous["week_code"] + " · " + date.fromisoformat(str(previous["report_date"])).strftime("%d/%m/%Y")
                      if previous else "chưa có kỳ trước")

    def metric_panel(label, total_key, land_key, tarp_key, unit, symbol, amber=False):
        total = current_totals[total_key]
        before = previous_totals[total_key]
        tone = " amber" if amber else ""
        if previous and total is not None and before is not None:
            difference = total - before
            marker = "+" if difference > 0 else ""
            delta = f'<div class="metric-delta"><strong>{marker}{fmt_num(difference)}</strong> {esc(unit)} so với {esc(previous_label)}</div>'
        elif previous:
            delta = '<div class="metric-delta muted">Chưa đủ dữ liệu để so sánh</div>'
        else:
            delta = '<div class="metric-delta muted">Chưa có tuần trước để so sánh</div>'
        return f"""<article class="metric-panel{tone}">
          <div class="metric-top"><h3>{label}</h3>{icon(symbol)}</div>
          <div class="metric-middle"><div><div class="metric-number">{display_num(total)}</div><div class="metric-caption">{unit} · lũy tiến đến {esc(period_label)}</div>{delta}</div><div class="metric-symbol">{icon(symbol)}</div></div>
          <div class="metric-breakdown"><div><span>Muối đất</span><strong>{display_num(current_totals[land_key])} <small>{unit}</small></strong></div><div><span>Muối trải bạt</span><strong>{display_num(current_totals[tarp_key])} <small>{unit}</small></strong></div></div>
        </article>"""

    metrics = ''.join([
        metric_panel('DIỆN TÍCH SẢN XUẤT MUỐI', 'dien_tich', 'area_land', 'area_tarp', 'ha', 'area'),
        metric_panel('SẢN LƯỢNG MUỐI THU HOẠCH', 'san_luong', 'harvest_land', 'harvest_tarp', 'tấn', 'harvest'),
        metric_panel('SẢN LƯỢNG MUỐI TIÊU THỤ', 'sold_total', 'sold_land', 'sold_tarp', 'tấn', 'sold', True),
        metric_panel('SẢN LƯỢNG MUỐI CÒN LẠI', 'remaining_total', 'remaining_land', 'remaining_tarp', 'tấn', 'stock', True),
    ])
    options = ''.join(
        f'<option value="{period["week_code"]}" {"selected" if period["week_code"] == selected["week_code"] else ""}>'
        f'{esc(period["week_code"])}</option>'
        for period in periods
    )
    status_cards = ''.join([
        f'<a class="status-card draft" href="/salt/weekly"><span class="status-icon">{icon("file")}</span><div><div class="status-card-count">{len(periods)}</div><div class="status-card-label">Tuần có dữ liệu</div></div></a>',
        f'<a class="status-card approved" href="/salt/weekly?week={selected["week_code"]}"><span class="status-icon">{icon("check")}</span><div><div class="status-card-count">{len(rows)}</div><div class="status-card-label">Đơn vị trong kỳ</div></div></a>',
        f'<a class="status-card submitted" href="/salt/weekly?week={selected["week_code"]}"><span class="status-icon">{icon("chart")}</span><div><div class="status-card-count">{data["warning_rows"]}</div><div class="status-card-label">Dòng cảnh báo</div></div></a>',
        f'<a class="status-card returned" href="/salt/weekly?week={previous["week_code"] if previous else selected["week_code"]}"><span class="status-icon">{icon("return")}</span><div><div class="status-card-count">{previous["week_code"] if previous else "—"}</div><div class="status-card-label">Kỳ dùng để so sánh</div></div></a>',
    ])
    summary_rows = ''.join(
        f'<tr><td>{esc(row["unit_name"])}</td><td>{fmt_num(row["dien_tich"])}</td><td>{fmt_num(row["san_luong"])}</td><td>{fmt_num(row["sold_total"])}</td><td>{fmt_num(row["remaining_total"])}</td><td>{fmt_num(row["households"])}</td><td>{fmt_num(row["workers"])}</td></tr>'
        for row in sorted(rows, key=lambda item: item["unit_name"])
    ) or '<tr><td colspan="7" class="empty">Đơn vị này chưa có dữ liệu trong tuần.</td></tr>'
    scope = "Tất cả đơn vị" if is_chi_cuc_user(session) else session["unit_name"]
    body = f"""
    <div class="container">
      {take_flash(session)}
      <div class="page-head"><div><h1>Bảng giám sát</h1><div class="subtitle">Số liệu lũy tiến có hiệu lực của tuần và mức thay đổi so với tuần có dữ liệu gần nhất trước đó.</div></div><a class="btn primary" href="/import-excel">{icon('download')}Import báo cáo tuần</a></div>
      <div class="dashboard-toolbar"><form class="dashboard-period" method="get"><div class="field"><label>Kỳ đang xem</label><select name="week">{options}</select></div><button class="btn primary">Xem dashboard</button></form><div class="scope-label">{icon('location')}<span>Đơn vị: <strong>{esc(scope)}</strong></span></div></div>
      <div class="status-overview" aria-label="Tình trạng báo cáo tuần">{status_cards}</div>
      <section class="dashboard-panel">
        <div class="panel-heading"><h2 class="panel-title"><span class="panel-title-icon">{icon('chart')}</span>SỐ LIỆU SẢN XUẤT MUỐI</h2><div class="panel-meta">{esc(selected['week_code'])}<br>So sánh với: {esc(previous['week_code']) if previous else 'Chưa có tuần trước'}</div></div>
        <div class="dashboard-metrics">{metrics}</div>
        <div class="people-metrics"><div class="people-card"><div><div class="label">Số hộ làm muối</div><div class="value">{display_num(current_totals['households'])} <small>hộ</small></div></div>{icon('home')}</div><div class="people-card"><div><div class="label">Lao động làm muối</div><div class="value">{display_num(current_totals['workers'])} <small>người</small></div></div>{icon('users')}</div></div>
        <p class="dashboard-note">Số liệu có hiệu lực của tuần đang chọn. Chênh lệch bằng tuần đang xem trừ tuần có dữ liệu gần nhất trước đó; các lần upload cùng tuần không tạo thêm kỳ so sánh.</p>
      </section>
      <section class="dashboard-panel"><div class="panel-heading"><h2 class="panel-title"><span class="panel-title-icon">{icon('table')}</span>CHI TIẾT THEO ĐƠN VỊ</h2><a class="btn small" href="/salt/weekly?week={selected['week_code']}">Lịch sử Excel của tuần</a></div><div class="table-wrap"><table class="summary-table"><thead><tr><th>Đơn vị</th><th>Diện tích (ha)</th><th>Thu hoạch (tấn)</th><th>Tiêu thụ (tấn)</th><th>Còn lại (tấn)</th><th>Số hộ</th><th>Lao động</th></tr></thead><tbody>{summary_rows}</tbody></table></div></section>
    </div>"""
    return base_page("Tổng quan", body, session)


def legacy_dashboard_page(session):
    con = db_conn()
    if is_chi_cuc_user(session):
        approved_rows = con.execute("SELECT * FROM records WHERE status='approved' ORDER BY report_date DESC, id DESC").fetchall()
        status_counts = {s: con.execute("SELECT COUNT(*) FROM records WHERE status=?", (s,)).fetchone()[0] for s in STATUS_LABELS}
        standard_rows = con.execute(
            """SELECT s.*, d.TenDonVi unit_name
               FROM DN_SanLuongMuoi s
               LEFT JOIN DM_DonViHanhChinh d
                 ON d.Ma_DonViHanhChinh=s.Ma_DonViHanhChinh"""
        ).fetchall()
    else:
        approved_rows = con.execute("SELECT * FROM records WHERE unit_name=? AND status='approved' ORDER BY report_date DESC, id DESC", (session["unit_name"],)).fetchall()
        status_counts = {s: con.execute("SELECT COUNT(*) FROM records WHERE unit_name=? AND status=?", (session["unit_name"], s)).fetchone()[0] for s in STATUS_LABELS}
        standard_rows = con.execute(
            """SELECT s.*, d.TenDonVi unit_name
               FROM DN_SanLuongMuoi s
               JOIN DM_DonViHanhChinh d
                 ON d.Ma_DonViHanhChinh=s.Ma_DonViHanhChinh
               WHERE d.TenDonVi=?""",
            (session["unit_name"],),
        ).fetchall()

    # Operational details are kept in app.records. Use only the latest
    # cumulative report per unit/month so weekly cumulative reports are not
    # added together. Official area/production below always come from QD 5277.
    rows = []
    seen_periods = set()
    for row in approved_rows:
        if row["reporting_mode"] != "cumulative":
            continue
        key = (row["unit_name"], str(row["report_date"])[:7])
        if key not in seen_periods:
            seen_periods.add(key)
            rows.append(row)

    totals = {"area_total":0,"harvest_total":0,"sold_total":0,"remaining_total":0,"households":0,"workers":0}
    for r in rows:
        t = record_totals(r)
        totals["sold_total"] += t["sold_total"]
        totals["remaining_total"] += t["remaining_total"]
        totals["households"] += r["households"] or 0
        totals["workers"] += r["workers"] or 0
    totals["area_total"] = sum(float(r["DienTich"] or 0) for r in standard_rows)
    totals["harvest_total"] = sum(float(r["SanLuong"] or 0) for r in standard_rows)

    if is_chi_cuc_user(session):
        summaries = {}
        for r in standard_rows:
            name = r["unit_name"] or r["Ma_DonViHanhChinh"]
            item = summaries.setdefault(name, {"unit_name":name, "area_total":0,
                "harvest_total":0, "sold_total":0, "remaining_total":0,
                "households":0, "workers":0, "periods":set()})
            item["area_total"] += float(r["DienTich"] or 0)
            item["harvest_total"] += float(r["SanLuong"] or 0)
            item["periods"].add(str(r["Ma_ThoiGian"]))
        for r in rows:
            item = summaries.setdefault(r["unit_name"], {"unit_name":r["unit_name"],
                "area_total":0, "harvest_total":0, "sold_total":0,
                "remaining_total":0, "households":0, "workers":0, "periods":set()})
            detail = record_totals(r)
            item["sold_total"] += float(detail["sold_total"] or 0)
            item["remaining_total"] += float(detail["remaining_total"] or 0)
            item["households"] += float(r["households"] or 0)
            item["workers"] += float(r["workers"] or 0)
            item["periods"].add(str(r["report_date"])[:7])
        unit_summary = []
        for item in summaries.values():
            item["reports"] = len(item.pop("periods"))
            unit_summary.append(item)
        unit_summary.sort(key=lambda item: item["unit_name"])
    else:
        unit_summary = []
    con.close()

    breakdowns = {
        key: sum(float(r[key] or 0) for r in rows)
        for key in ('area_land', 'area_tarp', 'harvest_land', 'harvest_tarp',
                    'sold_land', 'sold_tarp', 'remaining_land', 'remaining_tarp')
    }

    def metric_panel(label, total_key, land_key, tarp_key, unit, symbol, amber=False):
        total = float(totals[total_key] or 0)
        land, tarp = breakdowns[land_key], breakdowns[tarp_key]
        tone = ' amber' if amber else ''
        if amber:
            share = max(0, min(100, land / total * 100)) if total > 0 else 0
            share_label = f'{fmt_num(share, 1)}%' if total > 0 else '—'
            visual = f'<div class="donut" style="--share:{share:.2f}%" role="img" aria-label="Tỷ trọng muối đất: {share_label}"><div class="donut-label"><strong>{share_label}</strong><small>Muối đất</small></div></div>'
        else:
            visual = f'<div class="metric-symbol">{icon(symbol)}</div>'
        return f"""<article class="metric-panel{tone}">
          <div class="metric-top"><h3>{label}</h3>{icon(symbol)}</div>
          <div class="metric-middle"><div><div class="metric-number">{fmt_num(total)}</div><div class="metric-caption">{unit} · tổng số đã duyệt</div></div>{visual}</div>
          <div class="metric-breakdown"><div><span>Muối đất</span><strong>{fmt_num(land)} <small>{unit}</small></strong></div><div><span>Muối trải bạt</span><strong>{fmt_num(tarp)} <small>{unit}</small></strong></div></div>
        </article>"""

    metrics = ''.join([
        metric_panel('DIỆN TÍCH SẢN XUẤT MUỐI', 'area_total', 'area_land', 'area_tarp', 'ha', 'area'),
        metric_panel('SẢN LƯỢNG MUỐI THU HOẠCH', 'harvest_total', 'harvest_land', 'harvest_tarp', 'tấn', 'harvest'),
        metric_panel('SẢN LƯỢNG MUỐI TIÊU THỤ', 'sold_total', 'sold_land', 'sold_tarp', 'tấn', 'sold', True),
        metric_panel('SẢN LƯỢNG MUỐI CÒN LẠI', 'remaining_total', 'remaining_land', 'remaining_tarp', 'tấn', 'stock', True),
    ])
    status_symbols = {'draft':'file', 'submitted':'send', 'approved':'check', 'returned':'return'}
    status_short = {'draft':'Bản nháp', 'submitted':'Chờ Chi cục duyệt', 'approved':'Đã duyệt', 'returned':'Cần chỉnh sửa'}
    statuses = ''.join(
        f'<a class="status-card {state}" href="/records?status={state}"><span class="status-icon">{icon(status_symbols[state])}</span><div><div class="status-card-count">{status_counts[state]}</div><div class="status-card-label">{status_short[state]}</div></div></a>'
        for state in STATUS_LABELS
    )
    scope = 'Tất cả đơn vị' if is_chi_cuc_user(session) else session['unit_name']
    latest = max((r['updated_at'] for r in rows), default='')
    updated_label = f'Cập nhật báo cáo: {esc(latest)}' if latest else 'Chưa có báo cáo được duyệt'
    summary_html = ''
    if is_chi_cuc_user(session):
        if unit_summary:
            rows_html = ''.join(
                f"<tr><td>{esc(r['unit_name'])}</td><td>{fmt_num(r['area_total'])}</td><td>{fmt_num(r['harvest_total'])}</td><td>{fmt_num(r['sold_total'])}</td><td>{fmt_num(r['remaining_total'])}</td><td>{fmt_num(r['households'])}</td><td>{fmt_num(r['workers'])}</td><td>{r['reports']}</td></tr>"
                for r in unit_summary
            )
            summary_content = f'<div class="table-wrap"><table class="summary-table"><thead><tr><th>Đơn vị</th><th>Diện tích (ha)</th><th>Thu hoạch (tấn)</th><th>Tiêu thụ (tấn)</th><th>Còn lại (tấn)</th><th>Số hộ</th><th>Lao động</th><th>Số kỳ</th></tr></thead><tbody>{rows_html}</tbody></table></div>'
        else:
            summary_content = f'<div class="summary-empty">{icon("table")}<div><strong>Chưa có số liệu để tổng hợp</strong><p>Báo cáo của các đơn vị sẽ xuất hiện tại đây sau khi được Chi cục duyệt.</p></div></div>'
        summary_html = f'<section class="dashboard-panel"><div class="panel-heading"><h2 class="panel-title"><span class="panel-title-icon">{icon("table")}</span>TỔNG HỢP THEO ĐƠN VỊ</h2><a class="btn small" href="/records?status=approved">Xem báo cáo</a></div>{summary_content}</section>'

    body = f"""
    <div class="container">
      {take_flash(session)}
      <div class="page-head"><div><h1>Bảng giám sát</h1><div class="subtitle">Theo dõi số liệu sản xuất muối và tình trạng báo cáo.</div></div><a class="btn primary" href="/records/new">{icon('plus')}Nhập số liệu</a></div>
      <div class="dashboard-toolbar"><nav class="view-tabs" aria-label="Chế độ xem"><a class="view-tab active" href="/dashboard" aria-current="page">Tổng quan</a><a class="view-tab" href="/records">Dữ liệu chi tiết</a></nav><div class="scope-label">{icon('location')}<span>Đơn vị: <strong>{esc(scope)}</strong></span></div></div>
      <div class="status-overview" aria-label="Tình trạng báo cáo">{statuses}</div>
      <section class="dashboard-panel">
        <div class="panel-heading"><h2 class="panel-title"><span class="panel-title-icon">{icon('chart')}</span>SỐ LIỆU SẢN XUẤT MUỐI</h2><div class="panel-meta">{len(rows)} báo cáo đã duyệt · Tất cả kỳ báo cáo<br>{updated_label}</div></div>
        <div class="dashboard-metrics">{metrics}</div>
        <div class="people-metrics"><div class="people-card"><div><div class="label">Số hộ làm muối</div><div class="value">{fmt_num(totals['households'])} <small>hộ</small></div></div>{icon('home')}</div><div class="people-card"><div><div class="label">Lao động làm muối</div><div class="value">{fmt_num(totals['workers'])} <small>người</small></div></div>{icon('users')}</div></div>
        <p class="dashboard-note">Các chỉ tiêu được cộng từ tất cả báo cáo đã duyệt; số hộ và lao động có thể lặp lại giữa các kỳ.</p>
      </section>
      {summary_html}
      <div class="actions"><a class="btn" href="/export.xlsx?status=approved">{icon('download')}Xuất Excel số liệu đã duyệt</a><a class="btn" href="/records">Tra cứu báo cáo</a></div>
    </div>
    """
    return base_page("Tổng quan", body, session)


def record_action_buttons(session, r):
    actions = [f'<a class="btn small" href="/records/{r["id"]}">Xem</a>']
    if is_chi_cuc_user(session):
        actions.append(f'<a class="btn small" href="/records/{r["id"]}/edit">Sửa</a>')
        if r["status"] == STATUS_SUBMITTED:
            actions.append(f'<form style="display:inline" method="post" action="/records/{r["id"]}/approve">{csrf_input(session)}<button class="btn small ok">Duyệt</button></form>')
            actions.append(f'<a class="btn small warn" href="/records/{r["id"]}/return">Trả chỉnh sửa</a>')
    else:
        if r["status"] in (STATUS_DRAFT, STATUS_RETURNED):
            actions.append(f'<a class="btn small" href="/records/{r["id"]}/edit">Sửa</a>')
            actions.append(f'<form style="display:inline" method="post" action="/records/{r["id"]}/submit">{csrf_input(session)}<button class="btn small primary">Gửi Chi cục</button></form>')
    return '<div class="actions">' + ''.join(actions) + '</div>'


def records_page(session, query):
    """Legacy approval-list screen kept internal while weekly lookup is public."""
    filters = filters_from_query(query, session)
    con = db_conn()
    records = fetch_records(con, session, filters)
    units = get_units(con)
    con.close()

    unit_options = '<option value="">Tất cả đơn vị</option>' + ''.join(f'<option value="{esc(u)}" {"selected" if u==filters["unit"] else ""}>{esc(u)}</option>' for u in units)
    status_options = '<option value="">Tất cả trạng thái</option>' + ''.join(f'<option value="{s}" {"selected" if s==filters["status"] else ""}>{STATUS_LABELS[s]}</option>' for s in STATUS_LABELS)
    unit_field = '' if session["role"] == ROLE_UNIT else f'<div class="field"><label>Đơn vị</label><select name="unit">{unit_options}</select></div>'

    trs = []
    for r in records:
        t = record_totals(r)
        trs.append(f"""
        <tr>
          <td>{esc(r['report_date'])}</td><td>{esc(r['unit_name'])}</td><td><span class="status {esc(r['status'])}">{STATUS_LABELS[r['status']]}</span></td>
          <td>{fmt_num(t['area_total'])}</td><td>{fmt_num(t['harvest_total'])}</td><td>{fmt_num(t['sold_total'])}</td><td>{fmt_num(t['remaining_total'])}</td><td>{fmt_num(r['households'])}</td><td>{fmt_num(r['workers'])}</td><td>{fmt_num(t['avg_yield'],2)}</td>
          <td>{record_action_buttons(session,r)}</td>
        </tr>""")
    table_body = ''.join(trs) if trs else '<tr><td colspan="11" class="empty">Chưa có dữ liệu phù hợp.</td></tr>'

    query_string = '&'.join(f"{quote(k)}={quote(v)}" for k,v in filters.items() if v)
    standard_button = '<a class="btn" href="/standard-data">Dữ liệu chuẩn hóa</a>' if is_chi_cuc_user(session) else ''
    body = f"""
    <div class="container">
      {take_flash(session)}
      <div class="page-head"><div><h1>Dữ liệu báo cáo</h1><div class="subtitle">Tra cứu, lọc và xuất dữ liệu theo kỳ báo cáo/đơn vị/trạng thái.</div></div><div class="actions"><a class="btn primary" href="/records/new">+ Nhập số liệu</a><a class="btn ok" href="/export.xlsx?{query_string}">Xuất Excel</a>{standard_button}</div></div>
      <div class="card"><form class="toolbar" method="get" action="/records">
        <div class="field"><label>Từ khóa</label><input name="q" value="{esc(filters['q'])}" placeholder="Đơn vị, ghi chú..."></div>
        {unit_field}
        <div class="field"><label>Trạng thái</label><select name="status">{status_options}</select></div>
        <div class="field"><label>Từ ngày</label><input type="date" name="from" value="{esc(filters['from'])}"></div>
        <div class="field"><label>Đến ngày</label><input type="date" name="to" value="{esc(filters['to'])}"></div>
        <button class="btn primary">Tra cứu</button><a class="btn" href="/records">Xóa lọc</a>
      </form></div>
      <div class="card"><div class="table-wrap"><table class="table"><thead><tr><th>Kỳ báo cáo</th><th>Đơn vị</th><th>Trạng thái</th><th>Diện tích (ha)</th><th>Thu hoạch (tấn)</th><th>Tiêu thụ (tấn)</th><th>Còn lại (tấn)</th><th>Số hộ</th><th>Lao động</th><th>Năng suất BQ</th><th>Thao tác</th></tr></thead><tbody>{table_body}</tbody></table></div></div>
    </div>
    """
    return base_page("Dữ liệu báo cáo", body, session)


def standard_data_page(session):
    """Technical QD 5277 view; intentionally excluded from end-user navigation."""
    if not is_chi_cuc_user(session):
        return None

    con = db_conn()
    rows = con.execute(
        """
        SELECT
            s.Ma_SanLuongMuoi,
            s.Ma_DonViHanhChinh,
            d.TenDonVi,
            s.Ma_ThoiGian,
            s.PhuongPhapSX,
            s.DienTich,
            s.SanLuong,
            s.GiaBanBinhQuan
        FROM DN_SanLuongMuoi s
        LEFT JOIN DM_DonViHanhChinh d
               ON d.Ma_DonViHanhChinh = s.Ma_DonViHanhChinh
        ORDER BY s.Ma_ThoiGian DESC, d.TenDonVi ASC, s.PhuongPhapSX ASC
        """
    ).fetchall()
    con.close()

    def display_time(code):
        text = str(code or "")
        if len(text) == 8 and text.isdigit():
            return f"{text[6:8]}/{text[4:6]}/{text[0:4]}"
        return text

    trs = []
    for r in rows:
        code_class = "muted" if str(r["Ma_DonViHanhChinh"] or "").startswith("TMP") else ""
        trs.append(f"""
        <tr>
          <td>{r['Ma_SanLuongMuoi']}</td>
          <td><span class="{code_class}">{esc(r['Ma_DonViHanhChinh'])}</span></td>
          <td>{esc(r['TenDonVi'] or '')}</td>
          <td>{esc(display_time(r['Ma_ThoiGian']))}</td>
          <td>{esc(r['PhuongPhapSX'])}</td>
          <td>{fmt_num(r['DienTich'], 2)}</td>
          <td>{fmt_num(r['SanLuong'], 2)}</td>
          <td>{fmt_num(r['GiaBanBinhQuan'], 2)}</td>
        </tr>
        """)

    table_body = ''.join(trs) if trs else '<tr><td colspan="8" class="empty">Chưa có báo cáo tháng được duyệt để xuất bản theo QĐ 5277.</td></tr>'

    body = f"""
    <div class="container">
      {take_flash(session)}
      <div class="page-head">
        <div>
          <h1>Dữ liệu chuẩn hóa</h1>
          <div class="subtitle">
            Dữ liệu báo cáo tháng đã duyệt, xuất bản theo cấu trúc DN_SanLuongMuoi của QĐ 5277.
          </div>
        </div>
        <div class="actions">
          <a class="btn" href="/records">← Dữ liệu báo cáo</a>
        </div>
      </div>

      <div class="notice info">
        Tổng cộng <b>{len(rows)}</b> bản ghi chuẩn hóa. Sheet tuần chỉ dùng cho theo dõi và so sánh;
        chưa tự ghi vào đây để tránh hiểu số liệu lũy tiến tuần thành báo cáo tháng.
      </div>

      <div class="card">
        <div class="table-wrap">
          <table class="summary-table">
            <thead>
              <tr>
                <th>ID</th>
                <th>Mã đơn vị</th>
                <th>Đơn vị</th>
                <th>Kỳ số liệu</th>
                <th>Phương pháp SX</th>
                <th>Diện tích (ha)</th>
                <th>Sản lượng (tấn)</th>
                <th>Giá bán BQ (đ/kg)</th>
              </tr>
            </thead>
            <tbody>{table_body}</tbody>
          </table>
        </div>
      </div>
    </div>
    """
    return base_page("Dữ liệu chuẩn hóa", body, session)


def input_field(name, label, value="", typ="number", step="0.1", extra=""):
    attrs = f'type="{typ}" name="{name}" value="{esc(value)}"'
    if typ == "number":
        attrs += f' min="0" step="{step}"'
    return f'<div class="field"><label>{esc(label)}</label><input {attrs} {extra}></div>'


def record_form_page(session, record=None, error=""):
    con = db_conn()
    units = get_units(con)
    con.close()
    is_edit = record is not None
    v = lambda key, default="": record[key] if record is not None and key in record.keys() else default
    unit_name = session["unit_name"] if session["role"] == ROLE_UNIT else v("unit_name", units[0] if units else "")
    unit_select = ''.join(f'<option value="{esc(u)}" {"selected" if u==unit_name else ""}>{esc(u)}</option>' for u in units)
    unit_html = f'<div class="field"><label>Đơn vị xã/phường</label><select name="unit_name" required>{unit_select}</select></div>' if is_chi_cuc_user(session) else f'<div class="field"><label>Đơn vị xã/phường</label><input value="{esc(unit_name)}" readonly></div><input type="hidden" name="unit_name" value="{esc(unit_name)}">'
    action = f'/records/{record["id"]}/edit' if is_edit else '/records/new'
    notice = f'<div class="notice err">{esc(error)}</div>' if error else ''
    reviewer = ""
    if is_edit and v("reviewer_note"):
        reviewer = f'<div class="notice info"><b>Ý kiến Chi cục:</b> {esc(v("reviewer_note"))}</div>'
    body = f"""
    <div class="container">
      <div class="page-head">
        <div>
          <h1>{'Cập nhật' if is_edit else 'Nhập'} số liệu</h1>
          <div class="subtitle">Các cột “Cộng” và “Năng suất bình quân” được hệ thống tự tính.</div>
        </div>
        <div class="actions">
          <a class="btn ok" href="/import-excel">{icon('download')} Nhập dữ liệu từ Excel</a>
          <a class="btn" href="/records">← Danh sách</a>
        </div>
      </div>
      {notice}{reviewer}
      <form class="card" method="post" action="{action}">{csrf_input(session)}
        <div class="grid">
          <div class="field"><label>Ngày/kỳ chốt số liệu</label><input type="date" name="report_date" value="{esc(v('report_date', date.today().isoformat()))}" required></div>
          {unit_html}
        </div>

        <div class="section-title">1. Diện tích sản xuất muối (ha)</div><div class="grid">
          {input_field('area_land','Muối đất',v('area_land',0))}
          {input_field('area_tarp','Muối trải bạt',v('area_tarp',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="area_total" readonly></div>
        </div>

        <div class="section-title">2. Sản lượng muối thu hoạch (tấn)</div><div class="grid">
          {input_field('harvest_land','Muối đất',v('harvest_land',0))}
          {input_field('harvest_tarp','Muối trải bạt',v('harvest_tarp',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="harvest_total" readonly></div>
        </div>

        <div class="section-title">3. Sản lượng muối tiêu thụ (tấn)</div><div class="grid">
          {input_field('sold_land','Muối đất',v('sold_land',0))}
          {input_field('sold_tarp','Muối trải bạt',v('sold_tarp',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="sold_total" readonly></div>
        </div>

        <div class="section-title">4. Sản lượng còn lại (tấn)</div><div class="grid">
          {input_field('remaining_land','Muối đất',v('remaining_land',0))}
          {input_field('remaining_tarp','Muối trải bạt',v('remaining_tarp',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="remaining_total" readonly></div>
        </div>

        <div class="section-title">5. Sản lượng muối chế biến (tấn)</div><div class="grid">
          {input_field('processed_fine','Muối tinh',v('processed_fine',0))}
          {input_field('processed_iodized','Muối I-ốt',v('processed_iodized',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="processed_total" readonly></div>
        </div>

        <div class="section-title">6. Hộ, lao động, giá bán và năng suất</div><div class="grid">
          {input_field('households','Số hộ làm muối (hộ)',v('households',0),'number','1')}
          {input_field('workers','Số lao động làm muối (người)',v('workers',0),'number','1')}
          {input_field('price_land','Giá bán muối đất (đồng/kg)',v('price_land',''),'text','','placeholder="Ví dụ 1.000 - 1.500"')}
          {input_field('price_tarp','Giá bán muối trải bạt (đồng/kg)',v('price_tarp',''),'text','','placeholder="Ví dụ 1.200 - 1.800"')}
          <div class="field"><label>Năng suất BQ (tấn/ha, tự tính)</label><input class="calc" id="avg_yield" readonly></div>
        </div>

        <div class="section-title">7. Thiệt hại do mưa trái mùa (tấn)</div><div class="grid">
          {input_field('damage_land','Muối đất',v('damage_land',0))}
          {input_field('damage_tarp','Muối trải bạt',v('damage_tarp',0))}
          <div class="field"><label>Cộng (tự tính)</label><input class="calc" id="damage_total" readonly></div>
        </div>

        <div class="section-title">8. Ghi chú</div><div class="field"><textarea name="note" placeholder="Ghi chú, giải trình số liệu...">{esc(v('note',''))}</textarea></div>
        <div class="actions" style="margin-top:16px"><button class="btn primary">Lưu dữ liệu</button><a class="btn" href="/records">Hủy</a></div>
      </form>
    </div>
    <script>
    function n(name){{const e=document.querySelector('[name="'+name+'"]'); return e?parseFloat((e.value||'0').replace(',','.'))||0:0;}}
    function calc(){{
      const pairs=[['area_land','area_tarp','area_total'],['harvest_land','harvest_tarp','harvest_total'],['sold_land','sold_tarp','sold_total'],['remaining_land','remaining_tarp','remaining_total'],['processed_fine','processed_iodized','processed_total'],['damage_land','damage_tarp','damage_total']];
      pairs.forEach(p=>document.getElementById(p[2]).value=(n(p[0])+n(p[1])).toFixed(1));
      const a=n('area_land')+n('area_tarp'), h=n('harvest_land')+n('harvest_tarp');
      document.getElementById('avg_yield').value=a?(h/a).toFixed(2):'0.00';
    }}
    document.querySelectorAll('input[type="number"]').forEach(e=>e.addEventListener('input',calc)); calc();
    </script>
    """
    return base_page("Nhập số liệu", body, session)


def can_access_record(session, r):
    if is_chi_cuc_user(session):
        return True
    con = db_conn()
    try:
        return r["unit_name"] in admin_units.report_names(con, session["unit_name"])
    finally:
        con.close()


def detail_page(session, rid):
    con = db_conn()
    r = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    if not r or not can_access_record(session, r):
        con.close(); return None
    logs = con.execute("SELECT a.*, u.username FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id WHERE record_id=? ORDER BY a.id DESC", (rid,)).fetchall()
    con.close()
    t = record_totals(r)
    rows = [
        ("Diện tích sản xuất muối (ha)", t["area_total"], r["area_land"], r["area_tarp"]),
        ("Sản lượng thu hoạch (tấn)", t["harvest_total"], r["harvest_land"], r["harvest_tarp"]),
        ("Sản lượng tiêu thụ (tấn)", t["sold_total"], r["sold_land"], r["sold_tarp"]),
        ("Sản lượng còn lại (tấn)", t["remaining_total"], r["remaining_land"], r["remaining_tarp"]),
        ("Thiệt hại do mưa trái mùa (tấn)", t["damage_total"], r["damage_land"], r["damage_tarp"]),
    ]
    trs = ''.join(f'<tr><td>{esc(a)}</td><td>{fmt_num(b)}</td><td>{fmt_num(c)}</td><td>{fmt_num(d)}</td></tr>' for a,b,c,d in rows)
    logs_html = ''.join(f'<tr><td>{esc(l["created_at"])}</td><td>{esc(l["username"] or "")}</td><td>{esc(l["action"])}</td><td>{esc(l["detail"])}</td></tr>' for l in logs) or '<tr><td colspan="4" class="empty">Chưa có nhật ký.</td></tr>'
    actions = record_action_buttons(session, r)
    body = f"""
    <div class="container">{take_flash(session)}
      <div class="page-head"><div><h1>{esc(r['unit_name'])} · {esc(r['report_date'])}</h1><div class="subtitle"><span class="status {esc(r['status'])}">{STATUS_LABELS[r['status']]}</span></div></div>{actions}</div>
      <div class="card"><div class="table-wrap"><table class="summary-table"><thead><tr><th>Chỉ tiêu</th><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th></tr></thead><tbody>{trs}</tbody></table></div>
      <div class="section-title">Chỉ tiêu khác</div><div class="grid">
        <div><b>Muối chế biến:</b> {fmt_num(t['processed_total'])} tấn<br><span class="muted">Muối tinh {fmt_num(r['processed_fine'])}; I-ốt {fmt_num(r['processed_iodized'])}</span></div>
        <div><b>Số hộ:</b> {fmt_num(r['households'])} hộ<br><b>Lao động:</b> {fmt_num(r['workers'])} người</div>
        <div><b>Giá muối đất:</b> {esc(r['price_land'] or '-')}<br><b>Giá trải bạt:</b> {esc(r['price_tarp'] or '-')}</div>
        <div><b>Năng suất BQ:</b> {fmt_num(t['avg_yield'],2)} tấn/ha</div>
      </div>
      <div class="section-title">Ghi chú đơn vị</div><div>{esc(r['note'] or '-')}</div>
      <div class="section-title">Ý kiến Chi cục</div><div>{esc(r['reviewer_note'] or '-')}</div>
      </div>
      <div class="card"><div class="section-title" style="margin-top:0">Nhật ký xử lý</div><div class="table-wrap"><table class="summary-table"><thead><tr><th>Thời gian</th><th>Tài khoản</th><th>Hành động</th><th>Nội dung</th></tr></thead><tbody>{logs_html}</tbody></table></div></div>
    </div>
    """
    return base_page("Chi tiết báo cáo", body, session)


def return_page(session, rid):
    con = db_conn(); r = con.execute("SELECT * FROM records WHERE id=?",(rid,)).fetchone(); con.close()
    if not r or not is_chi_cuc_user(session):
        return None
    body = f"""
    <div class="container"><div class="page-head"><div><h1>Yêu cầu chỉnh sửa</h1><div class="subtitle">{esc(r['unit_name'])} · {esc(r['report_date'])}</div></div></div>
    <form class="card" method="post" action="/records/{rid}/return">{csrf_input(session)}<div class="field"><label>Nội dung yêu cầu đơn vị cập nhật</label><textarea name="reviewer_note" required autofocus>{esc(r['reviewer_note'])}</textarea></div><div class="actions" style="margin-top:12px"><button class="btn warn">Gửi yêu cầu chỉnh sửa</button><a class="btn" href="/records/{rid}">Hủy</a></div></form></div>
    """
    return base_page("Yêu cầu chỉnh sửa", body, session)

def users_page(session):
    if not can_manage_users(session):
        return None

    con = db_conn()
    users = repositories.list_users(con)
    unit_fields = admin_unit_pages.account_fields(con)
    con.close()

    trs = ""
    for u in users:
        cannot_delete = u["id"] == session["user_id"]

        actions = f"""
        <div class="actions">
            <a class="btn small" href="/users/{u['id']}/edit">Sửa</a>
        """

        if not cannot_delete:
            actions += f"""
            <form method="post"
                  action="/users/{u['id']}/{'deactivate' if u['active'] else 'activate'}"
                  style="display:inline"
                  >
                {csrf_input(session)}
                <button class="btn small warn" type="submit">{"Ngưng kích hoạt" if u["active"] else "Kích hoạt lại"}</button>
            </form>
            """

        actions += "</div>"

        trs += f"""
        <tr>
            <td>{esc(u["username"])}</td>
            <td>{esc(ROLE_LABELS.get(u["role"], u["role"]))}</td>
            <td>{esc(u["unit_name"] or "")}</td>
            <td>{"Hoạt động" if u["active"] else "Ngưng hoạt động"}</td>
            <td>{actions}</td>
        </tr>
        """

    body = f"""
    <div class="container">
      {take_flash(session)}

      <div class="page-head">
        <div>
          <h1>Tài khoản đơn vị</h1>
          <div class="subtitle">
            Quản lý tài khoản Chi cục và các đơn vị xã/phường.
          </div>
        </div>
        {'<a class="btn" href="/ocop/access">Phân địa bàn OCOP</a>' if ocop_available() else ''}
      </div>

      <div class="card">
        <form method="post" action="/users/new">
          {csrf_input(session)}
          <div class="grid">

            <div class="field">
              <label>Tên đăng nhập</label>
              <input name="username" required>
            </div>

            <div class="field">
              <label>Mật khẩu ban đầu</label>
              <input type="password" name="password" required minlength="6">
            </div>

            <div class="field">
              <label>Vai trò</label>
              <select name="role">
                <option value="unit">Đơn vị xã/phường</option>
                <option value="staff">Chuyên viên Chi cục</option>
                <option value="admin">Quản trị Chi cục</option>
              </select>
            </div>

            {unit_fields}

          </div>

          <h2 class="section-title">Thông tin cá nhân</h2>
          {user_profiles.personal_fields({})}

          <button class="btn primary" style="margin-top:12px">
            + Thêm tài khoản
          </button>
        </form>
      </div>

      <div class="card">
        <div class="table-wrap">
          <table class="summary-table">
            <thead>
              <tr>
                <th>Tên đăng nhập</th>
                <th>Vai trò</th>
                <th>Đơn vị</th>
                <th>Trạng thái</th>
                <th>Thao tác</th>
              </tr>
            </thead>
            <tbody>
              {trs}
            </tbody>
          </table>
        </div>
      </div>

    </div>
    """

    return base_page("Tài khoản", body, session)


def user_edit_page(session, user, error="", profile_data=None):
    if not can_manage_users(session):
        return None

    notice = (
        f'<div class="notice err">{esc(error)}</div>'
        if error else ""
    )

    unit_selected = "selected" if user["role"] == ROLE_UNIT else ""
    staff_selected = "selected" if user["role"] == ROLE_STAFF else ""
    admin_selected = "selected" if user["role"] == ROLE_ADMIN else ""
    active_checked = "checked" if user["active"] else ""
    con = db_conn()
    try:
        unit_fields = admin_unit_pages.account_fields(con, user)
        profile = user_profiles.get_profile(con, user["id"])
        if profile_data:
            profile.update({field: profile_data[field] for field in user_profiles.PROFILE_FIELDS if field in profile_data})
    finally:
        con.close()

    body = f"""
    <div class="container">

      <div class="page-head">
        <div>
          <h1>Sửa tài khoản</h1>
          <div class="subtitle">
            Cập nhật thông tin tài khoản và đơn vị.
          </div>
        </div>

        <a class="btn" href="/users">← Quay lại</a>
      </div>

      {notice}

      <form class="card"
            method="post"
            action="/users/{user['id']}/edit">

        {csrf_input(session)}

        <div class="grid">

          <div class="field">
            <label>Tên đăng nhập</label>
            <input name="username"
                   value="{esc(user['username'])}"
                   required>
          </div>

          <div class="field">
            <label>Mật khẩu mới</label>
            <input type="password"
                   name="password"
                   minlength="6"
                   placeholder="Để trống nếu không đổi">
          </div>

          <div class="field">
            <label>Vai trò</label>
            <select name="role">
              <option value="unit" {unit_selected}>
                Đơn vị xã/phường
              </option>

              <option value="staff" {staff_selected}>Chuyên viên Chi cục</option>
              <option value="admin" {admin_selected}>
                Quản trị Chi cục
              </option>
            </select>
          </div>

          {unit_fields}

          <div class="field">
            <label>
              <input type="checkbox"
                     name="active"
                     value="1"
                     {active_checked}>
              Cho phép đăng nhập
            </label>
          </div>

        </div>

        <h2 class="section-title">Thông tin cá nhân</h2>
        {user_profiles.personal_fields(profile)}

        <div class="actions" style="margin-top:16px">
          <button class="btn primary" type="submit">
            Lưu thay đổi
          </button>

          <a class="btn" href="/users">Hủy</a>
        </div>

      </form>

    </div>
    """

    return base_page("Sửa tài khoản", body, session)

def legacy_import_excel_page(session):
    if not is_chi_cuc_user(session):
        return None

    body = f"""
    <div class="container">

        {take_flash(session)}

        <div class="page-head">
            <div>
                <h1>Import dữ liệu Excel</h1>

                <div class="subtitle">
                    Nhập dữ liệu sản xuất muối từ file Excel lịch sử.
                </div>
            </div>

            <a class="btn" href="/records">
                ← Dữ liệu báo cáo
            </a>
        </div>

        <form
            class="card"
            method="post"
            action="/import-excel"
            enctype="multipart/form-data"
        >

            {csrf_input(session)}

            <div class="grid">

                <div class="field">
                    <label>File Excel (.xlsx)</label>

                    <input
                        type="file"
                        name="excel_file"
                        accept=".xlsx"
                        required
                    >
                </div>


                <div class="field">
                    <label>Ngày/kỳ báo cáo</label>

                    <input
                        type="date"
                        name="report_date"
                        required
                    >
                </div>


                <div class="field">
                    <label>Tên sheet</label>

                    <input
                        name="sheet_name"
                        placeholder="Ví dụ: 03-02. Để trống = sheet đầu tiên"
                    >
                </div>


                <div class="field">
                    <label>Nếu dữ liệu đã tồn tại</label>

                    <select name="mode">

                        <option value="skip">
                            Bỏ qua dữ liệu đã có
                        </option>

                        <option value="update">
                            Cập nhật lại dữ liệu
                        </option>

                    </select>
                </div>

            </div>


            <div class="notice info" style="margin-top:16px">

                Hệ thống sẽ đọc các xã/phường trong cột B và tự bỏ qua
                dòng TỔNG CỘNG.

                <br><br>

                Các cột Cộng và Năng suất bình quân không cần import
                vì hệ thống tự tính.

            </div>


            <button
                class="btn primary"
                type="submit"
                style="margin-top:16px"
            >
                Import vào database
            </button>

        </form>

    </div>
    """

    return base_page(
        "Import Excel",
        body,
        session
    )

def legacy_import_excel_data(
    session,
    file_bytes,
    report_date,
    sheet_name="",
    mode="skip",
    filename=""
):
    if not is_chi_cuc_user(session):
        return False, "Không có quyền import dữ liệu."

    try:
        date.fromisoformat(report_date)
    except ValueError:
        return False, "Ngày báo cáo không hợp lệ."

    try:
        from openpyxl import load_workbook
    except Exception:
        return False, "Chưa cài openpyxl."

    try:
        wb = load_workbook(
            io.BytesIO(file_bytes),
            data_only=True
        )
    except Exception as e:
        return False, f"Không đọc được file Excel: {e}"


    if sheet_name:

        if sheet_name not in wb.sheetnames:
            return False, (
                f"Không tìm thấy sheet '{sheet_name}'. "
                f"Các sheet hiện có: {', '.join(wb.sheetnames)}"
            )

        ws = wb[sheet_name]

    else:
        ws = wb.active


    con = db_conn()

    inserted = 0
    updated = 0
    skipped = 0

    try:

        for row in range(1, ws.max_row + 1):

            # Cột B = tên xã/phường
            unit_name = excel_text(
                ws.cell(row=row, column=2).value
            )

            if not unit_name:
                continue


            unit_check = unit_name.casefold()


            # Bỏ dòng tổng cộng
            if "tổng cộng" in unit_check:
                continue


            # Chỉ lấy dòng đơn vị
            if not unit_check.startswith(
                (
                    "xã ",
                    "phường ",
                    "thị trấn "
                )
            ):
                continue

            official_unit = canonical_admin_unit(unit_name)
            if not official_unit:
                raise ValueError(f"Đơn vị chưa có trong danh mục chính thức: {unit_name}")
            unit_name, unit_code = official_unit


            record = {

                "report_date": report_date,

                "unit_name": unit_name,

                "created_by": session["user_id"],


                # D, E
                "area_land":
                    fnum(ws.cell(row=row, column=4).value),

                "area_tarp":
                    fnum(ws.cell(row=row, column=5).value),


                # G, H
                "harvest_land":
                    fnum(ws.cell(row=row, column=7).value),

                "harvest_tarp":
                    fnum(ws.cell(row=row, column=8).value),


                # J, K
                "sold_land":
                    fnum(ws.cell(row=row, column=10).value),

                "sold_tarp":
                    fnum(ws.cell(row=row, column=11).value),


                # M, N
                "remaining_land":
                    fnum(ws.cell(row=row, column=13).value),

                "remaining_tarp":
                    fnum(ws.cell(row=row, column=14).value),


                # P, Q
                "processed_fine":
                    fnum(ws.cell(row=row, column=16).value),

                "processed_iodized":
                    fnum(ws.cell(row=row, column=17).value),


                # R, S
                "households":
                    fint(ws.cell(row=row, column=18).value),

                "workers":
                    fint(ws.cell(row=row, column=19).value),


                # T, U
                "price_land":
                    excel_text(
                        ws.cell(row=row, column=20).value
                    ),

                "price_tarp":
                    excel_text(
                        ws.cell(row=row, column=21).value
                    ),


                # X, Y
                "damage_land":
                    fnum(ws.cell(row=row, column=24).value),

                "damage_tarp":
                    fnum(ws.cell(row=row, column=25).value),


                "status": STATUS_APPROVED,

                "note":
                    f"Import Excel: {filename} / {ws.title}",

                "created_at": now_text(),

                "updated_at": now_text(),

                "submitted_at": now_text(),

                "approved_at": now_text()
            }


            existing = con.execute(
                """
                SELECT id
                FROM records
                WHERE unit_name=?
                  AND report_date=?
                """,
                (
                    unit_name,
                    report_date
                )
            ).fetchone()


            # =====================
            # ĐÃ CÓ DỮ LIỆU
            # =====================
            if existing:

                if mode != "update":
                    existing_record = con.execute(
                        "SELECT * FROM records WHERE id=?",
                        (existing["id"],),
                    ).fetchone()
                    if existing_record and existing_record["status"] == STATUS_APPROVED:
                        sync_standard_salt_record(con, existing_record)
                    skipped += 1
                    continue


                record["id"] = existing["id"]


                con.execute(
                    """
                    UPDATE records
                    SET
                        area_land=:area_land,
                        area_tarp=:area_tarp,

                        harvest_land=:harvest_land,
                        harvest_tarp=:harvest_tarp,

                        sold_land=:sold_land,
                        sold_tarp=:sold_tarp,

                        remaining_land=:remaining_land,
                        remaining_tarp=:remaining_tarp,

                        processed_fine=:processed_fine,
                        processed_iodized=:processed_iodized,

                        households=:households,
                        workers=:workers,

                        price_land=:price_land,
                        price_tarp=:price_tarp,

                        damage_land=:damage_land,
                        damage_tarp=:damage_tarp,

                        status=:status,

                        updated_at=:updated_at,
                        approved_at=:approved_at

                    WHERE id=:id
                    """,
                    record
                )
                con.execute(
                    "UPDATE records SET reporting_mode='cumulative',ma_don_vi_hanh_chinh=? WHERE id=?",
                    (unit_code, existing["id"]),
                )


                add_audit(
                    con,
                    existing["id"],
                    session["user_id"],
                    "Import Excel",
                    f"Cập nhật dữ liệu từ {filename} / {ws.title}"
                )

                updated_record = con.execute(
                    "SELECT * FROM records WHERE id=?",
                    (existing["id"],),
                ).fetchone()
                sync_standard_salt_record(con, updated_record)

                updated += 1


            # =====================
            # DỮ LIỆU MỚI
            # =====================
            else:

                cur = con.execute(
                    """
                    INSERT INTO records
                    (
                        report_date,
                        unit_name,
                        created_by,

                        area_land,
                        area_tarp,

                        harvest_land,
                        harvest_tarp,

                        sold_land,
                        sold_tarp,

                        remaining_land,
                        remaining_tarp,

                        processed_fine,
                        processed_iodized,

                        households,
                        workers,

                        price_land,
                        price_tarp,

                        damage_land,
                        damage_tarp,

                        status,
                        note,

                        created_at,
                        updated_at,

                        submitted_at,
                        approved_at
                    )

                    VALUES
                    (
                        :report_date,
                        :unit_name,
                        :created_by,

                        :area_land,
                        :area_tarp,

                        :harvest_land,
                        :harvest_tarp,

                        :sold_land,
                        :sold_tarp,

                        :remaining_land,
                        :remaining_tarp,

                        :processed_fine,
                        :processed_iodized,

                        :households,
                        :workers,

                        :price_land,
                        :price_tarp,

                        :damage_land,
                        :damage_tarp,

                        :status,
                        :note,

                        :created_at,
                        :updated_at,

                        :submitted_at,
                        :approved_at
                    )
                    """,
                    record
                )
                con.execute(
                    "UPDATE records SET reporting_mode='cumulative',ma_don_vi_hanh_chinh=? WHERE id=?",
                    (unit_code, cur.lastrowid),
                )


                add_audit(
                    con,
                    cur.lastrowid,
                    session["user_id"],
                    "Import Excel",
                    f"Import từ {filename} / {ws.title}"
                )

                inserted_record = con.execute(
                    "SELECT * FROM records WHERE id=?",
                    (cur.lastrowid,),
                ).fetchone()
                sync_standard_salt_record(con, inserted_record)

                inserted += 1


        con.commit()


    except Exception as e:

        con.rollback()
        con.close()

        return False, f"Lỗi khi import: {e}"


    con.close()


    return True, (
        f"Import hoàn tất. "
        f"Thêm mới: {inserted}; "
        f"Cập nhật: {updated}; "
        f"Bỏ qua: {skipped}."
    )

def import_excel_page(session, query=""):
    if not is_chi_cuc_user(session):
        return None
    con = db_conn()
    batches = con.execute(
        """SELECT b.*,u.username FROM salt_import_batches b
           JOIN users u ON u.id=b.imported_by ORDER BY b.id DESC LIMIT 8"""
    ).fetchall()
    con.close()
    history = ''.join(
        f"<tr><td>{esc(b['week_code'])}</td><td>{esc(b['filename'])}</td>"
        f"<td>{esc(b['sheet_name'])}</td><td>{b['imported_rows']}</td>"
        f"<td>{b['updated_rows']}</td><td>{b['skipped_rows']}</td><td>{esc(b['username'])}</td></tr>"
        for b in batches
    ) or '<tr><td colspan="7" class="empty">Chưa có đợt import tuần.</td></tr>'
    upload = session.get("weekly_import_upload")
    if upload:
        sheet_names = upload.get("sheet_names", [])
        selected_sheet = upload.get("selected_sheet") or upload.get("recommended_sheet") or (sheet_names[0] if sheet_names else "")
        options = "".join(
            f'<option value="{esc(name)}"{" selected" if name == selected_sheet else ""}>{esc(name)}</option>'
            for name in sheet_names
        )
        try:
            detected = detect_weekly_period(upload["content"], upload["filename"], selected_sheet)
        except WeeklyImportError as exc:
            detected = {"report_date": None, "hint": str(exc), "period_detected": False}
        if detected.get("report_date"):
            detected_day = date.fromisoformat(detected["report_date"])
            period_html = f'''
              <div class="detected-period" role="status">
                <div><span class="detected-period-label">Kỳ báo cáo</span><strong>Tuần {detected_day.isocalendar().week} / {detected_day.isocalendar().year}</strong></div>
                <div><span class="detected-period-label">Số liệu lũy tiến đến</span><strong>{detected_day.strftime("%d/%m/%Y")}</strong></div>
                <small>Đã nhận diện từ {esc(detected.get("source", "workbook"))}; bạn vẫn có thể đổi sheet.</small>
              </div>'''
            date_field = ""
        else:
            period_html = f'<div class="notice info">{esc(detected.get("hint") or "Không tự xác định được ngày báo cáo từ file này.")}</div>'
            date_field = f'<div class="field"><label>Số liệu lũy tiến đến ngày</label><input type="date" name="report_date" value="{esc(upload.get("report_date", ""))}" required></div>'
        selection_form = f'''
      <form class="card js-loading-form" method="post" action="/import-excel/preview">
        {csrf_input(session)}
        <div class="weekly-step"><span>2</span><div><strong>Chọn sheet và kỳ báo cáo</strong><small>Hệ thống dùng đúng sheet bạn chọn để xem trước.</small></div></div>
        <div class="field sheet-picker"><label for="weekly-sheet">Sheet</label><select id="weekly-sheet" name="sheet_name" required>{options}</select></div>
        <div class="import-file-name"><span>File Excel</span><strong>{esc(upload["filename"])}</strong></div>
        {period_html}{date_field}
        <div class="notice info">Cột B phải là xã/phường chính thức. C/F là tổng diện tích và sản lượng; D/E, G/H là chi tiết nền đất và nền trải bạt.</div>
        <div class="actions"><button class="btn primary" type="submit" data-loading-text="Đang đọc và kiểm tra Excel...">Xem trước</button><a class="btn" href="/import-excel?reset=1">Hủy file này</a></div>
      </form>'''
        workflow = selection_form
        page_heading = '<h1>Chọn sheet báo cáo Diêm nghiệp</h1><div class="subtitle">Bước 2: chọn sheet, xem kỳ báo cáo đã nhận diện rồi kiểm tra dữ liệu.</div>'
    else:
        workflow = f'''
      <form class="card import-card js-loading-form" method="post" action="/import-excel" enctype="multipart/form-data">
        {csrf_input(session)}
        <div class="weekly-step"><span>1</span><div><strong>Chọn file Excel</strong><small>Hệ thống sẽ đọc danh sách sheet trước, chưa ghi database.</small></div></div>
        <div class="field"><label for="weekly-excel-file">File Excel (.xlsx)</label><input id="weekly-excel-file" type="file" name="excel_file" accept=".xlsx" required></div>
        <div class="actions"><button class="btn primary" type="submit" data-loading-text="Đang đọc workbook...">Tiếp tục</button></div>
      </form>'''
        page_heading = '<h1>Import báo cáo Diêm nghiệp</h1><div class="subtitle">Bước 1: chọn file Excel. Tên sheet và kỳ báo cáo sẽ được hệ thống nhận diện sau đó.</div>'
    body = f"""
    <div class="container weekly-page">
      {take_flash(session)}
      <div class="page-head"><div>{page_heading}</div>
        <a class="btn" href="/records">Tra cứu báo cáo Diêm nghiệp</a></div>
      {workflow}
      <section class="card"><h2 class="section-title">Lịch sử import gần đây</h2>
        <div class="table-wrap"><table class="summary-table"><thead><tr><th>Tuần</th><th>File</th><th>Sheet</th><th>Mới</th><th>Cập nhật</th><th>Bỏ qua</th><th>Người import</th></tr></thead><tbody>{history}</tbody></table></div>
      </section>
    </div>"""
    return base_page("Import báo cáo Diêm nghiệp", body, session, active_path="/import-excel")


def import_preview_page(session):
    if not is_chi_cuc_user(session):
        return None
    preview = session.get("weekly_import_preview")
    if not preview:
        set_flash(session, "err", "Chưa có dữ liệu xem trước. Hãy chọn file Excel.")
        return None
    rows_html = []
    for row in preview["rows"]:
        data = row["canonical"]
        messages = row["errors"] + row["warnings"]
        label = {"valid":"Hợp lệ", "warning":"Cảnh báo", "error":"Lỗi"}[row["status"]]
        rows_html.append(
            f'<tr class="preview-{row["status"]}"><td>{row["excel_row"]}</td>'
            f'<td><strong>{esc(data["unit_name"])}</strong><small>{esc(data.get("ma_don_vi_hanh_chinh") or "Chưa có mã")}</small></td>'
            f'<td>{fmt_num(data["dien_tich"],2)}</td><td>{fmt_num(data["san_luong"],2)}</td>'
            f'<td><span class="validation-badge {row["status"]}">{label}</span>'
            f'<div class="validation-message">{esc("; ".join(messages))}</div></td></tr>'
        )
    blocked = preview["error_rows"] > 0 or preview.get("period_blocked")
    period_day = date.fromisoformat(str(preview["report_date"]))
    period_label = f"Tuần {period_day.isocalendar().week} / {period_day.isocalendar().year}"
    period_warning_html = "".join(
        f'<div class="notice warn"><strong>Kiểm tra kỳ báo cáo:</strong> {esc(message)}</div>'
        for message in preview.get("period_warnings", [])
    )
    detected_note = (
        "Kỳ báo cáo được tự nhận diện từ workbook."
        if preview.get("period_detected")
        else "Kỳ báo cáo được nhập bổ sung vì workbook chưa đủ thông tin nhận diện."
    )
    confirm = f"""
      <form class="confirm-import" method="post" action="/import-excel/confirm">
        {csrf_input(session)}
        <div class="field"><label>Nếu xã/tuần đã tồn tại</label><select name="mode"><option value="skip">Giữ dữ liệu hiện tại</option><option value="update">Cập nhật bằng file này</option></select></div>
        <button class="btn primary" type="submit" {'disabled' if blocked else ''}>Xác nhận import</button>
    </form>""" if not blocked else '<div class="notice err">Cần sửa các dòng lỗi hoặc kiểm tra lại kỳ báo cáo rồi tải lại file.</div>'
    body = f"""
    <div class="container weekly-page">
      {take_flash(session)}
      <div class="page-head"><div><h1>Xem trước import Diêm nghiệp</h1>
        <div class="subtitle">Bước 3: kiểm tra {esc(preview['filename'])} · {esc(preview['sheet_name'])}</div></div>
        <a class="btn" href="/import-excel?reset=1">Chọn file khác</a></div>
      <div class="detected-period preview-period" role="status">
        <div><span class="detected-period-label">Sheet</span><strong>{esc(preview['sheet_name'])}</strong></div>
        <div><span class="detected-period-label">Kỳ báo cáo</span><strong>{esc(period_label)}</strong></div>
        <div><span class="detected-period-label">Số liệu lũy tiến đến</span><strong>{period_day.strftime('%d/%m/%Y')}</strong></div>
        <small>{detected_note}</small>
      </div>
      {period_warning_html}
      <div class="validation-summary">
        <div><strong>{len(preview['rows'])}</strong><span>Tổng số xã</span></div>
        <div class="valid"><strong>{preview['valid_rows']}</strong><span>Hợp lệ</span></div>
        <div class="warning"><strong>{preview['warning_rows']}</strong><span>Cảnh báo</span></div>
        <div class="error"><strong>{preview['error_rows']}</strong><span>Lỗi</span></div>
      </div>
      <section class="card"><div class="table-wrap"><table class="summary-table preview-table"><thead><tr><th>Dòng Excel</th><th>Đơn vị</th><th>Diện tích (ha)</th><th>Sản lượng (tấn)</th><th>Kết quả</th></tr></thead><tbody>{''.join(rows_html)}</tbody></table></div></section>
      {confirm}
    </div>"""
    return base_page("Xem trước import Diêm nghiệp", body, session, active_path="/import-excel")


def _weekly_excel_value(value):
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return fmt_num(value, 2)
    token = str(value)
    try:
        parsed = datetime.fromisoformat(token)
        return parsed.strftime("%d/%m/%Y")
    except ValueError:
        return token


_WEEKLY_LAND_COLUMNS = (4, 7, 10, 13, 24)
_WEEKLY_TARP_COLUMNS = (5, 8, 11, 14, 25)


def _weekly_raw_value(raw, column):
    return raw.get(str(column), {}).get("value")


def _weekly_numeric_value(raw, column):
    value = _weekly_raw_value(raw, column)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return 0.0


def _weekly_status_badge(label, kind="neutral"):
    return f'<span class="weekly-status {esc(kind)}">{esc(label)}</span>'


def _weekly_type_status(raw, detail_columns, price_column):
    """Return the current-period production state only.

    Price is intentionally not used to infer whether a foundation was in
    production during the selected reporting period.
    """
    has_activity = any(_weekly_numeric_value(raw, column) > 0 for column in detail_columns)
    if has_activity:
        return "active", "Có dữ liệu"
    return "inactive", "Không phát sinh kỳ này"


def _weekly_price_state(raw, price_column, detail_columns):
    value = _weekly_raw_value(raw, price_column)
    type_kind, _ = _weekly_type_status(raw, detail_columns, price_column)
    if value not in (None, "", "-") and type_kind == "inactive":
        return "reference", "Giá tham chiếu"
    if value in (None, "", "-") and type_kind == "active":
        return "pending", "Chưa nhập"
    if value not in (None, "", "-"):
        return "active", "Có giá"
    return "inactive", "Không phát sinh"


def _weekly_type_badge(raw, label, detail_columns, price_column):
    kind, state = _weekly_type_status(raw, detail_columns, price_column)
    price_kind, _ = _weekly_price_state(raw, price_column, detail_columns)
    badge_kind = "review" if price_kind == "reference" else kind
    return (
        f'<span class="weekly-type-badge {badge_kind}">'
        f'<strong>{esc(label)}</strong><small>{esc(state)}</small></span>'
    )


def _weekly_price_display(raw, price_column, detail_columns):
    value = _weekly_raw_value(raw, price_column)
    if value not in (None, "", "-"):
        price_kind, price_label = _weekly_price_state(raw, price_column, detail_columns)
        extra = f' title="{esc(price_label)}: giá được giữ độc lập với số liệu sản xuất kỳ này"' if price_kind == "reference" else ""
        return f'<span class="weekly-price-value"{extra}>{esc(_weekly_excel_value(value))}</span>'
    type_kind, _ = _weekly_type_status(raw, detail_columns, price_column)
    if type_kind == "active":
        return _weekly_status_badge("Chưa nhập", "pending")
    return _weekly_status_badge("Không phát sinh", "inactive")


def _weekly_metric_display(raw, column, active):
    value = _weekly_raw_value(raw, column)
    if value in (None, "", "-"):
        return _weekly_status_badge("Chưa nhập", "pending") if active else _weekly_status_badge("—", "inactive")
    return esc(_weekly_excel_value(value))


def _weekly_type_detail(raw, label, detail_columns, price_column):
    kind, state = _weekly_type_status(raw, detail_columns, price_column)
    price_kind, _ = _weekly_price_state(raw, price_column, detail_columns)
    card_kind = "review" if price_kind == "reference" else kind
    if kind == "inactive":
        return (
            f'<div class="weekly-breakdown-card {card_kind}">'
            f'<div class="weekly-breakdown-heading"><strong>{esc(label)}</strong>'
            f'{_weekly_status_badge(state, card_kind)}</div>'
            f'<div class="weekly-breakdown-empty"><span>Diện tích, sản lượng, tiêu thụ và còn lại</span>'
            f'{_weekly_status_badge("—", "inactive")}</div>'
            f'<div class="weekly-breakdown-price-note">Giá: {_weekly_price_display(raw, price_column, detail_columns)}</div>'
            '</div>'
        )
    labels = (("Diện tích", detail_columns[0]), ("Thu hoạch", detail_columns[1]),
              ("Đã tiêu thụ", detail_columns[2]), ("Còn lại", detail_columns[3]))
    metrics = "".join(
        f'<div><span>{esc(metric_label)}</span><strong>{_weekly_metric_display(raw, column, kind in ("active", "review"))}</strong></div>'
        for metric_label, column in labels
    )
    return (
        f'<div class="weekly-breakdown-card {kind}">'
        f'<div class="weekly-breakdown-heading"><strong>{esc(label)}</strong>'
        f'{_weekly_status_badge(state, "ok")}</div>'
        f'<div class="weekly-breakdown-metrics">{metrics}'
        f'<div><span>Giá</span><strong>{_weekly_price_display(raw, price_column, detail_columns)}</strong></div>'
        '</div></div>'
    )


def _weekly_detail_dialog(raw, dialog_id):
    unit_name = _weekly_raw_value(raw, 2) or "Xã/phường"
    detail_cards = (
        _weekly_type_detail(raw, "Muối đất", _WEEKLY_LAND_COLUMNS, 20)
        + _weekly_type_detail(raw, "Muối trải bạt", _WEEKLY_TARP_COLUMNS, 21)
    )
    review_types = []
    for label, detail_columns, price_column in (
        ("muối đất", _WEEKLY_LAND_COLUMNS, 20),
        ("muối trải bạt", _WEEKLY_TARP_COLUMNS, 21),
    ):
        kind, _ = _weekly_type_status(raw, detail_columns, price_column)
        price_kind, _ = _weekly_price_state(raw, price_column, detail_columns)
        if kind == "inactive" and price_kind == "reference":
            review_types.append(label)
    review_note = ""
    if review_types:
        review_note = (
            '<div class="weekly-dialog-warning">'
            '<strong>Giá được giữ độc lập</strong>'
            f'<p>Đã nhập giá {esc(" và ".join(review_types))}, nhưng kỳ này không có số liệu sản xuất tương ứng. '
            'Các ô chỉ tiêu được để “—”; giá vẫn được giữ lại để tham khảo và không dùng để suy ra phương pháp sản xuất.</p>'
            '</div>'
        )
    return (
        f'<dialog id="{esc(dialog_id)}" class="weekly-dialog">'
        '<div class="weekly-dialog-content">'
        '<div class="weekly-dialog-header">'
        f'<div><span class="weekly-dialog-eyebrow">Chi tiết loại hình</span><h3>{esc(unit_name)}</h3></div>'
        '<button type="button" class="weekly-dialog-close" data-weekly-dialog-close aria-label="Đóng">×</button>'
        '</div>'
        f'{review_note}<div class="weekly-breakdown">{detail_cards}</div>'
        '<div class="weekly-dialog-footer"><span>Ô trống được diễn giải theo trạng thái, không hiển thị nguyên dấu “-” của Excel.</span>'
        '<button type="button" class="btn" data-weekly-dialog-close>Đóng</button></div>'
        '</div></dialog>'
    )


def _weekly_total(raw_rows, column):
    values = []
    for raw in raw_rows:
        value = raw.get(str(column), {}).get("value")
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            values.append(float(value))
    if column in (20, 21):
        non_zero = [value for value in values if value]
        return sum(non_zero) / len(non_zero) if non_zero else None
    return sum(values)


def weekly_records_page(session, query):
    """Browse one imported workbook sheet as one report."""
    params = parse_qs(query)
    con = db_conn()
    batches = con.execute(
        """SELECT b.*,u.username FROM salt_import_batches b
           JOIN users u ON u.id=b.imported_by
           ORDER BY b.report_date DESC,b.id DESC"""
    ).fetchall()
    requested_batch = params.get("batch", [""])[0].strip()
    requested_unit = params.get("unit", [""])[0].strip()
    view = params.get("view", ["summary"])[0].strip().lower()
    if view not in ("summary", "full"):
        view = "summary"
    selected = next((batch for batch in batches if str(batch["id"]) == requested_batch), None)
    if not selected:
        requested_week = params.get("week", [""])[0].strip()
        selected = next((batch for batch in batches if batch["week_code"] == requested_week), None)
    if not selected and batches:
        selected = batches[0]
    # Do not carry a unit filter from a different sheet into this report.
    if selected and requested_unit:
        has_unit = con.execute(
            """SELECT 1 FROM salt_weekly_import_rows
                WHERE batch_id=? AND ma_don_vi_hanh_chinh=? LIMIT 1""",
            (selected["id"], requested_unit),
        ).fetchone()
        if not has_unit:
            requested_unit = ""

    rows = []
    if selected:
        sql = """SELECT excel_row,unit_name_raw,ma_don_vi_hanh_chinh,
                        validation_status,validation_messages_json,raw_data_json
                 FROM salt_weekly_import_rows WHERE batch_id=?"""
        args = [selected["id"]]
        if session["role"] == ROLE_UNIT:
            official = canonical_admin_unit(session["unit_name"])
            sql += " AND ma_don_vi_hanh_chinh=?"
            args.append(official[1] if official else "")
        elif requested_unit:
            sql += " AND ma_don_vi_hanh_chinh=?"
            args.append(requested_unit)
        sql += " ORDER BY excel_row"
        rows = con.execute(sql, args).fetchall()
    con.close()

    options = ''.join(
        f'<option value="{batch["id"]}" {"selected" if selected and batch["id"] == selected["id"] else ""}>'
        f'{esc(batch["sheet_name"])} · {esc(batch["week_code"])} · {esc(batch["filename"])}</option>'
        for batch in batches
    )
    # Unit options belong to the selected report, not to the whole official
    # administrative catalog.  The catalog is still used for the display name
    # and code, but a unit only appears when this sheet actually contains a
    # mapped row for it.
    unit_rows = []
    if selected:
        unit_con = db_conn()
        try:
            unit_rows = [dict(row) for row in unit_con.execute(
                """SELECT DISTINCT r.ma_don_vi_hanh_chinh AS code,
                           COALESCE(u.TenDonVi, r.unit_name_raw) AS name
                    FROM salt_weekly_import_rows r
                    LEFT JOIN DM_DonViHanhChinh u
                      ON u.Ma_DonViHanhChinh=r.ma_don_vi_hanh_chinh
                   WHERE r.batch_id=?
                     AND r.ma_don_vi_hanh_chinh IS NOT NULL
                     AND upper(r.ma_don_vi_hanh_chinh) NOT LIKE 'TMP%'
                   ORDER BY name, code""", (selected["id"],)
            ).fetchall()]
        finally:
            unit_con.close()
    if session["role"] == ROLE_UNIT:
        official = canonical_admin_unit(session["unit_name"])
        unit_rows = [u for u in unit_rows if official and u["code"] == official[1]]
    available_codes = {unit["code"] for unit in unit_rows}
    # A batch change should not leave a previous unit filter pointing at a
    # unit that is absent from the newly selected sheet.
    if requested_unit and requested_unit not in available_codes:
        requested_unit = ""
    unit_options = '<option value="">Tất cả xã/phường</option>' + ''.join(
        f'<option value="{esc(unit["code"])}" {"selected" if unit["code"] == requested_unit else ""}>{esc(unit["name"])}</option>'
        for unit in unit_rows
    )
    if not selected:
        report = '<section class="card empty">Chưa có sheet báo cáo tuần nào được import.</section>'
    else:
        raw_rows = [json.loads(row["raw_data_json"]) for row in rows]
        def raw_value(raw, column):
            value = _weekly_raw_value(raw, column)
            if value in (None, "", "-"):
                return _weekly_status_badge("Chưa nhập", "pending")
            return esc(_weekly_excel_value(value))

        summary_row_parts = []
        dialog_parts = []
        for row_index, raw in enumerate(raw_rows):
            dialog_id = f"weekly-detail-dialog-{row_index}"
            summary_row_parts.append(
                f'<tr><td><strong>{raw_value(raw, 2)}</strong></td>'
                f'<td><div class="weekly-type-cell">'
                f'{_weekly_type_badge(raw, "Muối đất", _WEEKLY_LAND_COLUMNS, 20)}'
                f'{_weekly_type_badge(raw, "Trải bạt", _WEEKLY_TARP_COLUMNS, 21)}'
                '</div></td>'
                f'<td>{raw_value(raw, 3)}</td>'
                f'<td>{raw_value(raw, 6)}</td><td>{raw_value(raw, 9)}</td>'
                f'<td>{raw_value(raw, 12)}</td><td>{raw_value(raw, 18)}</td>'
                f'<td>{raw_value(raw, 19)}</td>'
                f'<td>{_weekly_price_display(raw, 20, _WEEKLY_LAND_COLUMNS)}</td>'
                f'<td>{_weekly_price_display(raw, 21, _WEEKLY_TARP_COLUMNS)}</td>'
                f'<td><button type="button" class="weekly-detail-trigger" '
                f'data-weekly-dialog="{esc(dialog_id)}">Chi tiết</button></td></tr>'
            )
            dialog_parts.append(_weekly_detail_dialog(raw, dialog_id))
        summary_rows = "".join(summary_row_parts) or '<tr><td colspan="11" class="empty">Không có xã/phường phù hợp.</td></tr>'
        detail_dialogs = "".join(dialog_parts)
        summary_total = (
            '<tr class="sheet-total"><th>TỔNG CỘNG</th>'
            '<th><span class="weekly-total-label">Toàn bộ loại hình</span></th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 3)))}</th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 6)))}</th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 9)))}</th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 12)))}</th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 18)))}</th>'
            f'<th>{esc(_weekly_excel_value(_weekly_total(raw_rows, 19)))}</th>'
            f'<th>{_weekly_status_badge("Không tổng hợp", "neutral")}</th>'
            f'<th>{_weekly_status_badge("Không tổng hợp", "neutral")}</th>'
            '<th>—</th></tr>'
        )
        summary_table = f'''
          <div class="table-wrap"><table class="summary-table weekly-summary-table">
            <thead><tr><th>Xã/phường</th><th>Loại hình</th><th>Diện tích tổng (ha)</th><th>Sản lượng thu hoạch (tấn)</th><th>Đã tiêu thụ (tấn)</th><th>Còn lại (tấn)</th><th>Số hộ</th><th>Lao động</th><th>Giá muối đất</th><th>Giá muối trải bạt</th><th>Chi tiết</th></tr></thead>
            <tbody>{summary_total}{summary_rows}</tbody>
          </table></div>{detail_dialogs}
          <div class="weekly-summary-legend" aria-label="Quy ước hiển thị">
            <span><i class="legend-dot active"></i>Có dữ liệu</span>
            <span><i class="legend-dot pending"></i>Chưa nhập</span>
            <span><i class="legend-dot inactive"></i>Không phát sinh</span>
            <span><i class="legend-dot reference"></i>Giá tham chiếu</span>
            <span>Giá không cộng dồn ở dòng tổng; xem theo từng xã/phường.</span>
          </div>'''
        detail_rows = []
        for raw in raw_rows:
            cells = ''.join(f'<td>{esc(_weekly_excel_value(raw.get(str(column), {}).get("value")))}</td>' for column in range(1, 29))
            detail_rows.append(f'<tr>{cells}</tr>')
        total_cells = ['<th></th>', '<th>TỔNG CỘNG</th>']
        for column in range(3, 29):
            if column == 22:
                area, production = _weekly_total(raw_rows, 3), _weekly_total(raw_rows, 6)
                value = production / area if area else None
            elif column in tuple(range(3, 22)) + tuple(range(23, 27)):
                value = _weekly_total(raw_rows, column)
            else:
                value = None
            total_cells.append(f'<th>{esc(_weekly_excel_value(value))}</th>')
        report_date = selected["report_date"]
        try:
            report_date = date.fromisoformat(str(report_date)).strftime("%d/%m/%Y")
        except ValueError:
            report_date = str(report_date)
        export_query = urlencode({k: v for k, v in {"batch": str(selected["id"]), "unit": requested_unit}.items() if v})
        view_query = lambda selected_view: urlencode({k: v for k, v in {"batch": str(selected["id"]), "unit": requested_unit, "view": selected_view}.items() if v})
        view_switch = (
            '<div class="view-switch" role="tablist" aria-label="Chế độ xem báo cáo">'
            f'<a class="view-switch-link{" active" if view == "summary" else ""}" href="/records?{view_query("summary")}" role="tab"{" aria-selected=\"true\"" if view == "summary" else ""}>Bảng tổng hợp</a>'
            f'<a class="view-switch-link{" active" if view == "full" else ""}" href="/records?{view_query("full")}" role="tab"{" aria-selected=\"true\"" if view == "full" else ""}>Xem bảng Excel đầy đủ</a>'
            '</div>'
        )
        full_table = f'''
          <div class="excel-sheet-wrap"><table class="excel-sheet">
            <caption>Báo cáo tình hình sản xuất, chế biến, tiêu thụ niên vụ muối — lũy tiến đến ngày {esc(report_date)}</caption>
            <thead>
              <tr><th rowspan="2">TT</th><th rowspan="2">Thành phố/huyện/thị xã</th><th colspan="3">Diện tích sản xuất muối (ha)</th><th colspan="3">Sản lượng muối thu hoạch (tấn)</th><th colspan="3">Sản lượng muối tiêu thụ (tấn)</th><th colspan="3">Sản lượng còn lại (tấn)</th><th colspan="3">Sản lượng muối chế biến (tấn)</th><th rowspan="2">Số hộ làm muối</th><th rowspan="2">Số lao động</th><th colspan="2">Giá bán (đồng/kg)</th><th rowspan="2">Năng suất BQ</th><th colspan="3">Thiệt hại do mưa trái mùa</th><th rowspan="2">Diện tích mất trắng</th><th rowspan="2">Ghi chú</th><th rowspan="2">Thời gian kết thúc niên vụ</th></tr>
              <tr><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th><th>Cộng</th><th>Muối tinh</th><th>Muối I-ốt</th><th>Muối đất</th><th>Muối trải bạt</th><th>Cộng</th><th>Muối đất</th><th>Muối trải bạt</th></tr>
            </thead>
            <tbody><tr class="sheet-total">{''.join(total_cells)}</tr>{''.join(detail_rows)}</tbody>
          </table></div>'''
        report_table = summary_table if view == "summary" else full_table
        report = f"""
        <section class="sheet-report">
          <div class="sheet-meta">
            <div><span>Sheet báo cáo</span><strong>{esc(selected['sheet_name'])}</strong></div>
            <div><span>Số liệu đến ngày</span><strong>{esc(report_date)}</strong></div>
            <div><span>Tuần</span><strong>{esc(selected['week_code'])}</strong></div>
            <div><span>Số xã/phường</span><strong>{len(rows)}</strong></div>
          </div>
          {view_switch}
          {report_table}
          <div class="actions" style="margin-top:12px"><a class="btn ok" href="/records/export.xlsx?{export_query}">Xuất Excel dữ liệu đang xem</a></div>
          <p class="sheet-footnote">Nguồn: {esc(selected['filename'])} · Import bởi {esc(selected['username'])} lúc {esc(selected['imported_at'])}. Mỗi sheet là một báo cáo; mỗi xã/phường là một dòng trong báo cáo.</p>
        </section>"""

    body = f"""
    <div class="container weekly-page">
      {take_flash(session)}
      <div class="page-head"><div><h1>Tra cứu báo cáo Diêm nghiệp</h1><div class="subtitle">Tra cứu số liệu theo kỳ báo cáo và xã/phường.</div></div>
        {'<a class="btn primary" href="/import-excel">Import Excel</a>' if is_chi_cuc_user(session) else ''}</div>
      <form class="card weekly-filter" method="get"><div class="field sheet-picker"><label>Sheet đã import</label><select name="batch">{options}</select></div><div class="field"><label>Xã/phường</label><select name="unit">{unit_options}</select></div><button class="btn primary">Mở bảng</button><div class="weekly-progress"><strong>{len(batches)}</strong><span>sheet báo cáo</span></div></form>
      {report}
    </div>"""
    return base_page("Tra cứu báo cáo tuần", body, session, active_path="/records")


def export_weekly_xlsx(session, query=""):
    """Export raw rows of the selected imported sheet and current unit filter."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return None, None
    params = parse_qs(query or "")
    batch_id = params.get("batch", [""])[0].strip()
    requested_unit = params.get("unit", [""])[0].strip()
    try:
        batch_id = int(batch_id)
    except (TypeError, ValueError):
        return None, None
    con = db_conn()
    try:
        batch = con.execute("SELECT * FROM salt_import_batches WHERE id=?", (batch_id,)).fetchone()
        if not batch:
            return None, None
        sql = "SELECT raw_data_json FROM salt_weekly_import_rows WHERE batch_id=?"
        args = [batch_id]
        if session["role"] == ROLE_UNIT:
            official = canonical_admin_unit(session["unit_name"])
            sql += " AND ma_don_vi_hanh_chinh=?"
            args.append(official[1] if official else "")
        elif requested_unit:
            sql += " AND ma_don_vi_hanh_chinh=?"
            args.append(requested_unit)
        sql += " ORDER BY excel_row"
        raw_rows = [json.loads(row["raw_data_json"]) for row in con.execute(sql, args).fetchall()]
    finally:
        con.close()
    headers = ["TT", "Xã/phường", "Diện tích cộng (ha)", "Muối đất (ha)", "Muối trải bạt (ha)",
               "Sản lượng cộng (tấn)", "Muối đất (tấn)", "Muối trải bạt (tấn)", "Tiêu thụ cộng (tấn)",
               "Tiêu thụ muối đất (tấn)", "Tiêu thụ trải bạt (tấn)", "Còn lại cộng (tấn)",
               "Còn lại muối đất (tấn)", "Còn lại trải bạt (tấn)", "Chế biến cộng (tấn)",
               "Muối tinh (tấn)", "Muối I-ốt (tấn)", "Số hộ", "Lao động", "Giá muối đất",
               "Giá trải bạt", "Năng suất BQ", "Thiệt hại cộng", "Thiệt hại muối đất",
               "Thiệt hại trải bạt", "Diện tích mất trắng", "Ghi chú", "Kết thúc niên vụ"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Bao cao tuan"
    ws.append([f"BÁO CÁO TUẦN — {batch['sheet_name']}"])
    ws.append([f"Nguồn: {batch['filename']} · Số liệu đến ngày: {batch['report_date']} · Thời điểm xuất: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"])
    ws.append([f"Điều kiện lọc: Xã/phường = {requested_unit or 'Tất cả trong phạm vi được phép'}"])
    ws.append(headers)
    for index, raw in enumerate(raw_rows, 1):
        values = [raw.get(str(column), {}).get("value", "") for column in range(1, 29)]
        values[0] = index
        ws.append(values)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    thin = Side(style="thin", color="B8C2CC")
    ws[1][0].font = Font(bold=True, size=14, color="203149")
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="BE3A24")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 18 if index != 2 else 24
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:AB{max(ws.max_row, 4)}"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def change_password_page(session, error=""):
    notice = f'<div class="notice err">{esc(error)}</div>' if error else ''
    body = f"""<div class="container"><div class="page-head"><div><h1>Đổi mật khẩu</h1></div></div>{notice}<form class="card" method="post" action="/change-password">{csrf_input(session)}<div class="grid"><div class="field"><label>Mật khẩu hiện tại</label><input type="password" name="old" required></div><div class="field"><label>Mật khẩu mới</label><input type="password" name="new" minlength="8" required></div><div class="field"><label>Nhập lại mật khẩu mới</label><input type="password" name="new2" minlength="8" required></div></div><button class="btn primary" style="margin-top:12px">Đổi mật khẩu</button></form></div>"""
    return base_page("Đổi mật khẩu", body, session)


def save_record(session, data, rid=None):
    unit_name = session["unit_name"] if session["role"] == ROLE_UNIT else data.get("unit_name", "").strip()
    report_date = data.get("report_date", "").strip()
    if not unit_name or not report_date:
        return False, "Vui lòng chọn đơn vị và ngày/kỳ chốt số liệu."
    official_unit = canonical_admin_unit(unit_name)
    if not official_unit:
        return False, f"Đơn vị {unit_name} chưa có trong danh mục hành chính chính thức."
    unit_code = official_unit[1] if official_unit else None
    unit_name = official_unit[0] if official_unit else unit_name
    try:
        date.fromisoformat(report_date)
    except ValueError:
        return False, "Ngày báo cáo không hợp lệ."
    values = {k: fnum(data.get(k,0)) for k in NUMERIC_FIELDS}
    values["households"] = fint(data.get("households",0))
    values["workers"] = fint(data.get("workers",0))
    price_land = data.get("price_land", "").strip()
    price_tarp = data.get("price_tarp", "").strip()
    note = data.get("note", "").strip()
    con = db_conn()
    try:
        if rid is None:
            cur = con.execute(
                """INSERT INTO records(report_date,unit_name,created_by,area_land,area_tarp,harvest_land,harvest_tarp,sold_land,sold_tarp,remaining_land,remaining_tarp,processed_fine,processed_iodized,households,workers,price_land,price_tarp,damage_land,damage_tarp,status,note,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'draft',?,?,?)""",
                (report_date, unit_name, session["user_id"], values["area_land"], values["area_tarp"], values["harvest_land"], values["harvest_tarp"], values["sold_land"], values["sold_tarp"], values["remaining_land"], values["remaining_tarp"], values["processed_fine"], values["processed_iodized"], values["households"], values["workers"], price_land, price_tarp, values["damage_land"], values["damage_tarp"], note, now_text(), now_text())
            )
            rid = cur.lastrowid
            con.execute(
                "UPDATE records SET reporting_mode='cumulative',ma_don_vi_hanh_chinh=? WHERE id=?",
                (unit_code, rid),
            )
            add_audit(con, rid, session["user_id"], "Tạo báo cáo", "Lưu bản nháp")
        else:
            r = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
            if not r or not can_access_record(session, r):
                return False, "Không có quyền sửa báo cáo này."
            if session["role"] == ROLE_UNIT and r["status"] not in (STATUS_DRAFT, STATUS_RETURNED):
                return False, "Báo cáo đã gửi/duyệt nên đơn vị không thể sửa."
            con.execute(
                """UPDATE records SET report_date=?,unit_name=?,area_land=?,area_tarp=?,harvest_land=?,harvest_tarp=?,sold_land=?,sold_tarp=?,remaining_land=?,remaining_tarp=?,processed_fine=?,processed_iodized=?,households=?,workers=?,price_land=?,price_tarp=?,damage_land=?,damage_tarp=?,note=?,updated_at=? WHERE id=?""",
                (report_date, unit_name, values["area_land"], values["area_tarp"], values["harvest_land"], values["harvest_tarp"], values["sold_land"], values["sold_tarp"], values["remaining_land"], values["remaining_tarp"], values["processed_fine"], values["processed_iodized"], values["households"], values["workers"], price_land, price_tarp, values["damage_land"], values["damage_tarp"], note, now_text(), rid)
            )
            con.execute(
                "UPDATE records SET reporting_mode='cumulative',ma_don_vi_hanh_chinh=? WHERE id=?",
                (unit_code, rid),
            )
            add_audit(con, rid, session["user_id"], "Cập nhật báo cáo", "Chỉnh sửa số liệu")
        con.commit()
        return True, rid
    except INTEGRITY_ERRORS:
        con.rollback()
        return False, f"Đơn vị {unit_name} đã có báo cáo ngày {report_date}. Hãy mở báo cáo đó để cập nhật."
    finally:
        con.close()


def export_xlsx(session, filters):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return None, None
    con = db_conn(); records = fetch_records(con, session, filters); con.close()
    wb = Workbook(); ws = wb.active; ws.title = "Tong hop"
    headers1 = ["TT", "Đơn vị", "Diện tích sản xuất muối (ha)", "", "", "Sản lượng muối thu hoạch (tấn)", "", "", "Sản lượng muối tiêu thụ (tấn)", "", "", "Sản lượng còn lại (tấn)", "", "", "Sản lượng muối chế biến (tấn)", "", "", "Số hộ làm muối (hộ)", "Số lao động làm muối (người)", "Giá bán (đồng/kg)", "", "Năng suất BQ (tấn/ha)", "Thiệt hại do mưa trái mùa (tấn)", "", "", "Kỳ báo cáo", "Trạng thái"]
    headers2 = ["", "", "Cộng", "Muối đất", "Muối trải bạt", "Cộng", "Muối đất", "Muối trải bạt", "Cộng", "Muối đất", "Muối trải bạt", "Cộng", "Muối đất", "Muối trải bạt", "Cộng", "Muối tinh", "Muối I-ốt", "", "", "Muối đất", "Muối trải bạt", "", "Cộng", "Muối đất", "Muối trải bạt", "", ""]
    ws.append(headers1); ws.append(headers2)
    merges = [(1,3,1,5),(1,6,1,8),(1,9,1,11),(1,12,1,14),(1,15,1,17),(1,20,1,21),(1,23,1,25)]
    for a,b,c,d in merges: ws.merge_cells(start_row=a,start_column=b,end_row=c,end_column=d)
    for col in [1,2,18,19,22,26,27]: ws.merge_cells(start_row=1,start_column=col,end_row=2,end_column=col)
    thin = Side(style="thin", color="B8C2CC")
    for row in ws.iter_rows(min_row=1,max_row=2,min_col=1,max_col=27):
        for cell in row:
            cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True); cell.fill = PatternFill("solid", fgColor="D9EAF7"); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    sums = [0.0]*23
    for i,r in enumerate(records, start=1):
        t = record_totals(r)
        row = [i,r["unit_name"],t["area_total"],r["area_land"],r["area_tarp"],t["harvest_total"],r["harvest_land"],r["harvest_tarp"],t["sold_total"],r["sold_land"],r["sold_tarp"],t["remaining_total"],r["remaining_land"],r["remaining_tarp"],t["processed_total"],r["processed_fine"],r["processed_iodized"],r["households"],r["workers"],r["price_land"],r["price_tarp"],t["avg_yield"],t["damage_total"],r["damage_land"],r["damage_tarp"],r["report_date"],STATUS_LABELS[r["status"]]]
        ws.append(row)
    ws.freeze_panes = "C3"; ws.auto_filter.ref = f"A2:AA{max(2,ws.max_row)}"
    widths = {1:6,2:25,18:15,19:16,20:18,21:20,22:16,26:13,27:18}
    for c in range(1,28): ws.column_dimensions[get_column_letter(c)].width = widths.get(c,14)
    for row in ws.iter_rows(min_row=3):
        for cell in row: cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_ocop_xlsx(session, query=""):
    """Export the filtered OCOP catalogue without technical identifiers."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return None, None
    params = parse_qs(query or "")
    filters = {key: params.get(key, [""])[0].strip() for key in ("q", "unit", "group", "star")}
    con = db_conn()
    try:
        import ocop_services
        records = ocop_services.export_products(con, session, filters)
    finally:
        con.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "OCOP"
    headers = ["Xã/phường", "Chủ thể", "Loại hình chủ thể", "Người đại diện", "Điện thoại",
               "Tên sản phẩm", "Nhóm sản phẩm", "Hạng sao hiện tại", "Ngày công nhận gần nhất",
               "Số quyết định", "Cơ quan ban hành", "Ngày hết hạn"]
    ws.append(["DANH MỤC SẢN PHẨM OCOP"])
    ws.append([f"Thời điểm xuất: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"])
    filter_labels = {"q": "Từ khóa", "unit": "Xã/phường", "group": "Nhóm sản phẩm", "star": "Hạng sao"}
    criteria = "; ".join(f"{filter_labels[k]}: {filters[k]}" for k in filters if filters[k]) or "Không lọc"
    ws.append([f"Điều kiện lọc: {criteria}"])
    ws.append(headers)
    for row in records:
        rank = row.get("recognition_star") or row.get("current_star")
        ws.append([
            row.get("unit_name") or "", row.get("entity_name") or "", row.get("facility_type") or "",
            row.get("representative_name") or "", row.get("phone") or "", row.get("name") or "",
            row.get("product_group") or "", f"{rank} sao" if rank else "", row.get("latest_recognition_date") or "",
            row.get("latest_decision_number") or "", row.get("latest_decision_authority") or "",
            row.get("latest_expiry_date") or "",
        ])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    thin = Side(style="thin", color="B8C2CC")
    for cell in ws[1] + ws[2] + ws[3]:
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws[1][0].font = Font(bold=True, size=14, color="203149")
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="BE3A24")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [18, 28, 18, 22, 16, 28, 20, 15, 20, 18, 25, 18]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{max(ws.max_row, 4)}"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def ocop_access_page(con, session):
    import ocop_services as svc
    if svc.get_scope(con, session) is not None:
        raise svc.OcopError("Chỉ Chi cục được phân địa bàn OCOP.", 403)
    rows = con.execute("""SELECT u.id,u.username,u.unit_name,d.TenDonVi,d.Ma_DonViHanhChinh
        FROM users u LEFT JOIN user_admin_units m ON m.user_id=u.id
        LEFT JOIN DM_DonViHanhChinh d ON d.Ma_DonViHanhChinh=m.ma_don_vi_hanh_chinh
        WHERE u.role='unit' AND u.active=1 ORDER BY u.unit_name,u.username""").fetchall()
    units = svc.unit_options(con, session)
    forms = []
    for row in rows:
        options = '<option value="">Chọn địa bàn chính thức</option>'
        for unit in units:
            selected = ' selected' if unit['code'] == row['Ma_DonViHanhChinh'] else ''
            options += f'<option value="{esc(unit["code"])}"{selected}>{esc(unit["name"])} ({esc(unit["code"])})</option>'
        forms.append(f'''<tr><td>{esc(row['username'])}</td><td>{esc(row['unit_name'])}</td>
          <td><form method="post" action="/ocop/access" class="actions">{csrf_input(session)}
          <input type="hidden" name="user_id" value="{row['id']}">
          <label class="field">Địa bàn OCOP<select name="unit_code" required>{options}</select></label>
          <button class="btn primary" type="submit">Lưu địa bàn</button></form></td></tr>''')
    return "Phân địa bàn OCOP", f'''<div class="container ocop-page">{take_flash(session)}
      <div class="page-head"><h1>Phân địa bàn OCOP</h1><a class="btn" href="/users">Tài khoản</a></div>
      <div class="notice info">Quyền OCOP dùng mã hành chính. Tài khoản chưa được gán địa bàn chưa thể truy cập hồ sơ OCOP.
      Việc gán địa bàn cập nhật tên đơn vị của tài khoản và phạm vi chung cho hệ thống; dữ liệu báo cáo cũ được giữ nguyên.</div>
      <div class="card table-wrap"><table class="summary-table"><thead><tr><th>Tài khoản</th><th>Tên đơn vị hiện tại</th>
      <th>Địa bàn được phép quản lý</th></tr></thead><tbody>{''.join(forms)}</tbody></table></div></div>'''


def save_ocop_access(con, session, data):
    import ocop_services as svc
    if not can_manage_users(session):
        raise svc.OcopError("Chỉ Chi cục được phân địa bàn OCOP.", 403)
    try:
        uid = int(data.get("user_id", ""))
    except (ValueError, TypeError):
        raise svc.OcopError("Tài khoản không hợp lệ.")
    code = str(data.get("unit_code", "")).strip()
    with con:
        con.execute("BEGIN")
        current = get_user(con, session["user_id"])
        if not current or not current["active"] or not can_manage_users(dict(current)):
            raise svc.OcopError("Chỉ quản trị được phân địa bàn OCOP.", 403)
        svc.get_scope(con, session)
        user = con.execute("SELECT id FROM users WHERE id=? AND active=1 AND role='unit'", (uid,)).fetchone()
        if not user:
            raise svc.OcopError("Chọn tài khoản đang hoạt động và địa bàn chính thức.")
        try:
            _, code = admin_units.account_unit(con,ROLE_UNIT,{"unit_code":code})
        except admin_units.CatalogError as exc:
            raise svc.OcopError(str(exc),exc.status) from exc
        admin_units.assign_account(con,session['user_id'],uid,code)
    invalidate_user_sessions(uid)


class Handler(BaseHTTPRequestHandler):
    server_version = "SaltData/1.0"

    def handle_admin_units(self, path, query, session, data=None):
        if path.rstrip("/") != "/admin-units" and not path.startswith("/admin-units/"):
            return False
        con = db_conn()
        parts = [p for p in path.split("/") if p]
        code = parts[1] if len(parts) == 3 else None
        try:
            admin_units.require_admin(con,session)
            if data is None:
                if len(parts) == 1:
                    body = admin_unit_pages.listing(con,session,query,csrf_input(session),take_flash(session))
                elif len(parts) == 3 and parts[2] == "edit":
                    body = admin_unit_pages.form(con,session,csrf_input(session),code)
                else:
                    raise admin_units.CatalogError("Không tìm thấy trang.",404)
                self.send_html(base_page("Danh mục đơn vị hành chính",body,session,active_path="/admin-units"))
            else:
                if len(parts) == 3 and parts[2] == "edit":
                    con.execute("BEGIN")
                    admin_units.update_unit(con,session,data,code)
                else:
                    raise admin_units.CatalogError("Thao tác không được hỗ trợ.",405)
                con.commit()
                set_flash(session,"ok","Đã lưu danh mục hành chính.")
                self.redirect("/admin-units")
        except (admin_units.CatalogError, *INTEGRITY_ERRORS) as exc:
            con.rollback()
            status = exc.status if isinstance(exc,admin_units.CatalogError) else 409
            message = str(exc) if isinstance(exc,admin_units.CatalogError) else "Mã hoặc dữ liệu đơn vị bị trùng/không hợp lệ."
            body = f'<div class="container"><div class="notice err">{esc(message)}</div><a class="btn" href="/admin-units">Quay lại danh mục</a></div>'
            if status in (400,409) and data is not None and len(parts) == 3 and parts[2] == "edit":
                body = admin_unit_pages.form(con,session,csrf_input(session),code,message,data)
            self.send_html(base_page("Danh mục đơn vị hành chính",body,session,active_path="/admin-units"),status)
        finally:
            con.close()
        return True

    def log_message(self, format, *args):
        sys.stdout.write("%s - - [%s] %s\n" % (self.client_address[0], self.log_date_time_string(), format%args))

    def send_html(self, content, status=200, extra_headers=None):
        data = content.encode("utf-8")
        self.send_response(status); self.send_header("Content-Type","text/html; charset=utf-8"); self.send_header("Content-Length",str(len(data))); self.send_header("X-Frame-Options","DENY"); self.send_header("X-Content-Type-Options","nosniff")
        if extra_headers:
            for k,v in extra_headers: self.send_header(k,v)
        self.end_headers(); self.wfile.write(data)

    def redirect(self, location, headers=None):
        self.send_response(303); self.send_header("Location", location)
        if headers:
            for k,v in headers: self.send_header(k,v)
        self.end_headers()

    def require_session(self):
        sid,s = get_session(self)
        if not s:
            self.redirect("/"); return None, None
        s["user_id"] = int(s["user_id"])
        return sid,s

    def handle_ocop(self, path, query, session, data=None):
        """OCOP dispatch stays separate from the existing salt routes."""
        if not (path == "/ocop" or path.startswith("/ocop/")):
            return False
        import ocop_services as svc
        import ocop_pages
        con = db_conn()
        active = "/".join(path.split("/")[:3]) if path != "/ocop" else path
        if path == "/ocop/access":
            active = "/users"
        helpers = {"esc": esc, "csrf_input": csrf_input, "icon": icon}
        manual_committed = False
        try:
            if not ocop_available(con):
                raise svc.OcopError("Phân hệ OCOP chưa được khởi tạo trên cơ sở dữ liệu này.", 503)
            if path == "/ocop/manual":
                ocop_registry.require_internal(con, session)
                if data is None:
                    body = ocop_manual.page(con, session, csrf_input(session), query)
                    self.send_html(base_page("Nhập dữ liệu OCOP", body, session, active_path=path))
                else:
                    result = ocop_manual.save(con, session, data, metadata={
                        "ip_address": self.client_address[0],
                        "user_agent": re.sub(r"(?i)(password|token|cookie|authorization)\s*[:=]\s*\S+", "[redacted]", self.headers.get("User-Agent", ""))[:1000],
                    })
                    manual_committed = True
                    self.redirect(f'/ocop/manual?saved={result["product_id"]}')
                return True
            if path != "/ocop/access":
                svc.get_scope(con, session)
            if data is None:
                page = ocop_access_page(con, session) if path == "/ocop/access" else ocop_pages.render(path, query, session, con, helpers)
                if page is None:
                    raise svc.OcopError("Không tìm thấy trang OCOP.", 404)
                title, body = page
                body = body.replace('<div class="container ocop-page">', '<div class="container ocop-page">' + take_flash(session), 1)
                self.send_html(base_page(title, body, session, active_path=active))
                return True
            if path == "/ocop/access":
                save_ocop_access(con, session, data)
                destination = path
            else:
                parts = path.strip("/").split("/")
                if len(parts) not in (3, 4) or parts[1] not in ("entities", "products", "applications"):
                    raise svc.OcopError("Không tìm thấy thao tác OCOP.", 404)
                section = parts[1]
                singular = {"entities": "entity", "products": "product", "applications": "application"}[section]
                if len(parts) == 3 and parts[2] == "new":
                    result = getattr(svc, f"create_{singular}")(con, session, data)
                    rid = result["id"] if isinstance(result, dict) else result
                    destination = f"/ocop/{section}/{rid}"
                elif len(parts) == 4 and parts[2].isdigit():
                    rid, action = int(parts[2]), parts[3]
                    if action == "edit":
                        getattr(svc, f"update_{singular}")(con, session, rid, data)
                    elif action == "archive" and section in ("entities", "products"):
                        getattr(svc, f"archive_{singular}")(con, session, rid, data)
                    elif section == "applications" and action in ("submit", "start-review", "return", "eligible", "cancel"):
                        svc.application_action(con, session, rid, action, data)
                    else:
                        raise svc.OcopError("Không tìm thấy thao tác OCOP.", 404)
                    destination = f"/ocop/{section}/{rid}"
                else:
                    raise svc.OcopError("Không tìm thấy thao tác OCOP.", 404)
            set_flash(session, "ok", "Đã lưu thay đổi OCOP.")
            self.redirect(destination)
        except ocop_registry.RegistryError as exc:
            con.rollback()
            if path == "/ocop/manual":
                if data is not None:
                    ocop_registry.record_failure(con, session)
                body = ocop_manual.page(con, session, csrf_input(session), query, data=data, error=exc)
                self.send_html(base_page("Nhập dữ liệu OCOP", body, session, active_path=path), exc.status)
            else:
                self.send_html(base_page("Thông báo OCOP", f'<div class="container"><div class="notice err">{esc(exc)}</div></div>', session), exc.status)
        except svc.OcopError as exc:
            con.rollback()
            if path == "/ocop/manual" and data is not None and is_chi_cuc_user(session) and not manual_committed:
                ocop_registry.record_failure(con, session)
            status = getattr(exc, "status", 400)
            page = None
            if data is not None and status in (400, 409) and (path.endswith("/new") or path.endswith("/edit")):
                try:
                    page = ocop_pages.render(path, query, session, con, {**helpers, "form_data":data, "error":str(exc)})
                except (svc.OcopError, ValueError):
                    pass
            if page:
                title, body = page
            else:
                title = "Thông báo OCOP"
                body = f'<div class="container"><div class="notice err">{esc(exc)}</div><a class="btn" href="/ocop">Tổng quan OCOP</a></div>'
            self.send_html(base_page(title, body, session, active_path=active), status)
        except Exception:
            con.rollback()
            logging.exception("OCOP request failed")
            if path == "/ocop/manual" and data is not None and not manual_committed:
                ocop_registry.record_failure(con, session)
            self.send_html(base_page("Thông báo OCOP", '<div class="container"><div class="notice err">Không thể xử lý yêu cầu lúc này. Vui lòng thử lại hoặc liên hệ quản trị.</div></div>', session, active_path=active), 500)
        finally:
            con.close()
        return True

    def handle_ocop_t2_get(self, path, query, session):
        if path not in ("/ocop/import", "/ocop/import/preview", "/ocop/recognitions"):
            return False
        con = db_conn()
        try:
            if path == "/ocop/import":
                if parse_qs(query or "").get("reset", [""])[0] == "1":
                    session.pop("ocop_import_upload", None)
                    session.pop("ocop_import_preview", None)
                title, body, status = ocop_import_pages.import_page(session, con, {"csrf_input": csrf_input, "take_flash": take_flash, "esc": esc})
                self.send_html(base_page(title, body, session, active_path="/ocop/import"), status)
            elif path == "/ocop/import/preview":
                if not is_chi_cuc_user(session):
                    self.send_html(base_page("403", '<div class="container"><div class="notice err">Không có quyền.</div></div>', session, active_path="/ocop/import"), 403)
                    return True
                preview = session.get("ocop_import_preview")
                upload = session.get("ocop_import_upload")
                if not preview and upload:
                    preview = ocop_import.parse_ocop_workbook(upload["content"], upload["filename"], admin_units.unit_lookup(con), upload["selected_sheet"])
                    session["ocop_import_preview"] = preview
                    session.pop("ocop_import_upload", None)
                if not preview:
                    set_flash(session, "err", "Dữ liệu xem trước đã hết hạn. Hãy chọn lại file.")
                    self.redirect("/ocop/import")
                    return True
                title, body, status = ocop_import_pages.preview_page(session, preview, {"csrf_input": csrf_input, "take_flash": take_flash, "esc": esc})
                self.send_html(base_page(title, body, session, active_path="/ocop/import"), status)
            else:
                title, body, status = ocop_import_pages.recognitions_page(session, con, query, {"csrf_input": csrf_input, "take_flash": take_flash, "esc": esc})
                self.send_html(base_page(title, body, session, active_path="/ocop/recognitions"), status)
        except Exception as exc:
            logging.exception("OCOP T2 GET failed")
            self.send_html(base_page("Thông báo OCOP", f'<div class="container"><div class="notice err">{esc(exc)}</div><a class="btn" href="/ocop">Quay lại OCOP</a></div>', session, active_path="/ocop/import"), 400 if isinstance(exc, ocop_import.OcopImportError) else 500)
        finally:
            con.close()
        return True

    def handle_ocop_t2_post(self, path, session):
        if path not in ("/ocop/import", "/ocop/import/preview", "/ocop/import/confirm"):
            return False
        if not is_chi_cuc_user(session):
            self.send_html(base_page("403", '<div class="container"><div class="notice err">Chỉ tài khoản Chi cục được import OCOP.</div></div>', session, active_path="/ocop/import"), 403)
            return True
        try:
            if path == "/ocop/import":
                data, uploaded_files = parse_multipart_form(self)
                if not check_csrf(session, data):
                    raise ocop_import.OcopImportError("Phiên làm việc không hợp lệ. Vui lòng tải lại trang.")
                uploaded = uploaded_files.get("excel_file")
                if not uploaded:
                    raise ocop_import.OcopImportError("Vui lòng chọn file Excel OCOP.")
                filename = uploaded.get("filename") or ""
                try:
                    inspected = inspect_workbook(uploaded.get("content", b""), filename, "Loc")
                except WorkbookInspectionError as exc:
                    raise ocop_import.OcopImportError(str(exc)) from exc
                session.pop("ocop_import_preview", None)
                session["ocop_import_upload"] = {
                    "content": uploaded["content"],
                    "filename": filename,
                    "sheet_names": inspected["sheet_names"],
                    "recommended_sheet": inspected["recommended_sheet"],
                    "selected_sheet": inspected["recommended_sheet"],
                }
                self.redirect("/ocop/import")
                return True
            if path == "/ocop/import/preview":
                data = parse_body(self)
                if not check_csrf(session, data):
                    raise ocop_import.OcopImportError("Phiên làm việc không hợp lệ. Vui lòng tải lại trang.")
                upload = session.get("ocop_import_upload")
                if not upload:
                    set_flash(session, "err", "Chưa có file Excel. Hãy chọn file trước.")
                    self.redirect("/ocop/import")
                    return True
                sheet_name = str(data.get("sheet_name", "")).strip()
                if sheet_name not in upload.get("sheet_names", []):
                    raise ocop_import.OcopImportError("Sheet được chọn không còn trong workbook. Hãy chọn lại.")
                upload["selected_sheet"] = sheet_name
                con = db_conn()
                try:
                    preview = ocop_import.parse_ocop_workbook(upload["content"], upload["filename"], admin_units.unit_lookup(con), sheet_name)
                finally:
                    con.close()
                session["ocop_import_preview"] = preview
                session.pop("ocop_import_upload", None)
                self.redirect("/ocop/import/preview")
                return True
            data = parse_body(self)
            if not check_csrf(session, data):
                self.send_html(base_page("Lỗi", '<div class="container"><div class="notice err">Phiên làm việc không hợp lệ. Vui lòng tải lại trang.</div></div>', session, active_path="/ocop/import"), 403)
                return True
            preview = session.get("ocop_import_preview")
            upload = session.get("ocop_import_upload")
            if not preview and upload:
                con = db_conn()
                try:
                    preview = ocop_import.parse_ocop_workbook(upload["content"], upload["filename"], admin_units.unit_lookup(con), upload["selected_sheet"])
                finally:
                    con.close()
            if not preview:
                set_flash(session, "err", "Dữ liệu xem trước đã hết hạn. Hãy chọn lại file.")
                self.redirect("/ocop/import/preview")
                return True
            con = db_conn()
            try:
                con.execute("BEGIN")
                result = ocop_import.commit_ocop_preview(con, session, preview, data.get("mode", "publish"))
                con.commit()
            except Exception:
                con.rollback()
                raise
            finally:
                con.close()
            session.pop("ocop_import_upload", None)
            session.pop("ocop_import_preview", None)
            if result.get("already_published"):
                set_flash(session, "ok", "File này đã được đồng bộ trước đó; không tạo dữ liệu trùng.")
            elif data.get("mode") == "stage_only":
                set_flash(session, "ok", f"Đã lưu staging đợt #{result['batch_id']} để đối chiếu; chưa ghi dữ liệu OCOP chính thức.")
            else:
                set_flash(session, "ok", f"Đã đồng bộ {result['published_products']} sản phẩm OCOP từ đợt #{result['batch_id']}.")
            self.redirect("/ocop/recognitions" if data.get("mode") != "stage_only" else "/ocop/import")
        except ocop_import.OcopImportError as exc:
            set_flash(session, "err", str(exc))
            self.redirect("/ocop/import/preview" if path.endswith("confirm") else "/ocop/import")
        except Exception:
            logging.exception("OCOP T2 POST failed")
            set_flash(session, "err", "Không thể đồng bộ OCOP; transaction đã được hoàn tác.")
            self.redirect("/ocop/import/preview")
        return True

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path

        # =========================
        # ASSETS: CSS / JS / ẢNH
        # =========================
        if path in ASSET_FILES:
            filename, content_type = ASSET_FILES[path]

            try:
                with open(
                    os.path.join(BASE_DIR, "assets", filename),
                    "rb"
                ) as asset_file:
                    content = asset_file.read()

            except FileNotFoundError:
                self.send_error(404, "Asset not found")
                return

            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(content)))

            self.send_header(
                "Cache-Control",
                "no-cache"
                if filename not in {"quoc-huy.png", "LOGO-CCPTNT-TP.HCM_.jpg"}
                else "public, max-age=86400"
            )

            self.send_header(
                "X-Content-Type-Options",
                "nosniff"
            )

            self.end_headers()
            self.wfile.write(content)
            return


        # =========================
        # KIỂM TRA SESSION
        # =========================
        sid, session = get_session(self)


        # =========================
        # TRANG ĐĂNG NHẬP
        # =========================
        if path in ("/", "/index", "/index.html"):

            if session:
                self.redirect("/dashboard")

            else:
                self.send_html(login_page())

            return


        # =========================
        # ĐĂNG XUẤT
        # =========================
        if path == "/logout":

            if sid:
                with SESSION_LOCK:
                    SESSIONS.pop(sid, None)

            self.redirect(
                "/",
                [
                    (
                        "Set-Cookie",
                        "salt_session=; Path=/; Max-Age=0; HttpOnly; SameSite=Lax"
                    )
                ]
            )

            return


        # =========================
        # YÊU CẦU ĐĂNG NHẬP
        # =========================
        sid, session = self.require_session()

        if not session:
            return

        if path == "/ocop/export.xlsx":
            content, mime = export_ocop_xlsx(session, parsed.query)
            if content is None:
                self.send_html(base_page("Lỗi", '<div class="container"><div class="notice err">Máy chủ chưa có openpyxl.</div></div>', session), 500)
                return
            self.send_response(200)
            self.send_header("Content-Type", mime)
            self.send_header("Content-Disposition", 'attachment; filename="tra_cuu_ocop.xlsx"')
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
            return

        if self.handle_ocop_t2_get(path, parsed.query, session):
            return


        # =========================
        # DASHBOARD
        # =========================
        if (path == "/users" or path.startswith("/users/") or path.rstrip("/") == "/ocop/access") and not can_manage_users(session):
            self.send_html(base_page("403", '<div class="container"><div class="notice err">Không có quyền quản lý tài khoản.</div></div>', session), 403)
            return
        if self.handle_admin_units(path, parsed.query, session):
            return
        if self.handle_ocop(path, parsed.query, session):
            return
        if path == "/dashboard":
            self.send_html(landing_page(session, parsed.query))
            return


        # =========================
        # DANH SÁCH BÁO CÁO
        # =========================
        if path in ("/records/export.xlsx", "/salt/weekly/export.xlsx"):
            content, mime = export_weekly_xlsx(session, parsed.query)
            if content is None:
                self.send_html(base_page("Lỗi", '<div class="container"><div class="notice err">Không tìm thấy sheet hoặc máy chủ chưa có openpyxl.</div></div>', session), 404)
                return
            self.send_response(200)
            self.send_header("Content-Type", mime)
            self.send_header("Content-Disposition", 'attachment; filename="bao_cao_muoi_tuan.xlsx"')
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
            return

        if path == "/records":
            # The public salt screen is the weekly sheet lookup. The old
            # approval list remains available only through internal links.
            self.send_html(weekly_records_page(session, parsed.query))
            return

        if path == "/standard-data":
            p = standard_data_page(session)
            self.send_html(
                p if p else base_page(
                    "403",
                    "<div class='container'><div class='notice err'>Không có quyền.</div></div>",
                    session
                ),
                200 if p else 403
            )
            return

        if path == "/salt/weekly":
            self.send_html(weekly_records_page(session, parsed.query))
            return


        # =========================
        # NHẬP SỐ LIỆU
        # =========================
        if path == "/records/new":

            self.send_html(
                record_form_page(session)
            )

            return

        if path == "/import-excel":

            if parsed.query and parse_qs(parsed.query).get("reset", [""])[0] == "1":
                session.pop("weekly_import_upload", None)
                session.pop("weekly_import_preview", None)

            p = import_excel_page(session, parsed.query)

            self.send_html(
                p
                if p
                else base_page(
                    "403",
                    """
                    <div class='container'>
                        <div class='notice err'>
                            Không có quyền.
                        </div>
                    </div>
                    """,
                    session
                ),
                200 if p else 403
            )

            return
        if path == "/import-excel/preview":
            p = import_preview_page(session)
            if p:
                self.send_html(p)
            else:
                self.redirect("/import-excel")
            return
        # =========================
        # TÀI KHOẢN ĐƠN VỊ
        # =========================
        if path == "/users":

            p = users_page(session)

            self.send_html(
                p
                if p
                else base_page(
                    "403",
                    """
                    <div class='container'>
                        <div class='notice err'>
                            Không có quyền.
                        </div>
                    </div>
                    """,
                    session
                ),
                200 if p else 403
            )

            return


        # =========================
        # ĐỔI MẬT KHẨU
        # =========================
        if path == "/change-password":

            self.send_html(
                change_password_page(session)
            )

            return


        # =========================
        # XUẤT EXCEL
        # =========================
        if path == "/export.xlsx":

            filters = filters_from_query(
                parsed.query,
                session
            )

            content, mime = export_xlsx(
                session,
                filters
            )

            if content is None:

                self.send_html(
                    base_page(
                        "Lỗi",
                        """
                        <div class='container'>
                            <div class='notice err'>
                                Máy chủ chưa có openpyxl.
                                Cài bằng: pip install openpyxl
                            </div>
                        </div>
                        """,
                        session
                    ),
                    500
                )

                return


            self.send_response(200)

            self.send_header(
                "Content-Type",
                mime
            )

            self.send_header(
                "Content-Disposition",
                'attachment; filename="tong_hop_so_lieu_muoi.xlsx"'
            )

            self.send_header(
                "Content-Length",
                str(len(content))
            )

            self.end_headers()

            self.wfile.write(content)

            return


        # =========================
        # TÁCH URL
        # =========================
        parts = [
            p
            for p in path.split("/")
            if p
        ]


        # =========================================
        # MỞ TRANG SỬA TÀI KHOẢN
        # URL: /users/9/edit
        # =========================================
        if (
            len(parts) == 3
            and parts[0] == "users"
            and parts[1].isdigit()
            and parts[2] == "edit"
        ):

            # Chỉ Chi cục được sửa tài khoản
            if not can_manage_users(session):

                self.send_html(
                    base_page(
                        "403",
                        """
                        <div class='container'>
                            <div class='notice err'>
                                Không có quyền.
                            </div>
                        </div>
                        """,
                        session
                    ),
                    403
                )

                return


            uid = int(parts[1])


            con = db_conn()

            user = get_user(con, uid)

            con.close()


            if not user:

                self.send_html(
                    base_page(
                        "404",
                        """
                        <div class='container'>
                            <div class='notice err'>
                                Không tìm thấy tài khoản.
                            </div>
                        </div>
                        """,
                        session
                    ),
                    404
                )

                return


            self.send_html(
                user_edit_page(
                    session,
                    user
                )
            )

            return


        # =========================================
        # CHI TIẾT / SỬA BÁO CÁO
        # =========================================
        if (
            len(parts) >= 2
            and parts[0] == "records"
            and parts[1].isdigit()
        ):

            rid = int(parts[1])


            # =========================
            # XEM CHI TIẾT BÁO CÁO
            # /records/1
            # =========================
            if len(parts) == 2:

                p = detail_page(
                    session,
                    rid
                )

                self.send_html(
                    p
                    if p
                    else base_page(
                        "404",
                        """
                        <div class='container'>
                            <div class='notice err'>
                                Không tìm thấy hoặc không có quyền.
                            </div>
                        </div>
                        """,
                        session
                    ),
                    200 if p else 404
                )

                return


            # =========================
            # SỬA BÁO CÁO
            # /records/1/edit
            # =========================
            if (
                len(parts) == 3
                and parts[2] == "edit"
            ):

                con = db_conn()

                r = con.execute(
                    "SELECT * FROM records WHERE id=?",
                    (rid,)
                ).fetchone()

                con.close()


                if (
                    not r
                    or not can_access_record(
                        session,
                        r
                    )
                    or (
                        session["role"] == ROLE_UNIT
                        and r["status"]
                        not in (
                            STATUS_DRAFT,
                            STATUS_RETURNED
                        )
                    )
                ):

                    self.send_html(
                        base_page(
                            "403",
                            """
                            <div class='container'>
                                <div class='notice err'>
                                    Không có quyền sửa báo cáo này.
                                </div>
                            </div>
                            """,
                            session
                        ),
                        403
                    )

                    return


                self.send_html(
                    record_form_page(
                        session,
                        r
                    )
                )

                return


            # =========================
            # TRẢ BÁO CÁO CHỈNH SỬA
            # /records/1/return
            # =========================
            if (
                len(parts) == 3
                and parts[2] == "return"
            ):

                p = return_page(
                    session,
                    rid
                )

                self.send_html(
                    p
                    if p
                    else base_page(
                        "403",
                        """
                        <div class='container'>
                            <div class='notice err'>
                                Không có quyền.
                            </div>
                        </div>
                        """,
                        session
                    ),
                    200 if p else 403
                )

                return


        # =========================
        # KHÔNG TÌM THẤY TRANG
        # =========================
        self.send_html(
            base_page(
                "404",
                """
                <div class='container'>
                    <div class='notice err'>
                        Không tìm thấy trang.
                    </div>
                </div>
                """,
                session
            ),
            404
        )

        return
    def do_POST(self):

        parsed = urlparse(self.path)
        path = parsed.path

        # OCOP import has its own multipart/preview/commit flow.
        if path in ("/ocop/import", "/ocop/import/preview", "/ocop/import/confirm"):
            sid, session = self.require_session()
            if session:
                self.handle_ocop_t2_post(path, session)
            return

        # Reject oversized OCOP forms before reading their body. Evidence uploads
        # belong to a later phase and are deliberately not accepted here.
        if path == "/ocop" or path.startswith("/ocop/"):
            try:
                length = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                length = -1
            if length < 0 or length > 256 * 1024 or self.headers.get("Transfer-Encoding"):
                self.close_connection = True
                self.send_html(base_page("Yêu cầu không hợp lệ", '<div class="container">Dữ liệu gửi không hợp lệ hoặc quá lớn.</div>'), 413)
                return
            if not self.headers.get("Content-Type", "").lower().startswith("application/x-www-form-urlencoded"):
                self.close_connection = True
                self.send_html(base_page("Yêu cầu không hợp lệ", '<div class="container">Định dạng biểu mẫu chưa được hỗ trợ.</div>'), 415)
                return

        uploaded_files = {}


        if "multipart/form-data" in self.headers.get(
            "Content-Type",
            ""
        ):

            data, uploaded_files = parse_multipart_form(self)

        else:

            data = parse_body(self)
        if path=="/login":
            username=data.get("username","").strip(); password=data.get("password","")
            con=db_conn(); u=find_user_by_username(con, username); con.close()
            if not u or not verify_password(password,u["password_hash"]): self.send_html(login_page("Tên đăng nhập hoặc mật khẩu không đúng."),401); return
            if not u["active"]: self.send_html(login_page("Tài khoản không hoạt động. Vui lòng liên hệ Chi cục để được xử lý."),401); return
            sid=new_session(u); self.redirect("/dashboard",[("Set-Cookie",f"salt_session={sid}; Path=/; HttpOnly; SameSite=Lax")]); return
        sid,session=self.require_session()
        if not session: return
        if not check_csrf(session,data): self.send_html(base_page("Lỗi","<div class='container'><div class='notice err'>Phiên làm việc không hợp lệ. Vui lòng tải lại trang.</div></div>",session),400); return
        if (path == "/users" or path.startswith("/users/") or path.rstrip("/") == "/ocop/access") and not can_manage_users(session):
            self.send_html(base_page("403", '<div class="container"><div class="notice err">Không có quyền quản lý tài khoản.</div></div>', session), 403)
            return
        if self.handle_admin_units(path, parsed.query, session, data):
            return
        if self.handle_ocop(path, parsed.query, session, data):
            return
        if path == "/import-excel":
            if not is_chi_cuc_user(session):
                self.send_html(base_page("403", '<div class="container"><div class="notice err">Không có quyền import.</div></div>', session), 403)
                return
            uploaded = uploaded_files.get("excel_file")
            if not uploaded:
                set_flash(session, "err", "Vui lòng chọn file Excel.")
                self.redirect("/import-excel")
                return
            filename = uploaded.get("filename") or ""
            try:
                inspected = inspect_workbook(uploaded.get("content", b""), filename)
            except WorkbookInspectionError as exc:
                set_flash(session, "err", str(exc))
                self.redirect("/import-excel")
                return
            # A new file replaces any unfinished preview/upload from the same
            # session; only the selected workbook is retained for step 2.
            session.pop("weekly_import_preview", None)
            session["weekly_import_upload"] = {
                "content": uploaded["content"],
                "filename": filename,
                "sheet_names": inspected["sheet_names"],
                "recommended_sheet": inspected["recommended_sheet"],
                "selected_sheet": inspected["recommended_sheet"],
            }
            self.redirect("/import-excel")
            return
        if path == "/import-excel/preview":
            if not is_chi_cuc_user(session):
                self.send_html(base_page("403", '<div class="container"><div class="notice err">Không có quyền import.</div></div>', session), 403)
                return
            upload = session.get("weekly_import_upload")
            if not upload:
                set_flash(session, "err", "Chưa có file Excel. Hãy chọn file trước.")
                self.redirect("/import-excel")
                return
            sheet_name = str(data.get("sheet_name", "")).strip()
            if sheet_name not in upload.get("sheet_names", []):
                set_flash(session, "err", "Sheet được chọn không còn trong workbook. Hãy chọn lại.")
                self.redirect("/import-excel")
                return
            report_date = str(data.get("report_date", "")).strip()
            upload["selected_sheet"] = sheet_name
            upload["report_date"] = report_date
            con = db_conn()
            try:
                preview = parse_weekly_workbook(
                    upload["content"], report_date, sheet_name, upload["filename"],
                    admin_units.unit_lookup(con),
                )
            except WeeklyImportError as exc:
                set_flash(session, "err", str(exc))
                self.redirect("/import-excel")
                return
            finally:
                con.close()
            session["weekly_import_preview"] = preview
            session.pop("weekly_import_upload", None)
            self.redirect("/import-excel/preview")
            return
        if path == "/import-excel/confirm":
            if not is_chi_cuc_user(session):
                self.send_html(base_page("403", "<div class='container'><div class='notice err'>Không có quyền import.</div></div>", session), 403)
                return
            preview = session.get("weekly_import_preview")
            if not preview:
                set_flash(session, "err", "Dữ liệu xem trước đã hết hạn. Hãy chọn lại file.")
                self.redirect("/import-excel")
                return
            con = db_conn()
            try:
                con.execute("BEGIN")
                admin_units.validate_weekly_units(con, preview)
                result = commit_weekly_preview(con, session, preview, data.get("mode", "skip"), now_text())
                con.commit()
            except (WeeklyImportError, admin_units.CatalogError) as exc:
                con.rollback()
                set_flash(session, "err", str(exc))
                self.redirect("/import-excel/preview")
                return
            except INTEGRITY_ERRORS:
                con.rollback()
                set_flash(session, "err", "Dữ liệu trùng hoặc không còn hợp lệ. Hãy tải lại file để kiểm tra.")
                self.redirect("/import-excel/preview")
                return
            finally:
                con.close()
            session.pop("weekly_import_preview", None)
            set_flash(session, "ok", f"Đã import tuần {preview['week_code']}: {result['inserted']} mới, {result['updated']} cập nhật, {result['skipped']} bỏ qua.")
            self.redirect(f"/salt/weekly?week={quote(preview['week_code'])}")
            return
        if path=="/records/new":
            ok,res=save_record(session,data)
            if ok: set_flash(session,"ok","Đã lưu bản nháp."); self.redirect(f"/records/{res}")
            else: self.send_html(record_form_page(session,None,res),400)
            return
        if path=="/users/new":
            if not can_manage_users(session): self.send_html(base_page("403","<div class='container'><div class='notice err'>Không có quyền.</div></div>",session),403); return
            username = data.get("username", "").strip()
            password = data.get("password", "")
            role = data.get("role", ROLE_UNIT)
            if role not in ROLE_LABELS or len(password) < 6 or not username:
                self.send_html(base_page("Lỗi", '<div class="container">Thông tin tài khoản chưa hợp lệ.</div>', session),400)
                return
            con = db_conn()
            try:
                con.execute("BEGIN")
                admin_units.require_admin(con,session)
                unit_name, unit_code = admin_units.account_unit(con,role,data)
                uid = create_user(con,username,hash_password(password),role,unit_name,now_text())
                admin_units.assign_account(con,session["user_id"],uid,unit_code)
                user_profiles.save_account_profile(con,uid,data)
                con.commit()
                set_flash(session,"ok","Đã tạo tài khoản và phạm vi địa bàn.")
            except admin_units.CatalogError as exc:
                con.rollback()
                self.send_html(base_page("Lỗi",f'<div class="container"><div class="notice err">{esc(exc)}</div></div>',session),exc.status)
                return
            except user_profiles.ProfileError as exc:
                con.rollback()
                self.send_html(base_page("Lỗi",f'<div class="container"><div class="notice err">{esc(exc)}</div><a href="/users">Quay lại tài khoản</a></div>',session),400)
                return
            except INTEGRITY_ERRORS:
                con.rollback()
                self.send_html(base_page("Lỗi",'<div class="container">Tên đăng nhập đã tồn tại.</div>',session),409)
                return
            finally:
                con.close()
            self.redirect("/users")
            return
        if path=="/change-password":
            old=data.get("old",""); new=data.get("new",""); new2=data.get("new2","")
            con=db_conn(); u=get_user(con, session["user_id"])
            if not verify_password(old,u["password_hash"]): con.close(); self.send_html(change_password_page(session,"Mật khẩu hiện tại không đúng."),400); return
            if new!=new2 or len(new)<8: con.close(); self.send_html(change_password_page(session,"Mật khẩu mới phải trùng nhau và có ít nhất 8 ký tự."),400); return
            set_user_password(con, session["user_id"], hash_password(new)); con.commit(); con.close(); set_flash(session,"ok","Đã đổi mật khẩu."); self.redirect("/dashboard"); return
        parts=[p for p in path.split('/') if p]
        if (
            len(parts) == 3
            and parts[0] == "users"
            and parts[1].isdigit()
        ):
            if not can_manage_users(session):
                self.send_html(
                    base_page(
                        "403",
                        "<div class='container'><div class='notice err'>Không có quyền.</div></div>",
                        session
                    ),
                    403
                )
                return

            uid = int(parts[1])
            action = parts[2]

            if action in ("activate", "deactivate"):
                if uid == session["user_id"] and action == "deactivate":
                    self.send_html(base_page("Không thể ngưng kích hoạt", '<div class="container">Không thể ngưng kích hoạt tài khoản đang đăng nhập.</div>', session), 400)
                    return
                con = db_conn()
                try:
                    if not get_user(con, uid):
                        self.send_error(404)
                        return
                    repositories.set_user_active(con, uid, action == "activate")
                    con.commit()
                    invalidate_user_sessions(uid)
                finally:
                    con.close()
                set_flash(session, "ok", "Đã kích hoạt lại tài khoản." if action == "activate" else "Đã ngưng kích hoạt tài khoản.")
                self.redirect("/users")
                return

            # ======================
            # SỬA TÀI KHOẢN
            # ======================
            if action == "edit":

                con = db_conn()

                user = get_user(con, uid)

                if not user:
                    con.close()
                    self.redirect("/users")
                    return

                username = data.get("username", "").strip()
                unit_name = data.get("unit_name", "").strip()
                role = data.get("role", ROLE_UNIT)
                password = data.get("password", "")
                active = 1 if data.get("active") == "1" else 0

                if not username:
                    con.close()
                    self.send_html(
                        user_edit_page(
                            session,
                            user,
                            "Vui lòng nhập đầy đủ thông tin."
                        ),
                        400
                    )
                    return

                if role not in tuple(ROLE_LABELS):
                    role = ROLE_UNIT

                # Không cho admin đang đăng nhập tự khóa chính mình
                if uid == session["user_id"] and (not active or role != ROLE_ADMIN):
                    con.close()
                    self.send_html(user_edit_page(session, user, "Không thể tự ngưng kích hoạt hoặc hạ quyền tài khoản đang đăng nhập."), 400)
                    return

                try:
                    con.execute("BEGIN")
                    admin_units.require_admin(con,session)
                    unit_name, unit_code = admin_units.account_unit(con,role,data)
                    new_password_hash = None
                    if password:
                        if len(password) < 6:
                            con.close()
                            self.send_html(
                                user_edit_page(
                                    session,
                                    user,
                                    "Mật khẩu phải có ít nhất 6 ký tự."
                                ),
                                400
                            )
                            return

                        new_password_hash = hash_password(password)

                    update_user_account(
                        con, uid, username, role, unit_name, active, new_password_hash
                    )

                    admin_units.assign_account(con,session["user_id"],uid,unit_code)
                    user_profiles.save_account_profile(con,uid,data)
                    con.commit()
                    invalidate_user_sessions(uid)

                    set_flash(
                        session,
                        "ok",
                        "Đã cập nhật tài khoản."
                    )

                except admin_units.CatalogError as exc:
                    con.rollback()
                    self.send_html(user_edit_page(session,user,str(exc)),exc.status)
                    return
                except user_profiles.ProfileError as exc:
                    con.rollback()
                    self.send_html(user_edit_page(session,user,str(exc),profile_data=data),400)
                    return
                except INTEGRITY_ERRORS:
                    con.rollback()

                    set_flash(
                        session,
                        "err",
                        "Tên đăng nhập đã tồn tại."
                    )

                finally:
                    con.close()

                self.redirect("/users")
                return

            # ======================
            # XÓA / KHÓA TÀI KHOẢN
            # ======================
            if action == "delete":

                if uid == session["user_id"]:
                    set_flash(
                        session,
                        "err",
                        "Không thể xóa tài khoản đang đăng nhập."
                    )
                    self.redirect("/users")
                    return

                con = db_conn()

                user = get_user(con, uid)

                if not user:
                    con.close()
                    self.redirect("/users")
                    return

                record_count = con.execute(
                    "SELECT COUNT(*) FROM records WHERE created_by=?",
                    (uid,)
                ).fetchone()[0]

                log_count = con.execute(
                    "SELECT COUNT(*) FROM audit_logs WHERE user_id=?",
                    (uid,)
                ).fetchone()[0]

                # Nếu tài khoản đã có dữ liệu -> chỉ khóa
                if record_count > 0 or log_count > 0 or ocop_user_links(con, uid) > 0:

                    con.execute(
                        "UPDATE users SET active=0 WHERE id=?",
                        (uid,)
                    )

                    con.commit()
                    con.close()

                    set_flash(
                        session,
                        "ok",
                        "Tài khoản đã có dữ liệu nên hệ thống đã khóa thay vì xóa."
                    )

                else:

                    delete_user_account(con, uid)

                    con.commit()
                    con.close()

                    set_flash(
                        session,
                        "ok",
                        "Đã xóa tài khoản."
                    )

                invalidate_user_sessions(uid)
                self.redirect("/users")
                return
        if len(parts)>=3 and parts[0]=="records" and parts[1].isdigit():
            rid=int(parts[1]); action=parts[2]
            if action=="edit":
                con=db_conn(); r=con.execute("SELECT * FROM records WHERE id=?",(rid,)).fetchone(); con.close()
                if not r or not can_access_record(session,r): self.send_html(base_page("403","<div class='container'><div class='notice err'>Không có quyền.</div></div>",session),403); return
                ok,res=save_record(session,data,rid)
                if ok: set_flash(session,"ok","Đã cập nhật số liệu."); self.redirect(f"/records/{rid}")
                else: self.send_html(record_form_page(session,r,res),400)
                return
            con=db_conn(); r=con.execute("SELECT * FROM records WHERE id=?",(rid,)).fetchone()
            if not r or not can_access_record(session,r): con.close(); self.send_html(base_page("403","<div class='container'><div class='notice err'>Không có quyền.</div></div>",session),403); return
            if action=="submit":
                if session["role"]!=ROLE_UNIT or r["status"] not in (STATUS_DRAFT,STATUS_RETURNED): con.close(); self.send_html(base_page("Lỗi","<div class='container'><div class='notice err'>Không thể gửi báo cáo này.</div></div>",session),400); return
                con.execute("UPDATE records SET status='submitted',submitted_at=?,updated_at=? WHERE id=?",(now_text(),now_text(),rid)); add_audit(con,rid,session["user_id"],"Gửi Chi cục","Đề nghị kiểm tra/phê duyệt"); con.commit(); con.close(); set_flash(session,"ok","Đã gửi số liệu lên Chi cục."); self.redirect(f"/records/{rid}"); return
            if action=="approve":
                if not can_review_records(session) or r["status"]!=STATUS_SUBMITTED:
                    con.close()
                    self.send_html(base_page("Lỗi","<div class='container'><div class='notice err'>Không thể duyệt báo cáo này.</div></div>",session),400)
                    return
                try:
                    approved_at = now_text()
                    con.execute(
                        "UPDATE records SET status='approved',approved_at=?,reviewer_note='',updated_at=? WHERE id=?",
                        (approved_at, approved_at, rid),
                    )
                    approved_record = con.execute(
                        "SELECT * FROM records WHERE id=?",
                        (rid,),
                    ).fetchone()
                    sync_standard_salt_record(con, approved_record)
                    add_audit(
                        con, rid, session["user_id"],
                        "Phê duyệt",
                        "Chấp nhận số liệu và đồng bộ CSDL chuẩn DN_SanLuongMuoi",
                    )
                    con.commit()
                except Exception as e:
                    con.rollback()
                    con.close()
                    self.send_html(
                        base_page(
                            "Lỗi",
                            f"<div class='container'><div class='notice err'>Không thể chuẩn hóa dữ liệu khi phê duyệt: {esc(e)}</div></div>",
                            session,
                        ),
                        500,
                    )
                    return
                con.close()
                set_flash(session,"ok","Đã phê duyệt và đồng bộ dữ liệu chuẩn.")
                self.redirect(f"/records/{rid}")
                return
            if action=="return":
                if not can_review_records(session) or r["status"]!=STATUS_SUBMITTED: con.close(); self.send_html(base_page("Lỗi","<div class='container'><div class='notice err'>Không thể trả báo cáo này.</div></div>",session),400); return
                note=data.get("reviewer_note","").strip()
                if not note: con.close(); self.send_html(return_page(session,rid),400); return
                con.execute("UPDATE records SET status='returned',reviewer_note=?,updated_at=? WHERE id=?",(note,now_text(),rid)); add_audit(con,rid,session["user_id"],"Yêu cầu chỉnh sửa",note); con.commit(); con.close(); set_flash(session,"ok","Đã gửi yêu cầu chỉnh sửa cho đơn vị."); self.redirect(f"/records/{rid}"); return
            con.close()
        self.send_html(base_page("404","<div class='container'><div class='notice err'>Không tìm thấy thao tác.</div></div>",session),404)


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Hệ thống quản lý nghiệp vụ Chi cục Phát triển nông thôn")
    parser.add_argument("--host", default=HOST)
    parser.add_argument("--port", type=int, default=PORT)
    args = parser.parse_args()
    init_db()
    con_ready = db_conn()
    try:
        ready = con_ready.execute(
            "SELECT to_regclass('app.ocop_recognitions') AS recognitions, "
            "to_regclass('staging.ocop_import_batches') AS batches, "
            "to_regclass('staging.ocop_import_rows') AS import_rows"
        ).fetchone()
        if not all(ready.values()):
            raise RuntimeError("Thiếu migration 013_ocop_legacy_import.sql. Hãy áp dụng migration 013 trước khi chạy web.")
    finally:
        con_ready.close()
    try:
        con_check = db_conn()
        tmp_count = con_check.execute("SELECT COUNT(*) FROM DM_DonViHanhChinh WHERE Ma_DonViHanhChinh LIKE 'TMP%'").fetchone()[0]
        official_count = len(admin_units.units(con_check, active_only=True, communes_only=True))
        con_check.close()
        synced = sync_all_approved_records()
        print(f"Đã đồng bộ dữ liệu chuẩn từ {synced} báo cáo đã duyệt.")
        print(f"Danh mục hành chính chính thức: {official_count} đơn vị; mã TMP còn lại: {tmp_count}.")
    except Exception as e:
        print(f"Cảnh báo: chưa đồng bộ được dữ liệu chuẩn: {e}")
    server=ThreadingHTTPServer((args.host,args.port),Handler)
    print("="*72)
    print("HỆ THỐNG QUẢN LÝ NGHIỆP VỤ CHI CỤC PHÁT TRIỂN NÔNG THÔN")
    settings = load_settings()
    print(f"Database PostgreSQL: {settings.dbname} @ {settings.host}:{settings.port}")
    print(f"Đang chạy tại: http://127.0.0.1:{args.port}")
    print(f"Trong mạng LAN: http://<IP-máy-chủ>:{args.port}")
    print("Tài khoản demo Chi cục: chicuc / 123456")
    print("Nhấn Ctrl+C để dừng.")
    print("="*72)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nĐã dừng máy chủ.")
    finally:
        server.server_close()

if __name__=="__main__":
    main()
