"""Small, shared helpers for safe server-side Excel inspection."""

from __future__ import annotations

import io


MAX_FILE_BYTES = 20 * 1024 * 1024


class WorkbookInspectionError(ValueError):
    """Raised when an uploaded workbook cannot be inspected safely."""


def inspect_workbook(file_bytes: bytes, filename: str, preferred_sheet: str | None = None):
    """Validate an .xlsx upload and return only its workbook metadata.

    The workbook is opened in read-only mode and always closed before the
    function returns. No worksheet rows are parsed or retained here.
    """
    if not file_bytes:
        raise WorkbookInspectionError("File Excel rỗng.")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise WorkbookInspectionError("File Excel quá lớn (tối đa 20 MB).")
    if not str(filename or "").lower().endswith(".xlsx"):
        raise WorkbookInspectionError("Chỉ hỗ trợ file .xlsx.")
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise WorkbookInspectionError("Máy chủ chưa cài openpyxl.") from exc

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookInspectionError(f"Không đọc được file Excel: {exc}") from exc
    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            raise WorkbookInspectionError("Workbook không có sheet nào.")
        active_sheet = workbook.active.title if workbook.active else sheet_names[0]
        preferred = str(preferred_sheet or "").strip()
        recommended = (
            preferred if preferred in sheet_names
            else active_sheet if active_sheet in sheet_names
            else sheet_names[0]
        )
        return {
            "filename": str(filename or ""),
            "sheet_names": sheet_names,
            "recommended_sheet": recommended,
        }
    finally:
        workbook.close()
